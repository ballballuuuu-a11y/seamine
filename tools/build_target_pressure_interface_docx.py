"""生成《舰艇水压场特性仿真数据结构与函数说明》Word 文档。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

from build_minesweeping_interface_docx import (
    BLUE,
    CAUTION,
    DARK_BLUE,
    MUTED,
    NAVY,
    add_body,
    add_bullet,
    add_callout,
    add_code_block,
    add_heading,
    add_numbered,
    add_numbering_definitions,
    add_page_break,
    add_page_field,
    add_table,
    configure_page,
    configure_styles,
    set_paragraph_border,
    set_run_font,
)


# 文档输出路径。
PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DOCX = PROJECT_ROOT / "docs" / "舰艇水压场特性仿真数据结构与函数说明.docx"


def configure_header_footer(document: Document) -> None:
    """配置水压场接口说明的运行页眉和页脚。"""

    section = document.sections[0]
    section.first_page_header.paragraphs[0].text = ""

    # 后续页面显示简洁的文档名称和内容类别。
    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(3)
    header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = header.add_run("舰艇水压场特性仿真")
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    header.add_run("\t")
    right = header.add_run("数据结构与函数说明")
    set_run_font(right, size=8.5, color=MUTED)
    set_paragraph_border(header, side="bottom", color="D8E2E8", size=5)

    # 页脚只显示页码，避免干扰正文阅读。
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(footer)


def build_cover(document: Document) -> None:
    """创建技术参考手册首页。"""

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(4)
    kicker_run = kicker.add_run("调用接口参考")
    set_run_font(kicker_run, size=10.5, color=CAUTION, bold=True)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("舰艇水压场特性仿真")
    set_run_font(title_run, size=28, color=NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle_run = subtitle.add_run("数据结构、成员含义与函数输入输出说明")
    set_run_font(subtitle_run, size=14, color=DARK_BLUE)

    metadata = [
        ("对应类", "TargetPressureFieldModel"),
        ("接口文件", "TargetPressureFieldModel.h"),
        ("实现文件", "TargetPressureFieldModel.cpp"),
        ("坐标约定", "ENU：x 向东、y 向北、z 向上，水下 z 为负"),
        ("主要用途", "单点压力、固定观测点时序、二维/三维空间网格"),
        ("更新时间", "2026-07-30"),
    ]
    for label, value in metadata:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        label_run = paragraph.add_run(f"{label}：")
        set_run_font(label_run, size=10.5, color=NAVY, bold=True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=10.5, color=NAVY)

    rule = document.add_paragraph()
    rule.paragraph_format.space_before = Pt(9)
    rule.paragraph_format.space_after = Pt(18)
    set_paragraph_border(rule, side="bottom", color=BLUE, size=12)

    add_callout(
        document,
        "一句话理解",
        "先用 TargetParameter 描述舰艇和海洋环境，再调用 calculate、simulate 或 simulateGrid；每个计算点的结果都保存在 PressureFieldSample 中。",
    )
    add_page_break(document)


def build_document() -> Path:
    """构建完整的数据结构与函数说明文档。"""

    document = Document()
    document.core_properties.title = "舰艇水压场特性仿真数据结构与函数说明"
    document.core_properties.subject = "TargetPressureFieldModel 数据结构、成员与函数接口"
    document.core_properties.author = "seamine 项目组"
    document.core_properties.last_modified_by = "seamine 项目组"
    document.core_properties.keywords = "舰艇水压场, 数据结构, 函数接口, C++"
    document.core_properties.comments = "按当前工程源码与测试整理的通俗调用说明。"
    configure_styles(document)
    configure_page(document)
    configure_header_footer(document)
    bullet_id, decimal_id = add_numbering_definitions(document)
    build_cover(document)

    add_heading(document, "1. 先了解整体调用方式", level=1)
    add_body(
        document,
        "该模型根据目标类型、船长弗劳德数和水深弗劳德数自动选择工况。低速工况使用回转体偶极子和刚盖边界；水面舰艇在 H/L≤0.3 的亚临界条件下使用线性深度平均兴波近似。调用人员只需要正确填写目标参数、观测位置和采样方式。",
    )
    add_heading(document, "1.1 最常用的调用步骤", level=2)
    for text in [
        "创建 TargetParameter，填写目标尺寸、海深、航速、初始位置、航向和环境参数。",
        "使用带参数的构造函数创建模型，或对默认构造的空模型调用 setTarget。",
        "单点计算调用 calculate；固定观测点时序调用 simulate；空间分布调用 simulateGrid。",
        "从 PressureFieldSample 中读取实际工况、各压力分量、动态压力、总表压、目标位置和相对距离。",
    ]:
        add_numbered(document, text, decimal_id)

    add_heading(document, "1.2 四个必须记住的口径", level=2)
    add_callout(
        document,
        "配置状态",
        "默认构造函数只创建空模型，不能立即计算；必须先调用 setTarget，否则读取或计算时会抛出 std::logic_error。",
        accent=CAUTION,
    )
    add_callout(
        document,
        "坐标口径",
        "静水面为 z=0，海床为 z=-waterDepth。水下位置的 z 通常为负值，而 PressureFieldSample.observationDepth 输出正的水深数值。",
        accent=CAUTION,
    )
    add_callout(
        document,
        "压力口径",
        "dynamicPressure 等于本体、自由液面和海床修正压力之和；totalGaugePressure 等于静水表压与 dynamicPressure 之和，压力单位均为 Pa。",
        accent=CAUTION,
    )
    add_callout(
        document,
        "模型边界",
        "观测点不能在目标等效椭球内部、静水面以上、海床以下，也不能距离目标或镜像点过近。",
        accent=CAUTION,
    )

    add_heading(document, "2. 坐标、方向与单位约定", level=1)
    add_table(
        document,
        ["项目", "约定", "通俗解释"],
        [
            ["全局坐标", "x 东、y 北、z 上", "静水面 z=0，海床 z=-waterDepth"],
            ["航向角", "东向 0°，逆时针为正", "90° 表示向北航行"],
            ["目标坐标", "纵向沿舰首，横向指向左舷侧", "用于输出纵向和横向相对距离"],
            ["尺寸、位置和距离", "m", "长度、宽度、吃水、海深和坐标统一使用米"],
            ["航速", "m/s", "目标沿水平航向匀速直线运动"],
            ["时间", "s", "相对零时刻的仿真时间，可以为负"],
            ["采样率", "Hz", "每秒生成的样本数量"],
            ["压力", "Pa", "静水压力、动态压力和总表压统一使用帕"],
            ["海水密度", "kg/m³", "默认值 1025，代表常用海水密度"],
        ],
        [1900, 2800, 4660],
        font_size=8.7,
    )
    add_page_break(document)

    add_heading(document, "3. 数据结构总览", level=1)
    add_table(
        document,
        ["数据结构", "主要作用", "什么时候使用"],
        [
            ["TargetType", "区分水面舰艇和潜艇", "填写 TargetParameter.type"],
            ["PressureFieldRegime", "表示自动识别的水动力工况", "读取 regime() 或计算结果"],
            ["ModelOptions", "保存工况判定阈值", "调整低速、浅水和兴波衰减判据"],
            ["Vector3", "保存三维位置或三维位移", "目标位置、观测点、网格边界和结果"],
            ["TargetParameter", "完整描述目标、运动状态和海洋环境", "构造模型或调用 setTarget"],
            ["PressureFieldSample", "保存一个点、一个时刻的完整压力结果", "calculate、simulate、simulateGrid 的输出"],
            ["GridParameter", "描述规则二维或三维空间网格", "调用 simulateGrid"],
        ],
        [2500, 4000, 2860],
        font_size=8.7,
    )

    add_heading(document, "4. TargetType：目标类型", level=1)
    add_table(
        document,
        ["枚举值", "含义", "当前计算差异"],
        [
            ["SurfaceShip", "水面舰艇", "按 FrL、H/L 和 FrH 选择低速或浅水亚临界模型"],
            ["Submarine", "潜艇", "按低速条件和兴波衰减量选择无兴波模型"],
        ],
        [2600, 3300, 3460],
        font_size=9.0,
    )
    add_callout(
        document,
        "容易误解",
        "type 会改变工况选择和几何校验，但不会代替调用方填写尺寸与潜深。潜艇应优先填写 height，并把 initialPosition.z 与 waterDepth 设置为符合完整浸没条件的数值。",
    )

    add_heading(document, "5. Vector3：三维坐标或位移", level=1)
    add_table(
        document,
        ["成员", "类型", "含义", "单位"],
        [
            ["x", "double", "东向坐标或东向分量", "m"],
            ["y", "double", "北向坐标或北向分量", "m"],
            ["z", "double", "上向坐标或垂向分量", "m"],
        ],
        [1700, 1700, 4100, 1860],
        font_size=8.9,
    )
    add_body(
        document,
        "Vector3 在本模型中只用来表示坐标或位移，因此三个分量的单位统一为 m。例如 initialPosition 是零时刻目标中心坐标，observationPosition 是观测点坐标。",
    )
    add_page_break(document)

    add_heading(document, "6. TargetParameter：目标输入参数", level=1)
    add_body(
        document,
        "TargetParameter 是最重要的输入结构。setTarget 会先校验全部成员，再保存参数并重建等效水动力参数；如果校验失败，原来已经可用的模型配置不会被破坏。",
    )
    add_table(
        document,
        ["成员", "类型 / 默认值", "单位", "通俗含义与约束"],
        [
            ["type", "TargetType / SurfaceShip", "无", "目标类型；参与工况选择和几何校验"],
            ["length", "double / 0", "m", "目标长度 L，必须>0，且不能小于宽度或垂向尺度"],
            ["width", "double / 0", "m", "目标宽度 B，必须>0"],
            ["draft", "double / 0", "m", "水面舰艇吃水；潜艇旧接口兼容高度"],
            ["height", "double / 0", "m", "潜艇垂向高度；为0时回退使用 draft"],
            ["waterDepth", "double / 0", "m", "静水面到海床的海深 H，必须不小于目标垂向尺度"],
            ["velocity", "double / 0", "m/s", "沿目标纵轴的航速 U，必须≥0"],
            ["initialPosition", "Vector3 / {0,0,0}", "m", "零时刻目标几何中心的全局坐标"],
            ["headingDegrees", "double / 0", "°", "水平航向角；东向为 0°，逆时针为正"],
            ["waterDensity", "double / 1025", "kg/m³", "海水密度，必须为有限正数"],
            ["gravityAcceleration", "double / 9.80665", "m/s²", "重力加速度，必须为有限正数"],
            ["blockCoefficient", "double / 0.70", "0～1", "方形系数 Cb，用于估算排水体积，范围必须为 (0,1]"],
            ["minimumDistance", "double / 1", "m", "目标或镜像点允许的最小中心距离，必须>0"],
            ["modelOptions", "ModelOptions / 默认值", "无", "过渡中心0.10、半宽0.02、浅水 H/L 阈值0.30、临界区半宽0.05和兴波衰减阈值0.01"],
        ],
        [2500, 2400, 1100, 3360],
        font_size=8.1,
    )
    add_heading(document, "6.1 参数之间的关系", level=2)
    for text in [
        "等效排水体积按 V=Cb×length×width×垂向尺度估算；水面舰艇使用 draft，潜艇优先使用 height。",
        "水面舰艇中心 z 必须等于 -draft/2；潜艇顶部和底部必须都处在水体内。",
        "航速为 0 时目标可以正常计算，但 dynamicPressure 为 0，总表压等于静水压力。",
        "海深不仅决定允许的 z 范围，还参与 H/L、FrH、海床镜像和浅水工况判断。",
        "minimumDistance 是数值安全和模型适用性限制，不是传感器的物理探测距离。",
    ]:
        add_bullet(document, text, bullet_id)

    add_heading(document, "7. PressureFieldSample：单点计算结果", level=1)
    add_table(
        document,
        ["成员", "类型 / 单位", "通俗含义"],
        [
            ["observationPosition", "Vector3 / m", "传入的固定观测点全局坐标"],
            ["targetPosition", "Vector3 / m", "当前时刻目标几何中心的全局坐标"],
            ["time", "double / s", "当前样本对应的仿真时刻"],
            ["distance", "double / m", "观测点到当前目标中心的三维直线距离"],
            ["longitudinalOffset", "double / m", "从目标中心看，观测点沿舰首方向的相对距离"],
            ["lateralOffset", "double / m", "从目标中心看，观测点横向的相对距离"],
            ["observationDepth", "double / m", "观测点距静水面的正水深，等于 -observationPosition.z"],
            ["regime", "PressureFieldRegime", "本次计算自动选择的水动力工况"],
            ["lengthFroudeNumber", "double / 无量纲", "船长弗劳德数 U/sqrt(gL)"],
            ["depthFroudeNumber", "double / 无量纲", "水深弗劳德数 U/sqrt(gH)"],
            ["waveAttenuation", "double / 无量纲", "潜艇自由液面兴波指数衰减估算"],
            ["transitionWeight", "double / 无量纲", "兴波模型在0.08～0.12平滑过渡带内的权重"],
            ["hydrostaticPressure", "double / Pa", "仅由水深、密度和重力产生的静水表压"],
            ["bodyDynamicPressure", "double / Pa", "目标本体偶极子产生的压力变化"],
            ["freeSurfaceCorrectionPressure", "double / Pa", "刚盖或兴波自由液面修正"],
            ["wavePressure", "double / Pa", "浅水亚临界兴波压力，低速时为0"],
            ["seabedCorrectionPressure", "double / Pa", "海床及重复边界镜像修正"],
            ["dynamicPressure", "double / Pa", "各动态压力分量之和，可正可负"],
            ["totalGaugePressure", "double / Pa", "静水表压与动态压力之和"],
        ],
        [2850, 2200, 4310],
        font_size=8.35,
    )
    add_callout(
        document,
        "压力关系",
        "dynamicPressure = bodyDynamicPressure + freeSurfaceCorrectionPressure + seabedCorrectionPressure；totalGaugePressure = hydrostaticPressure + dynamicPressure。",
    )
    add_page_break(document)

    add_heading(document, "8. GridParameter：规则空间网格", level=1)
    add_table(
        document,
        ["成员", "类型 / 默认值", "含义与约束"],
        [
            ["minimum", "Vector3 / {0,0,0}", "网格最小坐标，每个分量都必须≤maximum 对应分量"],
            ["maximum", "Vector3 / {0,0,0}", "网格最大坐标；整个 z 范围必须位于水体内"],
            ["xCount", "size_t / 1", "x 方向采样点数，必须>0"],
            ["yCount", "size_t / 1", "y 方向采样点数，必须>0"],
            ["zCount", "size_t / 1", "z 方向采样点数，必须>0"],
        ],
        [2400, 2400, 4560],
        font_size=8.7,
    )
    add_body(
        document,
        "网格起点和终点都包含在结果中。某个方向的点数为 1 时，该方向只取 minimum 对应坐标。结果按 z、y、x 三层循环展开，也就是 x 变化最快、z 变化最慢。",
    )
    add_callout(
        document,
        "数量限制",
        "单次网格最多生成 1,000,000 个点。总点数为 xCount×yCount×zCount，超过限制会抛出 std::length_error。",
        accent=CAUTION,
    )

    add_heading(document, "9. 模型内部保存的状态", level=1)
    add_body(
        document,
        "以下成员属于类的内部状态，普通调用方不能直接访问。了解它们可以帮助理解为什么参数必须通过 setTarget 更新。",
    )
    add_table(
        document,
        ["成员", "类型 / 默认值", "作用"],
        [
            ["kMaximumGridPointCount", "size_t / 1,000,000", "单次空间网格最大点数"],
            ["kMaximumTimeSampleCount", "size_t / 1,000,000", "单次时间序列最大样本数"],
            ["m_target", "TargetParameter", "保存当前已校验并生效的完整目标参数"],
            ["m_dipoleStrengthPerSpeed", "double / 0", "缓存单位航速的等效势流偶极矩，单位 m³"],
            ["m_configured", "bool / false", "记录模型是否已成功设置参数"],
        ],
        [3000, 2600, 3760],
        font_size=8.6,
    )
    add_callout(
        document,
        "为什么不能直接修改 target()",
        "target() 返回常量引用，调用方只能读取。参数变化必须调用 setTarget，模型才能重新校验并更新 m_dipoleStrengthPerSpeed。",
    )
    add_page_break(document)

    add_heading(document, "10. 公共函数总览", level=1)
    add_table(
        document,
        ["函数", "主要作用", "返回值"],
        [
            ["TargetPressureFieldModel()", "创建尚未配置参数的空模型", "模型对象"],
            ["TargetPressureFieldModel(param)", "创建并立即配置可计算模型", "模型对象"],
            ["setTarget(param)", "校验并更新目标参数，同时重建水动力缓存", "void"],
            ["target()", "读取当前生效的目标参数", "const TargetParameter&"],
            ["regime()", "读取自动识别的水动力工况", "PressureFieldRegime"],
            ["calculate(position)", "计算零时刻单个水下观测点", "PressureFieldSample"],
            ["calculate(position, time)", "计算指定时刻单个水下观测点", "PressureFieldSample"],
            ["simulate(position, start, duration, rate)", "生成固定观测点的连续压力时序", "vector<PressureFieldSample>"],
            ["simulateGrid(grid, time)", "生成指定时刻的规则空间压力分布", "vector<PressureFieldSample>"],
        ],
        [3900, 3660, 1800],
        font_size=8.3,
    )

    add_heading(document, "11. 构造和配置函数", level=1)
    add_heading(document, "11.1 默认构造函数", level=2)
    add_code_block(
        document,
        "TargetPressureFieldModel(); // 创建空模型，必须先调用 setTarget 才能计算。",
    )
    add_table(
        document,
        ["项目", "说明"],
        [
            ["作用", "创建一个尚未配置目标参数的模型对象"],
            ["输入", "无"],
            ["输出", "TargetPressureFieldModel 对象"],
            ["异常", "构造时不抛出；若未配置就读取或计算，会抛出 std::logic_error"],
        ],
        [1800, 7560],
        font_size=8.8,
    )

    add_heading(document, "11.2 带参数构造函数", level=2)
    add_code_block(
        document,
        "explicit TargetPressureFieldModel(const TargetParameter& param); // 校验参数并创建可计算模型。",
    )
    add_table(
        document,
        ["项目", "说明"],
        [
            ["作用", "用完整参数创建模型，并立即计算单位航速的等效偶极矩"],
            ["输入", "param：完整 TargetParameter"],
            ["输出", "已配置、可立即调用 calculate、simulate 和 simulateGrid 的模型"],
            ["异常", "参数非法时抛出 std::invalid_argument"],
        ],
        [1800, 7560],
        font_size=8.8,
    )

    add_heading(document, "11.3 setTarget", level=2)
    add_code_block(
        document,
        "void setTarget(const TargetParameter& param); // 更新参数并重建水动力缓存。",
    )
    add_table(
        document,
        ["项目", "说明"],
        [
            ["作用", "先校验新参数，成功后保存参数并重建等效偶极子强度"],
            ["输入", "param：新的完整目标参数"],
            ["输出", "无；成功后新参数立即生效"],
            ["异常", "参数非法时抛出 std::invalid_argument，原配置保持不变"],
            ["使用建议", "一组固定参数只需设置一次，不要在每个采样点重复调用"],
        ],
        [1800, 7560],
        font_size=8.8,
    )

    add_heading(document, "11.4 target", level=2)
    add_code_block(
        document,
        "const TargetParameter& target() const; // 只读返回当前生效参数。",
    )
    add_table(
        document,
        ["项目", "说明"],
        [
            ["作用", "让其他模块读取模型当前使用的参数"],
            ["输入", "无"],
            ["输出", "当前 TargetParameter 的常量引用，调用方不能通过该引用修改模型"],
            ["异常", "模型尚未配置时抛出 std::logic_error"],
        ],
        [1800, 7560],
        font_size=8.8,
    )
    add_page_break(document)

    add_heading(document, "12. 单点计算函数", level=1)
    add_heading(document, "12.1 calculate(position)", level=2)
    add_code_block(
        document,
        "PressureFieldSample calculate(const Vector3& observationPosition) const; // 计算零时刻。",
    )
    add_table(
        document,
        ["项目", "说明"],
        [
            ["作用", "计算零时刻某个水下位置的静水压力、动态压力和总表压"],
            ["输入", "observationPosition：观测点全局坐标，单位 m"],
            ["输出", "一个 PressureFieldSample，time 固定为 0"],
            ["内部行为", "直接转调 calculate(observationPosition, 0.0)"],
            ["异常", "未配置、坐标非法、位置不在水体内、进入目标内部或距离过近时抛出相应异常"],
        ],
        [1800, 7560],
        font_size=8.7,
    )

    add_heading(document, "12.2 calculate(position, time)", level=2)
    add_code_block(
        document,
        """PressureFieldSample calculate(
    const Vector3& observationPosition,
    double timeSeconds) const; // 计算指定时刻。""",
    )
    add_table(
        document,
        ["项目", "说明"],
        [
            ["作用", "先计算目标在指定时刻的位置，再计算观测点的完整压力结果"],
            ["输入1", "observationPosition：固定观测点全局坐标，单位 m"],
            ["输入2", "timeSeconds：相对零时刻的时间，单位 s，可以为负"],
            ["目标运动", "目标保持深度不变，沿水平航向以 velocity 匀速直线运动"],
            ["输出", "一个 PressureFieldSample，包含当前目标位置、距离、相对偏移和三类压力"],
            ["异常", "未配置、坐标/时间非有限、观测点越界、位于目标内部或距离过近时抛出相应异常"],
        ],
        [1800, 7560],
        font_size=8.65,
    )
    add_body(
        document,
        "当前目标位置为：x=x0+cos(heading)×velocity×time，y=y0+sin(heading)×velocity×time，z=z0。目标不会自动转弯、升沉或改变航速。",
    )
    add_heading(document, "12.3 单点计算内部步骤", level=2)
    for text in [
        "检查模型是否已经配置，并检查坐标和时间是否为有限数值。",
        "根据初始位置、航向、航速和时间计算当前目标中心位置。",
        "把观测点相对位移旋转到目标坐标系，判断是否进入等效椭球内部。",
        "计算单位航速偶极矩、航速平方和海水密度形成的压力系数。",
        "计算 FrL、FrH 和潜艇兴波衰减量，按目标类型选择水动力工况。",
        "按所选工况分别计算本体、刚盖或兴波自由液面项以及海床修正并求和。",
        "计算静水压力，再得到总表压并填充 PressureFieldSample。",
    ]:
        add_numbered(document, text, decimal_id)

    add_heading(document, "13. 时间序列函数 simulate", level=1)
    add_code_block(
        document,
        """std::vector<PressureFieldSample> simulate(
    const Vector3& observationPosition,
    double startTimeSeconds,
    double durationSeconds,
    double sampleRateHz) const; // 在固定观测点生成时序。""",
    )
    add_table(
        document,
        ["参数", "输入/输出", "含义"],
        [
            ["observationPosition", "输入 / Vector3 / m", "固定压力传感器或观测点坐标"],
            ["startTimeSeconds", "输入 / double / s", "时序开始时刻，可以为负"],
            ["durationSeconds", "输入 / double / s", "持续时间，必须≥0"],
            ["sampleRateHz", "输入 / double / Hz", "每秒采样次数，必须>0"],
            ["返回值", "输出 / vector<PressureFieldSample>", "按时间先后排列的完整压力样本"],
        ],
        [2600, 2800, 3960],
        font_size=8.7,
    )
    add_body(
        document,
        "样本数量为 N=floor(durationSeconds×sampleRateHz)+1。第 k 个样本时刻为 startTimeSeconds+k/sampleRateHz。结果包含起点，并保留持续时间范围内最后一个完整采样点。",
    )
    add_callout(
        document,
        "数量限制",
        "单次时间序列最多生成 1,000,000 个样本。持续时间与采样率乘积过大时会抛出 std::length_error。",
        accent=CAUTION,
    )
    add_callout(
        document,
        "异常会中断整段时序",
        "simulate 内部逐点调用 calculate。如果目标运动过程中某个时刻使观测点进入目标内部或过近范围，函数会在该点抛出异常，不会返回不完整结果。",
        accent=CAUTION,
    )

    add_heading(document, "14. 空间网格函数 simulateGrid", level=1)
    add_code_block(
        document,
        """std::vector<PressureFieldSample> simulateGrid(
    const GridParameter& grid,
    double timeSeconds = 0.0) const; // 生成指定时刻规则网格。""",
    )
    add_table(
        document,
        ["项目", "说明"],
        [
            ["作用", "计算指定时刻二维平面或三维空间中的规则压力分布"],
            ["输入1", "grid：网格边界与三个方向采样点数"],
            ["输入2", "timeSeconds：需要计算的时刻，默认值为 0 s"],
            ["输出", "按 z、y、x 顺序展开的 PressureFieldSample 数组，x 变化最快"],
            ["异常", "网格或时间非法时抛出 invalid_argument；总点数过大时抛出 length_error；某点不适用时抛出 domain_error"],
        ],
        [1800, 7560],
        font_size=8.6,
    )
    add_callout(
        document,
        "二维网格",
        "需要计算固定深度平面时，可令 minimum.z=maximum.z，并设置 zCount=1。例如 z=-30 m 表示海面下 30 m 的水平切片。",
    )
    add_callout(
        document,
        "目标内部点",
        "simulateGrid 不会自动跳过目标内部或过近点。只要网格中存在一个不适用点，整个调用就会抛出异常。调用方应提前避开或分区计算。",
        accent=CAUTION,
    )
    add_page_break(document)

    add_heading(document, "15. 内部关键函数", level=1)
    add_body(
        document,
        "以下函数是模型内部实现细节，普通业务代码不应直接调用。它们负责参数校验、坐标换算、等效水动力参数和压力公式。",
    )
    add_table(
        document,
        ["内部函数", "作用", "何时调用"],
        [
            ["validateTarget", "检查尺寸、海深、位置、航速、航向、环境参数、方形系数和最小距离", "构造函数和 setTarget"],
            ["validateGrid", "检查网格边界、z 范围及三个方向点数", "simulateGrid"],
            ["rebuildHydrodynamicParameters", "根据排水体积和形状因子重建单位航速偶极矩", "setTarget 成功后"],
            ["ensureConfigured", "确认模型已经配置", "所有读取或计算接口"],
            ["isFinite / isFiniteVector", "检查数值和三维坐标是否为有限值", "参数校验和计算入口"],
            ["degreesToRadians", "把航向角从度转换为弧度", "运动和坐标旋转"],
            ["magnitude", "计算三维矢量的欧氏长度", "距离计算"],
            ["globalToBody", "把全局水平位移旋转到目标纵横坐标系", "calculate"],
            ["gridCoordinate", "在最小值与最大值之间生成规则网格坐标", "simulateGrid"],
            ["calculateLongitudinalFactor", "由长、宽、吃水计算细长椭球纵向几何因子", "重建水动力参数"],
            ["calculateImagePressure", "计算一个真实或镜像偶极子对观测点的动态压力贡献", "calculate 的镜像循环"],
        ],
        [3050, 4300, 2010],
        font_size=8.15,
    )
    add_heading(document, "15.1 内部压力关系", level=2)
    for text in [
        "排水体积：V=blockCoefficient×length×width×垂向尺度。",
        "动态压力单项具有 (3cos²θ-1)/r³ 的方向性和距离衰减规律。",
        "动态压力强度与 waterDensity、单位航速偶极矩和 velocity² 成正比。",
        "低速刚盖工况使用两组周期镜像；浅水亚临界工况以线性深度平均兴波项代替刚盖自由液面项。",
        "浅水判据默认采用 H/L≤0.3，亚临界条件要求 FrH 低于临界区下界。",
        "浅水亚临界工况在0.08<FrL<0.12内以 w=s²(3-2s) 平滑混合两套边界项，本体压力不重复混合。",
        "静水压力为 waterDensity×gravityAcceleration×observationDepth。",
    ]:
        add_bullet(document, text, bullet_id)

    add_heading(document, "16. 参数校验和异常处理", level=1)
    add_table(
        document,
        ["异常", "常见原因", "调用方建议"],
        [
            ["std::invalid_argument", "参数非有限、尺寸≤0、长度小于宽度/吃水、海深不足、航速<0、方形系数越界、时间或网格非法", "在界面层逐项校验并提示具体参数"],
            ["std::logic_error", "默认构造后没有调用 setTarget，或在未配置状态读取 target", "检查模型初始化顺序"],
            ["std::domain_error", "观测点在水体外、目标内部，或距离目标/镜像点小于 minimumDistance", "调整观测点、时间或网格范围，不要简单忽略"],
            ["std::length_error", "时间样本或网格点超过 1,000,000", "降低采样率、缩小网格或分批调用"],
        ],
        [2500, 4400, 2460],
        font_size=8.35,
    )
    add_heading(document, "16.1 TargetParameter 校验清单", level=2)
    for text in [
        "length、width 和按类型解释的垂向尺度必须合法，且长度不小于宽度和垂向尺度。",
        "waterDepth 必须为有限正数，并且不小于目标垂向尺度。",
        "initialPosition、velocity 和 headingDegrees 必须为有限值，velocity≥0。",
        "waterDensity 和 gravityAcceleration 必须为有限正数。",
        "blockCoefficient 必须在 (0,1] 范围内。",
        "minimumDistance 必须为有限正数。",
        "水面舰艇中心 z 必须等于 -draft/2；潜艇顶部不得高于静水面且底部不得低于海床。",
        "模型选择阈值必须位于各自规定的无量纲范围内。",
    ]:
        add_bullet(document, text, bullet_id)

    add_heading(document, "17. 完整调用示例", level=1)
    add_code_block(
        document,
        """using Model = TargetPressureFieldModel;

Model::TargetParameter target;
target.type = Model::TargetType::SurfaceShip;  // 设置为水面舰艇。
target.length = 100.0;                         // 舰长 L，单位 m。
target.width = 15.0;                           // 舰宽 B，单位 m。
target.draft = 5.0;                            // 吃水 T，单位 m。
target.waterDepth = 60.0;                      // 海深 H，单位 m。
target.velocity = 3.0;                         // 低速航速 U，单位 m/s。
target.initialPosition = {-180.0, 40.0, -2.5}; // 零时刻目标中心位置。
target.headingDegrees = 0.0;                   // 舰首朝东。
target.waterDensity = 1025.0;                  // 海水密度，单位 kg/m³。
target.gravityAcceleration = 9.80665;          // 重力加速度，单位 m/s²。
target.blockCoefficient = 0.70;                // 方形系数。
target.minimumDistance = 1.0;                  // 最小允许中心距离。

const Model model(target);                     // 创建可立即计算的模型。

const Model::Vector3 sensor{0.0, 0.0, -55.0};  // 固定水压观测点。
const auto instant = model.calculate(sensor, 20.0); // 计算第 20 s 的压力。

const double dynamic = instant.dynamicPressure;     // 目标动态压力，单位 Pa。
const double total = instant.totalGaugePressure;    // 总表压，单位 Pa。
const double distance = instant.distance;           // 目标中心距离，单位 m。

const auto series = model.simulate(
    sensor,
    0.0,                                      // 时序开始时刻。
    75.0,                                     // 持续时间，单位 s。
    20.0);                                    // 采样率，单位 Hz。

Model::GridParameter grid;
grid.minimum = {-500.0, -300.0, -30.0};      // 网格最小坐标。
grid.maximum = {500.0, 300.0, -30.0};        // 网格最大坐标。
grid.xCount = 101;                            // x 方向采样点数。
grid.yCount = 61;                             // y 方向采样点数。
grid.zCount = 1;                              // 固定深度二维平面。

const auto spatial = model.simulateGrid(grid, 20.0); // 生成第 20 s 的空间分布。""",
    )

    add_heading(document, "18. 常见问题", level=1)
    add_table(
        document,
        ["问题", "通俗回答"],
        [
            ["L、B、T 分别是什么？", "length 是长度 L，width 是宽度 B；水面舰艇以 draft 表示吃水 T，潜艇优先以 height 表示垂向高度。"],
            ["type 会改变计算吗？", "会参与工况选择和几何校验，但尺寸、潜深、方形系数仍需调用方填写。"],
            ["怎样判断浅水？", "默认以 H/L≤0.3 判为浅水，并结合 FrL 和 FrH 决定是否进入浅水亚临界模型。"],
            ["为什么需要过渡带？", "避免 FrL 经过0.10时突然更换公式造成压力跳变；默认在0.08～0.12内平滑混合。"],
            ["为什么水下 z 是负数，水深却是正数？", "z 表示向上的坐标；水下位于原点下方，所以为负。observationDepth 是距海面的距离，因此为正。"],
            ["为什么静止目标的动态压力为零？", "动态压力系数与 velocity² 成正比；velocity=0 时只剩静水压力。"],
            ["calculate(position) 计算什么时间？", "它固定计算零时刻，相当于 calculate(position, 0.0)。"],
            ["simulate 包含结束时刻吗？", "只有结束时刻正好落在完整采样点上才包含；样本数按 floor(duration×rate)+1 计算。"],
            ["simulateGrid 如何恢复三维数组？", "输出按 z、y、x 展开，x 变化最快，可按 [zIndex][yIndex][xIndex] 顺序恢复。"],
            ["为什么网格中一个点非法会全部失败？", "函数逐点调用 calculate，当前没有自动跳过或写入空值的机制。"],
            ["可以并行调用 calculate 吗？", "模型配置完成后计算函数为 const，通常适合只读并行；不要同时调用 setTarget 修改配置。"],
        ],
        [3100, 6260],
        font_size=8.4,
    )

    add_heading(document, "19. 调用方检查清单", level=1)
    for text in [
        "确认使用东-北-上坐标，静水面 z=0，水下 z 为负。",
        "确认尺寸、位置、距离使用 m，时间使用 s，航速使用 m/s，压力使用 Pa。",
        "默认构造后先调用 setTarget。",
        "确认目标等效排水体完整位于静水面和海床之间。",
        "单点和时序计算前确认观测点位于水体内，并避开目标内部和过近区域。",
        "时序计算前检查持续时间、采样率和最大样本数量。",
        "空间网格计算前检查边界顺序、z 范围、点数和 z、y、x 展开顺序。",
        "参数变化后必须通过 setTarget 生效，不能只修改原 TargetParameter 变量。",
        "分析目标特征时优先读取 dynamicPressure，并明确是否需要同时展示静水压力基线。",
    ]:
        add_bullet(document, text, bullet_id)

    add_callout(
        document,
        "总结",
        "TargetParameter 负责“描述目标和环境”，GridParameter 负责“描述空间范围”，PressureFieldSample 负责“承载单点结果”；calculate、simulate 和 simulateGrid 分别对应单点、时序和空间三类计算。",
    )

    # 保存最终 Word 文档。
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


def main() -> None:
    """生成并保存接口说明 Word 文档。"""

    output = build_document()
    print(f"已生成 Word 文档：{output}")


if __name__ == "__main__":
    main()
