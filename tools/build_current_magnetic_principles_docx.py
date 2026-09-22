"""生成当前舰艇磁场仿真原理 Word 文档。"""

from __future__ import annotations

import json
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORK_DIR = PROJECT_ROOT / "build" / "docx_magnetic_principles"
EQUATION_DIR = WORK_DIR / "equations"
OUTPUT_DOCX = PROJECT_ROOT / "docs" / "当前舰艇磁场仿真原理.docx"
FLOW_IMAGE = WORK_DIR / "magnetic_flow.png"
ARRAY_IMAGE = WORK_DIR / "dipole_array.png"

# 所有公式均以 LaTeX 源表达式维护，渲染后再插入 Word。
EQUATIONS = {
    "eq01": r"a=\frac{L}{2},\qquad b=\frac{W}{2},\qquad c=\frac{H}{2}",
    "eq02": r"V_{\mathrm e}=\frac{4\pi abc}{3}=\frac{\pi LWH}{6},\qquad V_{\mathrm m}=\eta V_{\mathrm e}",
    "eq03": r"q(\mathbf d)=\left(\frac{d_x^{\mathrm b}}{a}\right)^2+\left(\frac{d_y^{\mathrm b}}{b}\right)^2+\left(\frac{d_z^{\mathrm b}}{c}\right)^2",
    "eq04": r"\mathbf r_{\mathrm c}(t)=\mathbf r_0+\mathbf R\left[vt,0,0\right]^{\mathrm T}",
    "eq05": r"\mathbf m_{\mathrm r}^{\mathrm b}=V_{\mathrm m}\mathbf M_{\mathrm r},\qquad \mathbf m_{\mathrm r}=\mathbf R\mathbf m_{\mathrm r}^{\mathrm b}",
    "eq06": r"N_i=\frac{abc}{2}\int_0^{\infty}\frac{\mathrm ds}{(s+a_i^2)\sqrt{(s+a^2)(s+b^2)(s+c^2)}},\qquad \sum_{i=1}^{3}N_i=1",
    "eq07": r"\chi=\mu_{\mathrm r}-1,\qquad \mathbf H_{\mathrm g}^{\mathrm b}=\frac{\mathbf R^{\mathrm T}\mathbf B_{\mathrm g}}{\mu_0}",
    "eq08": r"m_{\mathrm i,k}^{\mathrm b}=V_{\mathrm m}\frac{\chi H_{\mathrm g,k}^{\mathrm b}}{1+N_k\chi},\qquad \mathbf m_{\mathrm i}=\mathbf R\mathbf m_{\mathrm i}^{\mathrm b}",
    "eq09": r"\mathbf G(\mathbf r)=\frac{\mu_0}{4\pi r^3}\left(3\hat{\mathbf r}\hat{\mathbf r}^{\mathrm T}-\mathbf I_3\right)",
    "eq10": r"\mathbf B(\mathbf r)=\mathbf G(\mathbf r)\mathbf m=\frac{\mu_0}{4\pi r^3}\left[3(\mathbf m\cdot\hat{\mathbf r})\hat{\mathbf r}-\mathbf m\right]",
    "eq11": r"\delta_i^{\mathrm b}=\left[a\xi_i,b\nu_i,c\zeta_i\right]^{\mathrm T},\qquad \xi_i^2+\nu_i^2+\zeta_i^2<1",
    "eq12": r"\widetilde w_i=1-(\xi_i^2+\nu_i^2+\zeta_i^2),\qquad w_i=\frac{\widetilde w_i}{\sum_j\widetilde w_j},\qquad \sum_i w_i=1",
    "eq13": r"\mathbf G_{\mathrm{array}}(\mathbf r)=\sum_{i=1}^{N_{\mathrm a}}w_i\mathbf G\left(\mathbf r-\mathbf R\delta_i^{\mathrm b}\right)",
    "eq14": r"\mathbf G_{\mathrm{hybrid}}=\mathbf G_0+\alpha\left(\mathbf G_{\mathrm{array}}-\mathbf G_0\right),\qquad 0\leq\alpha\leq1",
    "eq15": r"\mathbf B_{\mathrm{static}}=\mathbf G_{\mathrm{hybrid}}\mathbf m_{\mathrm r},\qquad \mathbf B_{\mathrm{induced}}=\mathbf G_{\mathrm{hybrid}}\mathbf m_{\mathrm i}",
    "eq16": r"\mathbf B_{\mathrm{macro}}=\mathbf G_0(\mathbf m_{\mathrm r}+\mathbf m_{\mathrm i}),\qquad \Delta\mathbf B_{\mathrm{local}}=(\mathbf G_{\mathrm{hybrid}}-\mathbf G_0)(\mathbf m_{\mathrm r}+\mathbf m_{\mathrm i})",
    "eq17": r"\mathbf B_{\mathrm{target}}=\mathbf B_{\mathrm{macro}}+\Delta\mathbf B_{\mathrm{local}}=\mathbf B_{\mathrm{static}}+\mathbf B_{\mathrm{induced}}",
    "eq18": r"\mathbf B_{\mathrm{env}}=\mathbf B_{\mathrm{main}}+\mathbf B_{\mathrm{crust}}+\mathbf B_{\mathrm{fluct}}+\mathbf B_{\mathrm{local}}+\mathbf B_{\mathrm{motional}}",
    "eq19": r"\mathbf B_{\mathrm{sensor}}=\mathbf B_{\mathrm{target}}+\mathbf B_{\mathrm{env}}",
    "eq20": r"\Delta B_{\mathrm{scalar}}=\|\mathbf B_{\mathrm{sensor}}\|-\|\mathbf B_{\mathrm{env}}\|",
    "eq21": r"N=\left\lfloor Tf_{\mathrm s}\right\rfloor+1,\qquad t_k=t_{\mathrm{start}}+\frac{k}{f_{\mathrm s}},\qquad k=0,\ldots,N-1",
    "eq22": r"\frac{\|\Delta\mathbf B_{\mathrm{local}}\|}{\|\mathbf B_{\mathrm{macro}}\|}=\mathcal O\left(\frac{D^2}{r^2}\right),\qquad r\gg D",
}

BLACK = "000000"
NAVY = "17365D"
BLUE_FILL = "DCE6F1"
PALE_BLUE = "F3F7FB"
LIGHT_GRAY = "D9D9D9"
MUTED = "595959"
WHITE = "FFFFFF"
BODY_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"


def set_run_font(run, size: float = 11, bold: bool = False, color: str = BLACK,
                 name: str = BODY_FONT, east_asia: str = CJK_FONT) -> None:
    """统一设置中英文字体和字号。"""

    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), east_asia)


def set_cell_fill(cell, color: str) -> None:
    """设置表格单元格底色。"""

    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, top: int = 100, start: int = 120,
                     bottom: int = 100, end: int = 120) -> None:
    """为表格单元格设置统一内边距。"""

    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = LIGHT_GRAY) -> None:
    """设置浅灰色表格外框和内部边框。"""

    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def repeat_header(row) -> None:
    """把表格首行设为跨页重复表头。"""

    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def configure_document(document: Document) -> None:
    """配置 Letter 页面、正文、标题和图注样式。"""

    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.2

    title = document.styles["Title"]
    title.font.name = BODY_FONT
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    title.paragraph_format.space_after = Pt(14)

    for style_name, size, before, after in (
        ("Heading 1", 16, 16, 7),
        ("Heading 2", 13, 12, 5),
        ("Heading 3", 11.5, 9, 4),
    ):
        style = document.styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = document.styles["Caption"]
    caption.font.name = BODY_FONT
    caption.font.size = Pt(9.5)
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False


def configure_footer(document: Document) -> None:
    """添加简洁页脚和页码域。"""

    footer = document.sections[0].footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    visible = OxmlElement("w:t")
    visible.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, visible, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color=MUTED)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    """添加标题并防止标题单独停留在页尾。"""

    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True


def add_body(document: Document, text: str, *, first_line: bool = True) -> None:
    """添加正文段落。"""

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line:
        paragraph.paragraph_format.first_line_indent = Inches(0.28)
    paragraph.paragraph_format.widow_control = True
    run = paragraph.add_run(text)
    set_run_font(run)


def add_bullet(document: Document, text: str) -> None:
    """添加项目符号段落。"""

    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run)


def add_numbered(document: Document, text: str) -> None:
    """添加编号步骤段落。"""

    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run)


def add_table(document: Document, headers: list[str], rows: list[list[str]],
              widths: list[float]) -> None:
    """添加带重复表头和交替底色的技术表格。"""

    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_run_font(run, size=9.5, bold=True, color=WHITE)

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_fill(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            set_run_font(run, size=9.2)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_equation(document: Document, equations: dict[str, str], equation_id: str,
                 number: int, explanation: str, width: float = 5.6) -> None:
    """插入由 LaTeX 渲染的公式图片和编号说明。"""

    path = EQUATION_DIR / f"{equation_id}.png"
    if not path.exists():
        raise FileNotFoundError(f"缺少公式图片：{path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.keep_together = True
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", f"LaTeX: {equations[equation_id]}")
    shape._inline.docPr.set("title", f"公式 {number}")
    caption = document.add_paragraph(style="Caption")
    caption.add_run(f"式 {number}  {explanation}")


def add_figure(document: Document, path: Path, caption_text: str,
               alt_text: str, width: float = 6.4) -> None:
    """插入技术示意图及替代文本。"""

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.keep_together = True
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", caption_text)
    caption = document.add_paragraph(caption_text, style="Caption")
    caption.paragraph_format.keep_with_next = False


def choose_chinese_font_path() -> Path:
    """选择可用中文字体文件，保证示意图文字正确显示。"""

    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError("未找到可用于示意图的中文字体")


def create_flow_figure() -> None:
    """生成目标场和环境场合成流程图。"""

    image = Image.new("RGB", (2200, 900), "white")
    draw = ImageDraw.Draw(image)
    font_path = choose_chinese_font_path()
    font = ImageFont.truetype(str(font_path), 34)
    small_font = ImageFont.truetype(str(font_path), 29)
    outline = "#17365D"

    def box(x: int, y: int, width: int, height: int, text: str,
            fill: str = "#EAF1F7") -> None:
        draw.rounded_rectangle((x, y, x + width, y + height), radius=14,
                               fill=fill, outline=outline, width=4)
        lines = text.split("\n")
        line_height = 43
        start_y = y + (height - line_height * len(lines)) / 2
        for index, line in enumerate(lines):
            bounds = draw.textbbox((0, 0), line, font=font)
            text_width = bounds[2] - bounds[0]
            draw.text((x + (width - text_width) / 2, start_y + index * line_height),
                      line, fill="black", font=font)

    def arrow(x1: int, y1: int, x2: int, y2: int) -> None:
        draw.line((x1, y1, x2, y2), fill=outline, width=5)
        angle = math.atan2(y2 - y1, x2 - x1)
        size = 18
        points = [
            (x2, y2),
            (x2 - size * math.cos(angle - 0.55), y2 - size * math.sin(angle - 0.55)),
            (x2 - size * math.cos(angle + 0.55), y2 - size * math.sin(angle + 0.55)),
        ]
        draw.polygon(points, fill=outline)

    box(60, 120, 300, 150, "目标几何\n磁性参数")
    box(450, 120, 320, 150, "椭球退磁\n宏观磁矩")
    box(860, 80, 360, 150, "中心影响矩阵\n宏观场")
    box(860, 360, 360, 150, "多偶极子阵列\n局部修整")
    box(1340, 220, 300, 150, "目标异常场")
    box(80, 650, 390, 150, "WMM EMAG2\n扰动与局部异常")
    box(600, 650, 340, 150, "海洋环境背景场")
    box(1810, 300, 310, 150, "传感器合成场", "#DDEBF7")
    arrow(360, 195, 450, 195)
    arrow(770, 195, 860, 155)
    arrow(770, 225, 860, 420)
    arrow(1220, 155, 1340, 275)
    arrow(1220, 435, 1340, 330)
    arrow(470, 725, 600, 725)
    arrow(940, 725, 1810, 405)
    arrow(1640, 295, 1810, 350)
    draw.text((1320, 535), "逐分量矢量相加", fill="#595959", font=small_font)
    image.save(FLOW_IMAGE, dpi=(180, 180))


def render_equations() -> None:
    """检查 MATLAB 已生成的 LaTeX 公式图片并写出源公式清单。"""

    EQUATION_DIR.mkdir(parents=True, exist_ok=True)
    missing = [key for key in EQUATIONS if not (EQUATION_DIR / f"{key}.png").exists()]
    if missing:
        raise FileNotFoundError(f"MATLAB 尚未生成公式图片：{missing}")
    with (EQUATION_DIR / "equations.json").open("w", encoding="utf-8") as stream:
        json.dump(
            [{"id": key, "latex": value} for key, value in EQUATIONS.items()],
            stream,
            ensure_ascii=False,
            indent=2,
        )


def create_array_figure() -> None:
    """生成椭球纵横平面上的阵列节点示意图。"""

    extent = 0.78
    longitudinal = 7
    transverse = 3
    x_values = [(-extent + 2 * extent * i / (longitudinal - 1)) for i in range(longitudinal)]
    y_values = [(-extent + 2 * extent * i / (transverse - 1)) for i in range(transverse)]
    nodes = [(x, y) for y in y_values for x in x_values if x * x + y * y < 1.0]
    weights = [1.0 - x * x - y * y for x, y in nodes]

    image = Image.new("RGB", (1800, 940), "white")
    draw = ImageDraw.Draw(image)
    font_path = choose_chinese_font_path()
    title_font = ImageFont.truetype(str(font_path), 42)
    label_font = ImageFont.truetype(str(font_path), 31)
    center_x, center_y = 760, 480
    radius_x, radius_y = 580, 330
    draw.ellipse((center_x - radius_x, center_y - radius_y,
                  center_x + radius_x, center_y + radius_y),
                 outline="#17365D", width=6)
    draw.line((center_x - radius_x, center_y, center_x + radius_x, center_y),
              fill="#BFBFBF", width=3)
    draw.line((center_x, center_y - radius_y, center_x, center_y + radius_y),
              fill="#BFBFBF", width=3)
    maximum_weight = max(weights)
    for (x_value, y_value), weight in zip(nodes, weights):
        x = center_x + int(x_value * radius_x)
        y = center_y - int(y_value * radius_y)
        radius = int(13 + 18 * weight / maximum_weight)
        blue = int(210 - 85 * weight / maximum_weight)
        fill = (95, 150, blue)
        draw.ellipse((x - radius, y - radius, x + radius, y + radius),
                     fill=fill, outline="#17365D", width=3)
    sensor_x, sensor_y = 1580, 575
    draw.polygon([(sensor_x, sensor_y - 22), (sensor_x + 22, sensor_y),
                  (sensor_x, sensor_y + 22), (sensor_x - 22, sensor_y)],
                 fill="#C00000")
    draw.line((center_x + 100, center_y - 90, sensor_x - 30, sensor_y - 10),
              fill="#595959", width=4)
    draw.text((1190, 350), "相对位移", fill="#595959", font=label_font)
    draw.text((sensor_x - 55, sensor_y + 35), "观测点", fill="#C00000", font=label_font)
    title = "椭球内部多偶极子阵列的平面示意"
    title_box = draw.textbbox((0, 0), title, font=title_font)
    draw.text(((1800 - (title_box[2] - title_box[0])) / 2, 30), title,
              fill="black", font=title_font)
    draw.text((570, 850), "舰体纵向归一化坐标", fill="black", font=label_font)
    draw.text((85, 430), "舰体横向", fill="black", font=label_font)
    image.save(ARRAY_IMAGE, dpi=(190, 190))


def add_cover(document: Document) -> None:
    """创建简洁封面。"""

    document.add_paragraph().paragraph_format.space_after = Pt(70)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 显式覆盖模板可能为 Title 样式附加的段落边框，保持封面标题简洁。
    paragraph_properties = title._p.get_or_add_pPr()
    paragraph_borders = OxmlElement("w:pBdr")
    bottom_border = OxmlElement("w:bottom")
    bottom_border.set(qn("w:val"), "nil")
    paragraph_borders.append(bottom_border)
    paragraph_properties.append(paragraph_borders)
    title_run = title.add_run("当前舰艇磁场仿真原理")
    set_run_font(title_run, size=26, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(26)
    subtitle_run = subtitle.add_run("椭球宏观场  多偶极子局部修整  HJC 环境合成")
    set_run_font(subtitle_run, size=15, color=MUTED)

    summary = document.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    summary.paragraph_format.left_indent = Inches(0.7)
    summary.paragraph_format.right_indent = Inches(0.7)
    summary.paragraph_format.line_spacing = 1.35
    summary.paragraph_format.space_after = Pt(32)
    summary_run = summary.add_run(
        "本文档说明当前工程中舰艇目标磁场与海洋环境磁场的完整计算链路。"
        "目标场由三轴椭球宏观磁矩确定整体尺度，再由舰体内部多偶极子阵列修整局部空间分布；"
        "HJC 在同一时间轴上计算地磁主场、地壳异常和扰动背景，最后与目标异常场逐分量相加。"
    )
    set_run_font(summary_run, size=11.5)

    metadata = document.add_table(rows=4, cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
    metadata.autofit = False
    set_table_borders(metadata)
    entries = [
        ("文档类型", "仿真原理与实现说明"),
        ("对应实现", "TargetMagneticFieldModel 与 HJC 0.5.0"),
        ("默认阵列", "7 × 3 × 3 候选节点  局部修整系数 0.65"),
        ("生成日期", "2026 年 9 月 21 日"),
    ]
    for row, (label, value) in zip(metadata.rows, entries):
        row.cells[0].width = Inches(1.45)
        row.cells[1].width = Inches(4.75)
        set_cell_fill(row.cells[0], BLUE_FILL)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=130, bottom=130)
        label_run = row.cells[0].paragraphs[0].add_run(label)
        set_run_font(label_run, size=10, bold=True)
        value_run = row.cells[1].paragraphs[0].add_run(value)
        set_run_font(value_run, size=10)
    document.add_page_break()


def add_contents(document: Document) -> None:
    """添加静态目录和阅读说明。"""

    add_heading(document, "文档内容", 1)
    contents = [
        "1  模型范围与场量定义",
        "2  坐标系与目标运动",
        "3  三轴椭球与宏观磁矩",
        "4  磁偶极影响矩阵",
        "5  多偶极子阵列局部修整",
        "6  HJC 海洋环境磁场合成",
        "7  采样 输出与数值流程",
        "8  默认参数与校验规则",
        "9  模型适用范围与标定建议",
        "10  验证结果与实现位置",
    ]
    for item in contents:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.space_after = Pt(5)
        run = paragraph.add_run(item)
        set_run_font(run, size=11)
    add_body(
        document,
        "公式均由 LaTeX 源表达式渲染，并嵌入到 Word 文档中。公式图片的替代文本保留对应 LaTeX，"
        "便于后续核对或重新排版。文档中的磁场矢量采用东 北 上坐标，目标坐标系采用纵 横 垂三个主轴。",
    )
    document.add_page_break()


def build_document() -> Path:
    """生成完整 DOCX 并返回输出路径。"""

    WORK_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    render_equations()
    equations = EQUATIONS
    create_flow_figure()
    create_array_figure()

    document = Document()
    configure_document(document)
    configure_footer(document)
    document.core_properties.title = "当前舰艇磁场仿真原理"
    document.core_properties.subject = "椭球宏观场 多偶极子局部修整 HJC 环境合成"
    document.core_properties.author = "seamine 项目组"
    document.core_properties.keywords = "舰艇磁场, 椭球退磁, 多偶极子阵列, HJC, LaTeX"

    add_cover(document)
    add_contents(document)

    add_heading(document, "1 模型范围与场量定义", 1)
    add_body(
        document,
        "当前磁场仿真把传感器读数拆成目标异常场和海洋环境背景场。目标异常场描述舰艇剩磁及地磁感应磁化产生的附加磁场，"
        "不重复包含地磁背景。HJC 在传感器位置计算环境背景，并与目标异常场按东 北 上三个方向逐分量相加。",
    )
    add_figure(
        document,
        FLOW_IMAGE,
        "图 1  当前磁场仿真计算链路",
        "目标几何和磁性参数先形成椭球宏观磁矩，再由多偶极子阵列修整目标场；WMM、EMAG2和扰动形成环境场，两者在传感器位置逐分量相加。",
    )
    add_table(
        document,
        ["场量", "含义", "是否包含环境背景", "输出单位"],
        [
            ["目标宏观场", "椭球总磁矩位于目标中心时产生的场", "否", "nT"],
            ["局部修整量", "阵列矩阵场相对中心宏观场的差值", "否", "nT"],
            ["目标异常场", "宏观场与局部修整量的矢量和", "否", "nT"],
            ["环境背景场", "主场 地壳异常 扰动 局部异常及规定运动场", "是", "nT"],
            ["传感器合成场", "目标异常场与环境背景场的矢量和", "是", "nT"],
        ],
        [1.2, 3.05, 1.35, 0.9],
    )
    add_body(
        document,
        "该拆分保证 signalOnly environmentOnly 和 totalField 可以分别查看。总场必须由三分量相加后再求模，"
        "不能把目标场模长和背景场模长直接相加。",
    )

    add_heading(document, "2 坐标系与目标运动", 1)
    add_heading(document, "2 1 坐标约定", 2)
    add_body(
        document,
        "全局坐标采用 ENU 约定：x 轴向东，y 轴向北，z 轴向上，水下位置的 z 坐标为负。"
        "目标坐标系的 x 轴沿舰首方向，y 轴沿横向，z 轴沿垂向。航向 俯仰 横滚共同构成从目标坐标系到全局坐标系的旋转矩阵 R。",
    )
    add_heading(document, "2 2 匀速直线运动", 2)
    add_equation(document, equations, "eq04", 1, "目标中心随时间更新")
    add_body(
        document,
        "式中 r0 为零时刻目标中心，v 为沿舰体纵轴的航速。默认演示的俯仰和横滚均为零。"
        "HJC 明确采用正俯仰使舰首向上的约定；独立目标模型现有旋转实现对非零俯仰的符号不同，"
        "因此非零俯仰工况应先统一旋转约定再比较两条链路。",
    )

    add_heading(document, "3 三轴椭球与宏观磁矩", 1)
    add_heading(document, "3 1 几何与适用域", 2)
    add_equation(document, equations, "eq01", 2, "舰艇尺寸与椭球半轴")
    add_equation(document, equations, "eq02", 3, "椭球体积与等效磁性体积")
    add_body(
        document,
        "L W H 分别为目标长度 宽度和高度，η 为磁性材料等效体积占比。模型只在等效椭球外部计算目标磁场。"
        "观测点相对目标中心的位移先旋转到目标坐标系，再计算椭球归一化坐标。",
    )
    add_equation(document, equations, "eq03", 4, "观测点的椭球归一化坐标")
    add_body(
        document,
        "当 q 小于或等于 1 时，观测点位于椭球内部或表面，计算会被拒绝。模型同时检查观测点到目标中心的距离是否小于 minimumDistance。",
    )

    add_heading(document, "3 2 剩磁宏观磁矩", 2)
    add_equation(document, equations, "eq05", 5, "剩磁强度转换为全局坐标磁矩")
    add_body(
        document,
        "Mr 是目标坐标系中的剩磁强度，单位为 A/m。磁性体积与剩磁强度相乘得到剩磁偶极矩，单位为 A·m²，随后由姿态矩阵旋转到全局坐标。",
    )

    add_heading(document, "3 3 地磁感应宏观磁矩", 2)
    add_body(
        document,
        "三轴椭球沿不同主轴的退磁效应不同。程序通过数值积分计算三个退磁因子，并归一化使其和为 1。",
    )
    add_equation(document, equations, "eq06", 6, "三轴椭球退磁因子")
    add_equation(document, equations, "eq07", 7, "磁化率与目标坐标系地磁场强度")
    add_equation(document, equations, "eq08", 8, "第 k 个主轴方向的感应磁矩")
    add_body(
        document,
        "μr 是等效相对磁导率，χ 是磁化率，Bg 是目标位置的地磁感应强度。独立模型使用参数中给定的固定地磁场；"
        "HJC 使用目标当前位置的 WMM 或 WMM 与 EMAG2 组合场，因此两条链路在同一目标参数下仍可能得到不同的感应磁矩。",
    )

    add_heading(document, "4 磁偶极影响矩阵", 1)
    add_body(
        document,
        "单个磁偶极子的三分量磁场可以写成一个 3×3 影响矩阵与磁矩矢量的乘积。矩阵形式便于对多个偶极子逐节点累加，也让宏观场和局部修整共用同一线性映射。",
    )
    add_equation(document, equations, "eq09", 9, "磁偶极三阶影响矩阵")
    add_equation(document, equations, "eq10", 10, "影响矩阵与经典磁偶极公式等价")
    add_body(
        document,
        "r 是源点到观测点的位移，r̂ 是单位方向，I3 是三阶单位矩阵，μ0 为真空磁导率。"
        "独立模型内部将特斯拉乘以 10^9 输出为 nT；HJC 内部统一使用 SI 单位 T，并在门面输出时转换为 nT。",
    )

    add_heading(document, "5 多偶极子阵列局部修整", 1)
    add_heading(document, "5 1 节点生成与权重", 2)
    add_body(
        document,
        "候选节点在目标坐标系的纵 横 垂三个方向均匀取样。当前默认候选阵列为 7×3×3，"
        "归一化坐标范围为负 0.78 到正 0.78。只保留位于单位椭球内部的节点，默认配置得到 27 个有效节点。",
    )
    add_equation(document, equations, "eq11", 11, "阵列节点的目标坐标偏移与保留条件")
    add_equation(document, equations, "eq12", 12, "节点平滑权重和归一化")
    add_figure(
        document,
        ARRAY_IMAGE,
        "图 2  椭球内部阵列节点与平滑权重示意",
        "椭球纵横截面内的多偶极子节点，节点越靠近中心权重越大，观测点位于椭球外部。",
        width=5.9,
    )
    add_body(
        document,
        "节点位置关于目标中心成对对称，归一化权重之和为 1。该构造保持阵列总磁矩等于椭球宏观磁矩，并消除一阶偏心项。"
        "如果极低分辨率偶数网格没有节点落入椭球，程序会退化为中心单节点。",
    )

    add_heading(document, "5 2 阵列影响矩阵", 2)
    add_equation(document, equations, "eq13", 13, "多偶极子阵列的加权影响矩阵")
    add_body(
        document,
        "每个节点使用同一个宏观磁矩方向，但只分配权重 wi 对应的磁矩份额。节点的目标坐标偏移先由 R 旋转到全局坐标，"
        "再计算该节点到观测点的影响矩阵。",
    )

    add_heading(document, "5 3 宏观场与局部修整的混合", 2)
    add_equation(document, equations, "eq14", 14, "宏观影响矩阵与阵列影响矩阵的混合")
    add_body(
        document,
        "G0 是总磁矩位于目标中心时的宏观影响矩阵，α 是局部修整系数。α 等于 0 时退回中心宏观场；"
        "α 等于 1 时完整采用归一化阵列场；当前默认值为 0.65。",
    )
    add_equation(document, equations, "eq15", 15, "剩磁场与感应场分别通过同一混合矩阵计算")
    add_equation(document, equations, "eq16", 16, "目标宏观场和局部修整量")
    add_equation(document, equations, "eq17", 17, "目标异常场的两种等价分解")
    add_body(
        document,
        "第一种分解用于解释物理来源：目标异常场等于剩磁场加感应场。第二种分解用于诊断阵列作用："
        "目标异常场等于椭球宏观场加局部修整量。程序逐分量验证这两个恒等关系。",
    )

    add_heading(document, "5 4 远场一致性", 2)
    add_equation(document, equations, "eq22", 18, "中心对称归一化阵列的远场相对修整量")
    add_body(
        document,
        "D 表示目标特征尺寸，r 表示目标中心到传感器的距离。由于节点权重和为 1 且加权位置一阶矩为零，"
        "阵列与中心宏观场的差异主要来自更高阶空间项。距离远大于目标尺寸时，局部修整量相对宏观场快速减小。",
    )

    add_heading(document, "6 HJC 海洋环境磁场合成", 1)
    add_heading(document, "6 1 环境背景组成", 2)
    add_equation(document, equations, "eq18", 19, "传感器位置的环境背景磁场")
    add_table(
        document,
        ["环境分量", "当前实现", "数据位置"],
        [
            ["主地磁场", "WMM2025 或 WMMHR2025", "目标位置与传感器位置分别求值"],
            ["地壳磁异常", "可选 EMAG2 数值格网", "沿主场方向投影为 ENU 三分量"],
            ["磁扰动", "规定时序或由 PSD 和随机种子生成", "只计入环境背景"],
            ["局部磁异常", "局部矢量格网或局部偶极子目录", "禁止来源不清的重复叠加"],
            ["运动磁场", "可选规定时序", "当前不自动推导 u×B"],
        ],
        [1.25, 2.9, 2.35],
    )
    add_body(
        document,
        "HJC 用目标位置的主场和可选地壳异常计算感应磁矩，用传感器位置的环境场形成背景。"
        "这样目标磁化与传感器背景分别对应正确的空间位置。环境资源缺失或覆盖不足时，计算会返回明确错误，不回退为零背景。",
    )

    add_heading(document, "6 2 传感器总场与标量异常", 2)
    add_equation(document, equations, "eq19", 20, "传感器位置的目标场与环境场合成")
    add_equation(document, equations, "eq20", 21, "磁力仪有符号标量异常")
    add_body(
        document,
        "标量异常是合成矢量模长与环境矢量模长之差，可能为正也可能为负。它不等于目标异常矢量的模长，"
        "也不能用两个模长直接相加得到。",
    )

    add_heading(document, "7 采样 输出与数值流程", 1)
    add_heading(document, "7 1 时间采样", 2)
    add_equation(document, equations, "eq21", 22, "固定传感器时序的样本数与采样时刻")
    add_body(
        document,
        "时序包含起点，并保留持续时间内的最后一个完整采样点。每个采样时刻先更新目标中心，再校验观测点是否进入目标椭球，随后计算目标场 环境场和合成场。",
    )

    add_heading(document, "7 2 计算步骤", 2)
    for step in (
        "校验目标尺寸 姿态 磁性参数 阵列规模和修整系数。",
        "计算椭球体积 磁性体积 退磁因子 剩磁矩和感应磁矩。",
        "建立目标内部阵列节点及归一化权重。",
        "对每个空间点或时刻建立中心矩阵和阵列矩阵，得到目标宏观场 局部修整量和目标异常场。",
        "HJC 计算目标位置和传感器位置的地磁环境，并生成完整背景分量。",
        "逐分量相加得到传感器总场，校验宏观场加修整量及目标场加环境场两个恒等关系。",
        "输出 CSV 和质量元数据，MATLAB 根据三分量计算模值并绘图。",
    ):
        add_numbered(document, step)

    add_heading(document, "7 3 主要输出字段", 2)
    add_table(
        document,
        ["字段组", "定义", "用途"],
        [
            ["static_b*_nt", "阵列修整后的剩磁场三分量", "剩磁来源分析"],
            ["induced_b*_nt", "阵列修整后的感应场三分量", "地磁磁化分析"],
            ["macro_b*_nt", "椭球宏观总磁矩场", "远场基线和模型诊断"],
            ["local_correction_b*_nt", "阵列场减宏观场后的修整量", "近场空间结构分析"],
            ["signal_b*_nt", "HJC 中的目标异常场", "目标信号独立查看"],
            ["environment_b*_nt", "HJC 环境背景场", "背景独立查看"],
            ["total_b*_nt", "目标场与环境场合成", "传感器三轴读数"],
            ["scalar_anomaly_nt", "总场模长减环境场模长", "标量磁力仪异常"],
        ],
        [1.75, 2.9, 1.85],
    )

    add_heading(document, "8 默认参数与校验规则", 1)
    add_table(
        document,
        ["参数", "当前默认值", "校验或物理意义"],
        [
            ["候选阵列", "7 × 3 × 3", "三个方向节点数均大于 0"],
            ["有效节点", "27", "节点需满足单位椭球内部条件"],
            ["阵列范围", "各主轴的 ±0.78", "节点不放在椭球表面"],
            ["局部修整系数", "0.65", "有限值且位于 0 到 1"],
            ["最大候选节点数", "4096", "限制单个采样点的计算量"],
            ["最小中心距离", "演示为 2 m", "同时执行椭球外部判定"],
            ["相对磁导率", "演示为 180", "必须大于或等于 1"],
            ["磁性材料比例", "演示为 0.035", "必须位于 0 到 1"],
        ],
        [1.65, 1.55, 3.3],
    )
    add_body(
        document,
        "目标尺寸 航速 姿态 地磁场和剩磁三分量必须为有限数值。HJC 还校验环境快照有效期 WMM 系数文件 EMAG2 覆盖范围和规定扰动时序长度。"
        "任何致命校验失败都会停止求解，不生成看似有效的零曲线。",
    )

    add_heading(document, "9 模型适用范围与标定建议", 1)
    add_heading(document, "9 1 当前模型能够描述的内容", 2)
    for item in (
        "三轴椭球形状对退磁因子和感应磁矩方向的影响。",
        "剩磁与地磁感应磁化对目标异常场的共同贡献。",
        "目标姿态 运动和传感器相对位置造成的三分量时空变化。",
        "多偶极子阵列相对中心宏观偶极子的近场空间修整。",
        "WMM EMAG2 扰动 局部异常与目标场在传感器位置的矢量合成。",
    ):
        add_bullet(document, item)

    add_heading(document, "9 2 当前模型没有包含的内容", 2)
    for item in (
        "海水和海床多层介质中的频域磁扩散及界面边界条件。",
        "舰艇实际舱段 消磁绕组 局部材料和结构应力形成的独立磁矩分布。",
        "磁滞 饱和 涡流和随航向历史变化的非线性磁化。",
        "传感器噪声 频响 标定误差 采样量化和后端检测算法。",
        "有限元边界解或实测磁场数据的自动反演。",
    ):
        add_bullet(document, item)

    add_heading(document, "9 3 工程标定建议", 2)
    add_body(
        document,
        "当前阵列节点位置和权重是保证对称性与远场一致性的平滑先验，不是实际舰艇磁矩分布的测量结果。"
        "工程应用可保持宏观总磁矩约束不变，用有限元结果或多测线实测数据反演各节点的磁矩向量。"
        "建议同时约束节点磁矩之和 一阶空间矩以及不同航向下的拟合残差，避免只改善单条测线而破坏远场一致性。",
    )

    add_heading(document, "10 验证结果与实现位置", 1)
    add_heading(document, "10 1 当前验证结果", 2)
    add_table(
        document,
        ["验证项", "结果", "说明"],
        [
            ["C++ 构建", "通过", "Release 配置成功编译"],
            ["CTest", "8 项全部通过", "包含独立模型 HJC 和环境集成测试"],
            ["MATLAB", "通过", "空间图和综合时序图实际执行"],
            ["有效阵列节点", "27", "默认 7 × 3 × 3 候选阵列"],
            ["空间目标场峰值", "678.80 nT", "默认水平观测面的当前演示结果"],
            ["HJC 宏观场峰值", "35.84 nT", "当前运动目标演示"],
            ["HJC 局部修整峰值", "7.03 nT", "当前运动目标演示"],
            ["HJC 目标异常范围", "0.175 至 29.309 nT", "当前 WMM 环境演示"],
        ],
        [1.65, 1.55, 3.3],
    )
    add_body(
        document,
        "这些数值用于确认当前实现和输出格式，不代表任意舰艇的实测磁特征。尺寸 磁导率 剩磁 地理位置和阵列标定变化后，结果会相应改变。",
    )

    add_heading(document, "10 2 代码与绘图位置", 2)
    add_table(
        document,
        ["文件", "职责"],
        [
            ["TargetMagneticFieldModel.cpp", "独立目标模型的椭球磁矩 阵列矩阵和空间时序计算"],
            ["HJC src wmm.cpp", "HJC 椭球磁矩 单偶极影响矩阵和阵列修整数值核"],
            ["HJC src module.cpp", "目标场 环境背景与传感器总场的同时间轴合成"],
            ["integration hjc MagneticFieldSimulation.cpp", "主工程到 HJC 的参数适配和结果一致性校验"],
            ["main.cpp", "CSV 输出 演示参数和运行摘要"],
            ["plotMagneticFieldResults.m", "宏观场 局部修整 环境场和总场的 MATLAB 可视化"],
        ],
        [2.65, 3.85],
    )

    OUTPUT_DOCX.unlink(missing_ok=True)
    document.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    output = build_document()
    print(f"已生成：{output}")
