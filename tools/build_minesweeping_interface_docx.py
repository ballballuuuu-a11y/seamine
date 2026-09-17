"""生成《扫雷设备三分量磁场仿真接口说明》Word 文档。"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# 文档输出路径与文档技能脚本路径。
PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DOCX = PROJECT_ROOT / "docs" / "扫雷设备三分量磁场仿真接口说明.docx"
DOCUMENT_SKILL_ROOT = Path(
    "C:/Users/Administrator/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.727.11326/skills/documents"
)
sys.path.insert(0, str(DOCUMENT_SKILL_ROOT / "scripts"))
from table_geometry import apply_table_geometry  # noqa: E402


# compact_reference_guide 预设及 memo_masthead 页首样式参数。
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 90, "bottom": 90, "start": 120, "end": 120}
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "607786"
LIGHT_FILL = "F4F6F9"
TABLE_FILL = "E8EEF5"
CODE_FILL = "F6F8FA"
CAUTION = "7A5A00"
WHITE = "FFFFFF"
BODY_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"


def rgb(value: str) -> RGBColor:
    """把十六进制颜色转换为 Word 使用的颜色对象。"""

    return RGBColor.from_string(value)


def set_run_font(
    run,
    *,
    name: str = BODY_FONT,
    east_asia: str = CJK_FONT,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    """统一设置中西文字体、字号、颜色和强调属性。"""

    run.font.name = name
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_fill(cell, fill: str) -> None:
    """设置表格单元格底色。"""

    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_paragraph_fill(paragraph, fill: str) -> None:
    """设置段落底色。"""

    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_paragraph_border(
    paragraph,
    *,
    side: str,
    color: str,
    size: int = 8,
    space: int = 4,
) -> None:
    """给段落增加指定方向的边框。"""

    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    border = borders.find(qn(f"w:{side}"))
    if border is None:
        border = OxmlElement(f"w:{side}")
        borders.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    """把表格首行标记为跨页重复表头。"""

    properties = row._tr.get_or_add_trPr()
    header = properties.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        properties.append(header)


def configure_styles(document: Document) -> None:
    """按紧凑型技术参考指南预设配置 Word 样式。"""

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(NAVY)
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.22

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 9),
        "Heading 2": (13, BLUE, 13, 6),
        "Heading 3": (11.5, DARK_BLUE, 9, 4),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    callout.font.name = BODY_FONT
    callout.font.size = Pt(10)
    callout.font.color.rgb = rgb(NAVY)
    callout._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    callout.paragraph_format.left_indent = Inches(0.16)
    callout.paragraph_format.right_indent = Inches(0.08)
    callout.paragraph_format.space_before = Pt(5)
    callout.paragraph_format.space_after = Pt(8)
    callout.paragraph_format.line_spacing = 1.18
    callout.paragraph_format.keep_together = True

    code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = CODE_FONT
    code.font.size = Pt(8.2)
    code.font.color.rgb = rgb(NAVY)
    code._element.rPr.rFonts.set(qn("w:ascii"), CODE_FONT)
    code._element.rPr.rFonts.set(qn("w:hAnsi"), CODE_FONT)
    code._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    code.paragraph_format.left_indent = Inches(0.14)
    code.paragraph_format.right_indent = Inches(0.08)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing = 1.05
    code.paragraph_format.keep_together = True


def configure_page(document: Document) -> None:
    """设置 Letter 纵向页面、页边距和页眉页脚距离。"""

    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.48)
    section.footer_distance = Inches(0.48)
    section.different_first_page_header_footer = True


def add_page_field(paragraph) -> None:
    """在页脚中添加 Word 页码域。"""

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=8.5, color=MUTED)


def configure_header_footer(document: Document) -> None:
    """配置运行页眉和页码页脚。"""

    section = document.sections[0]
    section.first_page_header.paragraphs[0].text = ""

    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(3)
    header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = header.add_run("扫雷设备三分量磁场仿真")
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    header.add_run("\t")
    right = header.add_run("接口说明 V1.0")
    set_run_font(right, size=8.5, color=MUTED)
    set_paragraph_border(header, side="bottom", color="D8E2E8", size=5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(footer)


def add_numbering_definitions(document: Document) -> tuple[int, int]:
    """创建真实项目符号和十进制编号定义。"""

    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(item.get(qn("w:abstractNumId")))
        for item in numbering.findall(qn("w:abstractNum"))
    ]
    number_ids = [
        int(item.get(qn("w:numId")))
        for item in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    number_id = max(number_ids, default=0) + 1

    def create_definition(
        definition_id: int,
        instance_id: int,
        number_format: str,
        label_text: str,
        font_name: str,
    ) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(definition_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), number_format)
        level.append(fmt)
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), label_text)
        level.append(text)
        paragraph_properties = OxmlElement("w:pPr")
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), "540")
        indent.set(qn("w:hanging"), "270")
        paragraph_properties.append(indent)
        level.append(paragraph_properties)
        run_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), font_name)
        fonts.set(qn("w:hAnsi"), font_name)
        run_properties.append(fonts)
        level.append(run_properties)
        abstract.append(level)
        numbering.append(abstract)

        number = OxmlElement("w:num")
        number.set(qn("w:numId"), str(instance_id))
        abstract_reference = OxmlElement("w:abstractNumId")
        abstract_reference.set(qn("w:val"), str(definition_id))
        number.append(abstract_reference)
        numbering.append(number)

    create_definition(abstract_id, number_id, "bullet", "•", "Symbol")
    create_definition(abstract_id + 1, number_id + 1, "decimal", "%1.", BODY_FONT)
    return number_id, number_id + 1


def apply_numbering(paragraph, number_id: int) -> None:
    """把段落绑定到指定编号定义。"""

    properties = paragraph._p.get_or_add_pPr()
    number_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(number_id))
    number_properties.append(level)
    number_properties.append(number)
    properties.append(number_properties)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    """添加真实 Word 标题。"""

    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True


def add_body(document: Document, text: str, *, bold_lead: str | None = None) -> None:
    """添加正文段落，可选择加粗开头标签。"""

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, size=10.5, color=NAVY, bold=True)
        body = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(body, size=10.5, color=NAVY)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5, color=NAVY)


def add_callout(
    document: Document,
    label: str,
    text: str,
    *,
    accent: str = BLUE,
) -> None:
    """添加说明或警告提示框。"""

    paragraph = document.add_paragraph(style="Callout")
    set_paragraph_fill(paragraph, LIGHT_FILL)
    set_paragraph_border(paragraph, side="left", color=accent, size=16, space=8)
    label_run = paragraph.add_run(f"{label}：")
    set_run_font(label_run, size=10, color=accent, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=10, color=NAVY)


def add_bullet(document: Document, text: str, number_id: int) -> None:
    """添加真实项目符号段落。"""

    paragraph = document.add_paragraph()
    apply_numbering(paragraph, number_id)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.2
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=NAVY)


def add_numbered(document: Document, text: str, number_id: int) -> None:
    """添加真实编号段落。"""

    paragraph = document.add_paragraph()
    apply_numbering(paragraph, number_id)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.2
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=NAVY)


def add_code_block(document: Document, code_text: str) -> None:
    """添加保留换行和缩进的代码块。"""

    paragraph = document.add_paragraph(style="Code Block")
    set_paragraph_fill(paragraph, CODE_FILL)
    set_paragraph_border(paragraph, side="left", color=BLUE, size=8, space=6)
    run = paragraph.add_run(code_text.rstrip())
    set_run_font(
        run,
        name=CODE_FONT,
        east_asia=CJK_FONT,
        size=8.2,
        color=NAVY,
    )


def add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    widths_dxa: list[int],
    *,
    font_size: float = 8.8,
) -> None:
    """添加具有精确列宽、重复表头和统一边距的表格。"""

    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError("表格列宽之和必须为 9360 DXA")
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Table Grid"
    header_row = table.rows[0]
    set_repeat_table_header(header_row)

    for index, header_text in enumerate(headers):
        cell = header_row.cells[index]
        set_cell_fill(cell, TABLE_FILL)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(header_text)
        set_run_font(run, size=font_size, color=NAVY, bold=True)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 1:
                set_cell_fill(cell, "F9FBFC")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.12
            if column_index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(value))
            set_run_font(run, size=font_size, color=NAVY)

    apply_table_geometry(
        table,
        widths_dxa,
        table_width_dxa=CONTENT_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa=CELL_MARGINS_DXA,
    )
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_page_break(document: Document) -> None:
    """添加明确分页。"""

    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def build_cover(document: Document) -> None:
    """创建技术规范使用的 memo_masthead 首页。"""

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("接口设计说明")
    set_run_font(run, size=10.5, color=CAUTION, bold=True)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("扫雷设备三分量磁场仿真")
    set_run_font(title_run, size=27, color=NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle_run = subtitle.add_run("重要数据结构、函数接口与集成约定")
    set_run_font(subtitle_run, size=14, color=DARK_BLUE)

    metadata = [
        ("文档版本", "V1.0"),
        ("接口状态", "设计建议稿，需完成代码实现后方可直接调用"),
        ("适用对象", "扫雷任务规划、仿真控制、可视化、数据输出和算法验证模块"),
        ("坐标约定", "ENU：x 向东、y 向北、z 向上"),
        ("更新时间", "2026-07-30"),
    ]
    for label, value in metadata:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        label_run = paragraph.add_run(f"{label}：")
        set_run_font(label_run, size=10.5, color=NAVY, bold=True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=10.5, color=NAVY)

    rule = document.add_paragraph()
    rule.paragraph_format.space_before = Pt(9)
    rule.paragraph_format.space_after = Pt(18)
    set_paragraph_border(rule, side="bottom", color=BLUE, size=12)

    add_callout(
        document,
        "设计结论",
        "建议新增独立的 MinesweepingMagneticFieldModel，不修改现有舰艇被动磁场模型的物理含义；公共调用方式继续保持“配置—单点计算—时序仿真—网格仿真”的一致风格。",
    )
    add_page_break(document)


def build_document() -> Path:
    """构建并保存完整接口说明文档。"""

    document = Document()
    document.core_properties.title = "扫雷设备三分量磁场仿真接口说明"
    document.core_properties.subject = "主动磁场源三分量仿真的数据结构与 C++ 调用接口"
    document.core_properties.author = "seamine 项目组"
    document.core_properties.last_modified_by = "seamine 项目组"
    document.core_properties.keywords = "扫雷设备, 三分量磁场, C++接口, 数据结构, 时序仿真"
    document.core_properties.comments = "面向其他功能模块调用的接口设计建议稿。"
    configure_styles(document)
    configure_page(document)
    configure_header_footer(document)
    bullet_id, decimal_id = add_numbering_definitions(document)
    build_cover(document)

    add_heading(document, "1. 文档目的与边界", level=1)
    add_body(
        document,
        "本文定义扫雷设备主动磁场三分量仿真的建议数据结构和公共函数接口，供任务规划、仿真调度、可视化、数据记录及算法验证模块调用。接口重点描述调用契约，不展开内部数值积分和求解器实现。"
    )
    add_heading(document, "1.1 与现有磁场模型的关系", level=2)
    add_table(
        document,
        ["项目", "现有舰艇磁场模型", "新增扫雷设备模型"],
        [
            ["磁源性质", "剩磁与地磁感应形成的被动磁源", "线圈、电缆及磁源阵列形成的主动磁源"],
            ["磁矩变化", "配置后通常保持不变", "随电流波形、时间和姿态动态变化"],
            ["主要输入", "尺寸、磁导率、剩磁、地磁", "磁源几何、电流波形、运动姿态和求解模式"],
            ["主要输出", "静磁场、感应磁场、综合异常场", "主动扫场、平台场、异常场、绝对场和敏感轴投影"],
            ["公共模式", "单点、时间序列、空间网格", "继续保持相同的三类调用方式"],
        ],
        [1700, 3600, 4060],
        font_size=8.8,
    )
    add_callout(
        document,
        "兼容原则",
        "可以复用三维矢量、规则网格、采样时序和结果容器的设计习惯，但不建议复用“静态磁矩/感应磁矩”字段表达主动线圈磁场。",
    )

    add_heading(document, "2. 坐标、单位与结果口径", level=1)
    add_table(
        document,
        ["对象", "约定", "说明"],
        [
            ["全局坐标", "ENU", "x 向东、y 向北、z 向上；水下位置 z 为负"],
            ["设备局部坐标", "x 前、y 左、z 上", "磁源安装位置和法向量均在设备局部坐标系给出"],
            ["位置与距离", "m", "所有几何量统一使用米"],
            ["时间与频率", "s、Hz", "相位统一使用度，角度输入均明确后缀 Degrees"],
            ["电流", "A", "线圈和电缆激励电流"],
            ["磁矩", "A·m²", "线圈远场等效磁矩"],
            ["磁场", "nT", "所有公共结果统一输出纳特"],
            ["异常场", "主动扫场 + 平台场", "不包含背景地磁"],
            ["绝对场", "背景地磁 + 异常场", "用于需要模拟绝对磁传感器读数的调用方"],
        ],
        [1700, 2550, 5110],
        font_size=8.8,
    )
    add_callout(
        document,
        "方向提醒",
        "磁源法向量表示线圈磁矩方向，不表示所有观测位置的磁场方向。调用方必须使用返回的 Bx、By、Bz 判断局部磁场方向。",
        accent=CAUTION,
    )
    add_page_break(document)

    add_heading(document, "3. 数据结构总览", level=1)
    add_table(
        document,
        ["类型", "职责", "主要使用方"],
        [
            ["Vector3", "表示位置、速度、法向量或磁场三分量", "所有模块"],
            ["Pose3D", "表示设备或磁源的位置与姿态", "运动模块、磁源模块"],
            ["WaveformType", "标识直流、正弦、脉冲或采样波形", "电流控制模块"],
            ["CurrentWaveformParameter", "定义电流随时间变化的参数", "电源/波形模块"],
            ["MagneticSourceParameter", "定义单个线圈、电缆或等效偶极源", "磁源配置模块"],
            ["DeviceParameter", "汇总平台、地磁、运动状态和全部磁源", "配置与任务规划模块"],
            ["ObservationPoint", "定义水雷、传感器或计算点", "任务规划与算法模块"],
            ["CalculationOptions", "选择求解模式和输出分项", "性能控制模块"],
            ["GridParameter", "定义二维或三维规则空间网格", "可视化与场景分析模块"],
            ["TimeSeriesRequest", "定义时间范围和采样率", "时序分析模块"],
            ["MagneticFieldSample", "返回单点单时刻的完整三分量结果", "所有结果消费模块"],
            ["SourceFieldContribution", "可选返回每个磁源的独立贡献", "诊断与标定模块"],
        ],
        [2400, 4300, 2660],
        font_size=8.6,
    )

    add_heading(document, "3.1 基础几何与姿态", level=2)
    add_table(
        document,
        ["结构 / 字段", "类型", "单位", "含义与约束"],
        [
            ["Vector3.x / y / z", "double", "随上下文", "必须为有限数；全局分量按东、北、上解释"],
            ["Pose3D.position", "Vector3", "m", "设备中心或磁源中心位置"],
            ["Pose3D.headingDegrees", "double", "°", "航向角，东向为 0°，逆时针为正"],
            ["Pose3D.pitchDegrees", "double", "°", "抬头为正"],
            ["Pose3D.rollDegrees", "double", "°", "绕设备纵轴旋转，右手定则"],
            ["DeviceParameter.velocityBodyMps", "Vector3", "m/s", "设备局部坐标系速度；支持纵向、横向和垂向运动"],
        ],
        [2850, 1250, 900, 4360],
        font_size=8.6,
    )
    add_code_block(
        document,
        """struct Vector3
{
    double x{0.0}; // 东向坐标、分量或局部纵向分量。
    double y{0.0}; // 北向坐标、分量或局部横向分量。
    double z{0.0}; // 上向坐标、分量或局部垂向分量。
};

struct Pose3D
{
    Vector3 position{};          // 全局位置，单位 m。
    double headingDegrees{0.0};  // 航向角，单位度。
    double pitchDegrees{0.0};    // 俯仰角，单位度。
    double rollDegrees{0.0};     // 横滚角，单位度。
};""",
    )

    add_heading(document, "3.2 电流波形", level=2)
    add_table(
        document,
        ["字段", "类型", "默认值", "说明"],
        [
            ["type", "WaveformType", "DirectCurrent", "直流、正弦、脉冲或外部采样波形"],
            ["amplitudeA", "double", "0", "正弦峰值或脉冲幅值，单位 A"],
            ["offsetA", "double", "0", "直流偏置，单位 A"],
            ["frequencyHz", "double", "0", "周期波形频率；非周期波形可为 0"],
            ["phaseDegrees", "double", "0", "正弦波初相位"],
            ["dutyCycle", "double", "0.5", "脉冲占空比，范围 (0,1]"],
            ["startTimeSeconds", "double", "0", "波形开始生效的仿真时刻"],
            ["sampleTimesSeconds", "vector<double>", "空", "外部采样波形的严格递增时间点；范围外电流按 0 处理"],
            ["sampleCurrentsA", "vector<double>", "空", "与采样时间一一对应，区间内采用线性插值"],
        ],
        [2600, 1800, 1200, 3760],
        font_size=8.5,
    )

    add_heading(document, "3.3 单个磁源", level=2)
    add_table(
        document,
        ["字段", "类型", "含义与约束"],
        [
            ["id", "string", "调用方提供的稳定唯一标识，用于分项输出和日志定位"],
            ["type", "MagneticSourceType", "EquivalentDipole、CircularCoil 或 PolylineCable"],
            ["localPose", "Pose3D", "磁源相对设备中心的安装位置和姿态"],
            ["turns", "double", "线圈匝数，必须为正；电缆模型通常取 1"],
            ["areaSquareMeters", "double", "等效偶极子或闭合线圈面积，单位 m²"],
            ["localNormal", "Vector3", "设备局部坐标系磁矩方向，配置时归一化"],
            ["polylinePoints", "vector<Vector3>", "折线电缆节点，至少两个点；单位 m"],
            ["waveform", "CurrentWaveformParameter", "该磁源的电流波形"],
            ["enabled", "bool", "允许任务规划模块临时关闭某个磁源"],
            ["minimumDistance", "double", "近场奇异保护距离，单位 m，必须为正"],
        ],
        [2350, 2300, 4710],
        font_size=8.5,
    )
    add_callout(
        document,
        "选择规则",
        "EquivalentDipole 适合远场快速计算；CircularCoil 和 PolylineCable 适合近场或几何细节明显的场景。Auto 求解模式可按观测距离与磁源尺寸自动切换。",
    )
    add_heading(document, "3.4 公共枚举", level=2)
    add_table(
        document,
        ["枚举", "取值", "含义"],
        [
            ["WaveformType", "DirectCurrent", "恒定电流，使用 offsetA"],
            ["WaveformType", "Sine", "正弦电流，使用幅值、偏置、频率和相位"],
            ["WaveformType", "Pulse", "周期脉冲，使用幅值、偏置、频率和占空比"],
            ["WaveformType", "Sampled", "调用方提供采样时间和电流，区间内线性插值"],
            ["MagneticSourceType", "EquivalentDipole", "远场等效磁偶极源"],
            ["MagneticSourceType", "CircularCoil", "闭合圆形或等效闭合线圈"],
            ["MagneticSourceType", "PolylineCable", "由折线节点描述的通电电缆"],
            ["FieldSolverMode", "Auto", "按距离和磁源尺寸自动选择算法"],
            ["FieldSolverMode", "Dipole", "强制使用磁偶极子近似"],
            ["FieldSolverMode", "BiotSavart", "强制使用毕奥—萨伐尔线段求和"],
        ],
        [2450, 2650, 4260],
        font_size=8.5,
    )
    add_page_break(document)

    add_heading(document, "4. 设备配置与观测请求", level=1)
    add_heading(document, "4.1 DeviceParameter", level=2)
    add_table(
        document,
        ["字段", "类型", "说明"],
        [
            ["initialPose", "Pose3D", "零时刻扫雷设备中心位置和姿态"],
            ["velocityBodyMps", "Vector3", "设备局部坐标系速度，单位 m/s"],
            ["geomagneticFieldNt", "Vector3", "背景地磁东、北、上三分量，单位 nT"],
            ["sources", "vector<MagneticSourceParameter>", "一个或多个主动磁源，id 不得重复"],
            ["includePlatformField", "bool", "是否叠加扫雷平台自身磁场"],
            ["platformStaticMomentAm2", "Vector3", "平台等效静态磁矩，单位 A·m²"],
            ["platformInducedMomentAm2", "Vector3", "平台等效感应磁矩，单位 A·m²"],
            ["minimumDeviceDistance", "double", "设备整体最小计算距离，单位 m"],
        ],
        [2800, 2700, 3860],
        font_size=8.6,
    )

    add_heading(document, "4.2 ObservationPoint", level=2)
    add_table(
        document,
        ["字段", "类型", "说明"],
        [
            ["id", "string", "水雷、传感器或观测点唯一标识"],
            ["position", "Vector3", "全局坐标，单位 m"],
            ["sensorAxis", "Vector3", "可选敏感轴；非零时自动归一化并计算轴向投影"],
            ["enabled", "bool", "批量仿真时是否启用该观测点"],
        ],
        [2500, 2100, 4760],
        font_size=8.7,
    )

    add_heading(document, "4.3 CalculationOptions", level=2)
    add_table(
        document,
        ["字段", "类型 / 建议默认值", "作用"],
        [
            ["solverMode", "FieldSolverMode::Auto", "自动、偶极子或毕奥—萨伐尔求解"],
            ["includeGeomagneticField", "true", "是否计算 absoluteFieldVector"],
            ["includePlatformField", "true", "是否叠加平台自身磁场"],
            ["returnPerSource", "false", "是否返回各磁源独立贡献；开启后内存占用增加"],
            ["cableSegmentLength", "1.0 m", "折线电缆数值离散长度"],
            ["calculateDerivative", "false", "是否计算 dB/dt；需要相邻时刻结果"],
        ],
        [2800, 2500, 4060],
        font_size=8.6,
    )

    add_heading(document, "4.4 GridParameter 与 TimeSeriesRequest", level=2)
    add_table(
        document,
        ["结构", "字段", "说明"],
        [
            ["GridParameter", "minimum / maximum", "网格全局最小和最大坐标，单位 m"],
            ["GridParameter", "xCount / yCount / zCount", "三个方向采样点数，均必须大于 0"],
            ["TimeSeriesRequest", "startTimeSeconds", "仿真开始时刻"],
            ["TimeSeriesRequest", "durationSeconds", "持续时间，必须非负"],
            ["TimeSeriesRequest", "sampleRateHz", "采样率，必须为有限正数"],
        ],
        [2300, 2700, 4360],
        font_size=8.7,
    )

    add_heading(document, "5. 结果结构 MagneticFieldSample", level=1)
    add_table(
        document,
        ["字段", "类型 / 单位", "语义"],
        [
            ["observationId", "string", "观测点稳定标识"],
            ["observationPosition", "Vector3 / m", "当前观测点全局坐标"],
            ["devicePose", "Pose3D", "当前时刻设备中心位置和姿态"],
            ["time", "double / s", "仿真时刻"],
            ["geomagneticFieldVector", "Vector3 / nT", "背景地磁三分量"],
            ["activeSourceFieldVector", "Vector3 / nT", "所有主动磁源的矢量和"],
            ["platformFieldVector", "Vector3 / nT", "扫雷平台自身磁场"],
            ["anomalyFieldVector", "Vector3 / nT", "主动扫场与平台场的矢量和"],
            ["absoluteFieldVector", "Vector3 / nT", "背景地磁与异常场的矢量和"],
            ["activeSourceMagnitude", "double / nT", "主动磁源场模值"],
            ["anomalyMagnitude", "double / nT", "异常场模值"],
            ["absoluteMagnitude", "double / nT", "绝对场模值"],
            ["sensorAxisProjection", "double / nT", "异常场在观测点敏感轴上的投影"],
            ["fieldDerivativeVector", "Vector3 / nT/s", "可选三分量变化率"],
            ["sourceContributions", "vector<SourceFieldContribution>", "可选每个磁源的独立贡献"],
        ],
        [2950, 2200, 4210],
        font_size=8.4,
    )
    add_callout(
        document,
        "结果口径",
        "anomalyFieldVector 不包含背景地磁；absoluteFieldVector 包含背景地磁。调用方在阈值判断、绘图和数据导出时必须明确使用哪一种口径。",
        accent=CAUTION,
    )
    add_heading(document, "5.1 辅助结果与校验结构", level=2)
    add_table(
        document,
        ["结构 / 字段", "类型", "说明"],
        [
            ["SourceFieldContribution.sourceId", "string", "产生该分项结果的磁源 id"],
            ["SourceFieldContribution.currentA", "double", "当前时刻该磁源电流，单位 A"],
            ["SourceFieldContribution.fieldVector", "Vector3", "该磁源产生的 Bx、By、Bz，单位 nT"],
            ["SourceFieldContribution.magnitude", "double", "该磁源场模值，单位 nT"],
            ["SourceFieldContribution.solverMode", "FieldSolverMode", "本次实际采用的求解模式"],
            ["ValidationIssue.fieldPath", "string", "错误字段路径，例如 sources[0].turns"],
            ["ValidationIssue.message", "string", "面向界面和日志的中文错误说明"],
            ["ValidationResult.valid", "bool", "全部配置合法时为 true"],
            ["ValidationResult.message", "string", "首个或汇总错误信息，便于简单调用方直接显示"],
            ["ValidationResult.issues", "vector<ValidationIssue>", "完整校验问题列表"],
        ],
        [3550, 2300, 3510],
        font_size=8.4,
    )
    add_page_break(document)

    add_heading(document, "6. 公共函数接口", level=1)
    add_table(
        document,
        ["函数", "返回值", "用途"],
        [
            ["构造函数 / setDevice", "void", "校验并保存完整设备配置"],
            ["device", "const DeviceParameter&", "读取当前配置"],
            ["calculate", "MagneticFieldSample", "计算单个观测点、单个时刻"],
            ["simulate", "vector<MagneticFieldSample>", "生成固定观测点时间序列"],
            ["simulateGrid", "vector<MagneticFieldSample>", "生成指定时刻规则空间网格"],
            ["simulatePoints", "vector<MagneticFieldSample>", "批量计算多个非规则观测点"],
            ["evaluateCurrent", "double", "查询指定磁源在某时刻的实际电流"],
            ["validateDevice", "ValidationResult", "在不改变模型状态时预校验配置"],
        ],
        [3000, 2800, 3560],
        font_size=8.6,
    )

    add_heading(document, "6.1 建议的类声明", level=2)
    add_code_block(
        document,
        """class MinesweepingMagneticFieldModel
{
public:
    /** 使用完整设备参数创建可立即计算的模型。 */
    explicit MinesweepingMagneticFieldModel(const DeviceParameter& parameter);

    /** 校验并更新设备参数；失败时保持原配置不变。 */
    void setDevice(const DeviceParameter& parameter);

    /** 返回当前生效的设备参数。 */
    const DeviceParameter& device() const;

    /** 计算单个观测点在指定时刻的三分量磁场。 */
    MagneticFieldSample calculate(
        const ObservationPoint& observation,
        double timeSeconds,
        const CalculationOptions& options = {}) const;

    /** 在固定观测点生成连续磁场时间序列。 */
    std::vector<MagneticFieldSample> simulate(
        const ObservationPoint& observation,
        const TimeSeriesRequest& request,
        const CalculationOptions& options = {}) const;

    /** 生成指定时刻的规则二维或三维空间磁场网格。 */
    std::vector<MagneticFieldSample> simulateGrid(
        const GridParameter& grid,
        double timeSeconds,
        const CalculationOptions& options = {}) const;

    /** 批量计算多个非规则观测点。 */
    std::vector<MagneticFieldSample> simulatePoints(
        const std::vector<ObservationPoint>& observations,
        double timeSeconds,
        const CalculationOptions& options = {}) const;

    /** 返回指定磁源在某一时刻的电流值，单位 A。 */
    double evaluateCurrent(
        const std::string& sourceId,
        double timeSeconds) const;

    /** 在不修改模型状态的情况下检查设备配置。 */
    static ValidationResult validateDevice(
        const DeviceParameter& parameter) noexcept;
};""",
    )

    add_heading(document, "6.2 函数调用契约", level=2)
    add_table(
        document,
        ["函数", "前置条件", "结果顺序 / 说明"],
        [
            ["calculate", "模型已配置；时间和坐标有限", "返回一个完整样本"],
            ["simulate", "持续时间≥0；采样率>0", "包含起点及持续时间内最后一个完整采样点"],
            ["simulateGrid", "边界有序；各方向点数>0", "按 z、y、x 顺序展开"],
            ["simulatePoints", "观测点 id 不重复；坐标有限", "结果顺序与输入观测点顺序一致"],
            ["evaluateCurrent", "sourceId 存在；时间有限", "只计算波形，不进行磁场求解"],
            ["setDevice", "全部磁源 id 唯一且参数合法", "先校验临时状态，成功后再替换当前配置"],
        ],
        [2200, 3300, 3860],
        font_size=8.6,
    )

    add_heading(document, "7. 推荐调用流程", level=1)
    for text in [
        "构建设备初始姿态、速度和背景地磁参数。",
        "依次配置每个线圈、电缆或等效偶极源及其电流波形。",
        "调用 validateDevice 预检查配置，并把错误反馈给界面或任务规划模块。",
        "构造 MinesweepingMagneticFieldModel，配置成功后只读共享。",
        "根据业务需要调用 calculate、simulate、simulateGrid 或 simulatePoints。",
        "根据异常场、绝对场或敏感轴投影选择正确结果字段。",
    ]:
        add_numbered(document, text, decimal_id)

    add_heading(document, "7.1 调用示例", level=2)
    add_code_block(
        document,
        """using Model = MinesweepingMagneticFieldModel;

Model::DeviceParameter device;
device.initialPose.position = {0.0, 0.0, -5.0};       // 设备初始位置，单位 m。
device.initialPose.headingDegrees = 0.0;              // 舰首朝东。
device.velocityBodyMps = {4.0, 0.0, 0.0};             // 沿设备纵轴匀速运动。
device.geomagneticFieldNt = {0.0, 30000.0, -40000.0}; // 背景地磁东、北、上分量。

Model::MagneticSourceParameter coil;
coil.id = "main_coil";                                // 磁源稳定标识。
coil.type = Model::MagneticSourceType::CircularCoil;  // 使用闭合线圈模型。
coil.turns = 20.0;                                    // 线圈匝数。
coil.areaSquareMeters = 12.0;                         // 等效线圈面积，单位 m²。
coil.localNormal = {1.0, 0.0, 0.0};                   // 磁矩沿设备纵轴。
coil.waveform.type = Model::WaveformType::Sine;       // 使用正弦激励。
coil.waveform.amplitudeA = 500.0;                     // 峰值电流，单位 A。
coil.waveform.frequencyHz = 2.0;                      // 激励频率，单位 Hz。
device.sources.push_back(coil);                       // 将磁源加入设备配置。

const auto validation = Model::validateDevice(device); // 在构造模型前进行参数检查。
if (!validation.valid)
{
    throw std::invalid_argument(validation.message);    // 把中文错误信息交给上层处理。
}

const Model model(device);                              // 创建只读计算模型。

Model::ObservationPoint mine;
mine.id = "mine_001";                                   // 水雷或观测点标识。
mine.position = {0.0, 30.0, -40.0};                    // 水雷位置，单位 m。
mine.sensorAxis = {0.0, 1.0, 0.0};                     // 磁传感器敏感轴指向北。

Model::TimeSeriesRequest request;
request.startTimeSeconds = 0.0;                         // 仿真起始时刻。
request.durationSeconds = 60.0;                         // 仿真持续 60 s。
request.sampleRateHz = 50.0;                            // 每秒采样 50 次。

Model::CalculationOptions options;
options.solverMode = Model::FieldSolverMode::Auto;      // 自动选择远场或近场算法。
options.includeGeomagneticField = true;                 // 同时返回绝对磁场。

const auto samples = model.simulate(mine, request, options); // 生成三分量时序。""",
    )
    add_page_break(document)

    add_heading(document, "8. 参数校验与异常约定", level=1)
    add_table(
        document,
        ["异常类型", "触发条件", "调用方处理建议"],
        [
            ["std::invalid_argument", "坐标非有限、频率为负、匝数或面积不合法、id 重复", "提示配置错误并拒绝开始仿真"],
            ["std::logic_error", "模型尚未成功配置", "检查初始化顺序"],
            ["std::domain_error", "观测点进入奇异保护距离或线缆几何退化", "调整网格、保护距离或求解模式"],
            ["std::length_error", "时间样本数或网格点数超过安全上限", "降低采样率、缩小网格或分批调用"],
            ["std::out_of_range", "evaluateCurrent 使用不存在的 sourceId", "刷新磁源列表并检查标识"],
            ["std::runtime_error", "数值积分未收敛或内部求解失败", "记录场景参数并降级到偶极子模式"],
        ],
        [2300, 4150, 2910],
        font_size=8.5,
    )
    add_heading(document, "8.1 必须执行的配置校验", level=2)
    for text in [
        "所有 double 和 Vector3 分量必须为有限数。",
        "磁源 id 和观测点 id 在各自集合中必须唯一。",
        "直流波形允许 frequencyHz=0；周期波形必须 frequencyHz>0。",
        "采样波形的时间和值数量必须一致，时间必须严格递增。",
        "线圈匝数、面积和最小距离必须为有限正数。",
        "折线电缆至少包含两个不重合节点。",
        "时间采样数和网格总点数不得超过模型安全上限。",
    ]:
        add_bullet(document, text, bullet_id)

    add_heading(document, "9. 并发、性能与扩展约定", level=1)
    add_body(
        document,
        "模型完成配置后，calculate、simulate、simulateGrid、simulatePoints 和 evaluateCurrent 均应设计为 const，只读调用可以并发执行。setDevice 会修改内部缓存，不得与计算函数并发调用。"
    )
    add_table(
        document,
        ["主题", "建议约定"],
        [
            ["线程安全", "配置完成后的 const 计算接口可并发；配置更新由调用方串行化"],
            ["求解缓存", "缓存磁源局部几何、线段离散结果和不随时间变化的旋转数据"],
            ["批量性能", "优先使用 simulatePoints 或 simulateGrid，避免调用方逐点重复创建模型"],
            ["远近场切换", "Auto 模式依据距离/尺寸比自动选择；阈值作为内部策略或高级选项"],
            ["扩展磁源", "新增磁源类型时扩展 MagneticSourceType 和内部求解策略，不改变结果结构"],
            ["轨迹扩展", "未来可增加 MotionProvider，以支持转弯、升沉和外部航迹输入"],
            ["序列化", "配置层可另行提供 JSON 适配器，核心计算接口不依赖具体配置文件格式"],
        ],
        [2200, 7160],
        font_size=8.7,
    )

    add_heading(document, "10. 与其他功能模块的调用关系", level=1)
    add_table(
        document,
        ["调用模块", "主要输入", "主要使用接口 / 结果"],
        [
            ["任务规划", "航迹、速度、设备配置、水雷候选位置", "simulatePoints、sensorAxisProjection、anomalyMagnitude"],
            ["可视化", "规则空间范围与显示时刻", "simulateGrid、activeSourceFieldVector、anomalyFieldVector"],
            ["时序分析", "固定观测点、持续时间和采样率", "simulate、Bx/By/Bz、fieldDerivativeVector"],
            ["数据导出", "任意计算结果集合", "MagneticFieldSample 全字段及 sourceContributions"],
            ["参数标定", "实测电流和磁场数据", "evaluateCurrent、returnPerSource=true"],
            ["仿真控制", "运行、暂停、批次和场景参数", "模型生命周期及异常状态"],
        ],
        [1900, 3450, 4010],
        font_size=8.6,
    )
    add_callout(
        document,
        "集成建议",
        "业务模块只依赖公共数据结构和 const 计算接口，不直接访问内部磁源求解器。这样可以在不影响上层调用的情况下替换偶极子、圆线圈或电缆算法。",
    )

    add_heading(document, "11. 接口验收清单", level=1)
    for text in [
        "单个直流等效偶极源在轴线和赤道面的方向、符号及 1/r³ 衰减正确。",
        "两个相反磁矩磁源能够按三分量矢量抵消，而不是按模值相加。",
        "正弦电流的 Bx、By、Bz 与电流保持相同频率和正确相位。",
        "设备平移和姿态旋转后，磁源位置、法向量和输出分量同步变化。",
        "simulate 的采样点数与起止时间契约和现有场模型保持一致。",
        "simulateGrid 按 z、y、x 顺序输出，调用方可稳定恢复规则网格。",
        "异常场与绝对场字段口径通过单元测试区分。",
        "非法配置不会破坏已经生效的模型状态。",
        "大网格和高采样率请求能够触发长度保护，不会意外耗尽内存。",
        "所有公共类型、成员字段和函数均提供中文注释和明确单位。",
    ]:
        add_bullet(document, text, bullet_id)

    add_heading(document, "12. 推荐实施顺序", level=1)
    for text in [
        "先实现 Vector3、Pose3D、波形、磁源、设备和结果结构及完整校验。",
        "实现 EquivalentDipole 与直流/正弦波形，打通 calculate 和 simulate。",
        "实现 simulateGrid、simulatePoints 和批量性能优化。",
        "加入 CircularCoil、PolylineCable 及 Auto 远近场切换。",
        "加入平台磁场、敏感轴投影、dB/dt 和单磁源分项输出。",
        "最后补充 JSON 适配、标定工具和复杂外部航迹接口。",
    ]:
        add_numbered(document, text, decimal_id)
    add_callout(
        document,
        "最终建议",
        "V1 接口优先保证三分量口径、单位、时间采样和结果顺序稳定；复杂线缆动力学、海底介质和外部轨迹可作为后续扩展，避免首版接口过度耦合。",
    )

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


def main() -> None:
    """生成并保存接口说明 Word 文档。"""

    output = build_document()
    print(f"已生成 Word 文档：{output}")


if __name__ == "__main__":
    main()
