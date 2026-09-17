"""生成通俗优化版的水下目标静电场与轴频电场仿真原理 Word 文档。"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# 文档采用 compact_reference_guide 预设，并沿用原文档的蓝色技术手册风格。
BODY_FONT = "Calibri"
CHINESE_FONT = "Microsoft YaHei"
MATH_FONT = "Cambria Math"
COLOR_NAVY = "17365D"
COLOR_BLUE = "2E74B5"
COLOR_DARK_BLUE = "1F4D78"
COLOR_TEAL = "1F8C85"
COLOR_GOLD = "D28B16"
COLOR_MUTED = "667085"
COLOR_TEXT = "26364A"
COLOR_TABLE_HEADER = "E8EEF5"
COLOR_TABLE_ALT = "F7F9FC"
COLOR_BORDER = "BAC6D3"
COLOR_INFO_FILL = "EEF5FB"
COLOR_CAUTION_FILL = "FFF6E5"
COLOR_KEY_FILL = "EAF7F4"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


# 每条公式对应的 LaTeX 源码，用于维护、审计和无障碍说明。
LATEX_FORMULAS = {
    "ΔV ≈ −E · L": r"\Delta V \approx -\mathbf{E}\cdot\mathbf{L}",
    "J = σE": r"\mathbf{J}=\sigma\mathbf{E}",
    "φ(r) = I / (4πσr)": r"\varphi(r)=\frac{I}{4\pi\sigma r}",
    "p = I · d": r"\mathbf{p}=I\mathbf{d}",
    "φ(r) = (p · r) / (4πσr³)": r"\varphi(\mathbf{r})=\frac{\mathbf{p}\cdot\mathbf{r}}{4\pi\sigma r^3}",
    "E(r) = [3r̂(p · r̂) − p] / (4πσr³)": r"\mathbf{E}(\mathbf{r})=\frac{3\hat{\mathbf{r}}(\mathbf{p}\cdot\hat{\mathbf{r}})-\mathbf{p}}{4\pi\sigma r^3}",
    "|E| = √(Eₓ² + Eᵧ² + E_z²)": r"\lvert\mathbf{E}\rvert=\sqrt{E_x^2+E_y^2+E_z^2}",
    "I_est ≈ j_corr · S_wet · η_damage": r"I_{\mathrm{est}}\approx j_{\mathrm{corr}}S_{\mathrm{wet}}\eta_{\mathrm{damage}}",
    "p_static ≈ I_est · d_eff": r"p_{\mathrm{static}}\approx I_{\mathrm{est}}d_{\mathrm{eff}}",
    "R(t) = R₀[1 + m cos(2πf_st + φ)]": r"R(t)=R_0\left[1+m\cos(2\pi f_st+\varphi)\right]",
    "f_s = n / 60": r"f_s=\frac{n}{60}",
    "p_shaft(t) = p₁cos(2πf_st+φ₁) + p₂cos(4πf_st+φ₂) + p₃cos(6πf_st+φ₃) + …": r"p_{\mathrm{shaft}}(t)=p_1\cos(2\pi f_st+\varphi_1)+p_2\cos(4\pi f_st+\varphi_2)+p_3\cos(6\pi f_st+\varphi_3)+\cdots",
    "r_T(t) = r_T0 + v̂vt": r"\mathbf{r}_T(t)=\mathbf{r}_{T0}+\hat{\mathbf{v}}vt",
    "R(t) = r_S − r_T(t)": r"\mathbf{R}(t)=\mathbf{r}_S-\mathbf{r}_T(t)",
    "R(t) = √[b² + (vt)²]": r"R(t)=\sqrt{b^2+(vt)^2}",
    "A(t) ∝ 1 / [b² + (vt)²]³ᐟ²": r"A(t)\propto\frac{1}{\left[b^2+(vt)^2\right]^{3/2}}",
    "E_total(t) = E_static(t) + E_shaft(t)": r"\mathbf{E}_{\mathrm{total}}(t)=\mathbf{E}_{\mathrm{static}}(t)+\mathbf{E}_{\mathrm{shaft}}(t)",
    "f_sample > 2f_max": r"f_{\mathrm{sample}}>2f_{\max}",
    "频率分辨率 Δf ≈ 1 / T": r"\Delta f\approx\frac{1}{T}",
    "t_CPA = 300 / 8 = 37.5 s": r"t_{\mathrm{CPA}}=\frac{300}{8}=37.5\,\mathrm{s}",
    "b = √(100² + 25²) ≈ 103.1 m": r"b=\sqrt{100^2+25^2}\approx103.1\,\mathrm{m}",
    "(317.2 / 103.1)³ ≈ 29": r"\left(\frac{317.2}{103.1}\right)^3\approx29",
}


def math_run(text: str, style: str = "i"):
    """创建一个 Office Math 字符节点，style 可取 i、p、b 或 bi。"""
    run = OxmlElement("m:r")
    run_properties = OxmlElement("m:rPr")
    math_style = OxmlElement("m:sty")
    math_style.set(qn("m:val"), style)
    run_properties.append(math_style)
    run.append(run_properties)
    word_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), MATH_FONT)
    fonts.set(qn("w:hAnsi"), MATH_FONT)
    fonts.set(qn("w:eastAsia"), MATH_FONT)
    word_properties.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "25")
    word_properties.append(size)
    complex_size = OxmlElement("w:szCs")
    complex_size.set(qn("w:val"), "25")
    word_properties.append(complex_size)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), COLOR_NAVY)
    word_properties.append(color)
    run.append(word_properties)
    value = OxmlElement("m:t")
    value.text = text
    if text != text.strip():
        value.set(qn("xml:space"), "preserve")
    run.append(value)
    return run


def math_sequence(*items):
    """把嵌套序列展平成一组 Office Math 节点。"""
    result = []
    for item in items:
        if isinstance(item, (list, tuple)):
            result.extend(math_sequence(*item))
        else:
            result.append(item)
    return result


def append_math_items(parent, items) -> None:
    """向指定 Office Math 容器追加节点。"""
    for item in math_sequence(items):
        parent.append(item)


def math_subscript(base, subscript):
    """创建下标结构。"""
    node = OxmlElement("m:sSub")
    base_node = OxmlElement("m:e")
    append_math_items(base_node, base)
    sub_node = OxmlElement("m:sub")
    append_math_items(sub_node, subscript)
    node.extend([base_node, sub_node])
    return node


def math_superscript(base, superscript):
    """创建上标结构。"""
    node = OxmlElement("m:sSup")
    base_node = OxmlElement("m:e")
    append_math_items(base_node, base)
    sup_node = OxmlElement("m:sup")
    append_math_items(sup_node, superscript)
    node.extend([base_node, sup_node])
    return node


def math_fraction(numerator, denominator):
    """创建带横线的分式。"""
    node = OxmlElement("m:f")
    numerator_node = OxmlElement("m:num")
    append_math_items(numerator_node, numerator)
    denominator_node = OxmlElement("m:den")
    append_math_items(denominator_node, denominator)
    node.extend([numerator_node, denominator_node])
    return node


def math_radical(body):
    """创建平方根结构。"""
    node = OxmlElement("m:rad")
    properties = OxmlElement("m:radPr")
    hide_degree = OxmlElement("m:degHide")
    hide_degree.set(qn("m:val"), "1")
    properties.append(hide_degree)
    degree = OxmlElement("m:deg")
    body_node = OxmlElement("m:e")
    append_math_items(body_node, body)
    node.extend([properties, degree, body_node])
    return node


def math_accent(body, character: str = "\u0302"):
    """创建帽子等重音结构。"""
    node = OxmlElement("m:acc")
    properties = OxmlElement("m:accPr")
    accent = OxmlElement("m:chr")
    accent.set(qn("m:val"), character)
    properties.append(accent)
    body_node = OxmlElement("m:e")
    append_math_items(body_node, body)
    node.extend([properties, body_node])
    return node


def math_delimiter(body, begin: str = "(", end: str = ")"):
    """创建可随内容自动伸缩的括号。"""
    node = OxmlElement("m:d")
    properties = OxmlElement("m:dPr")
    begin_character = OxmlElement("m:begChr")
    begin_character.set(qn("m:val"), begin)
    end_character = OxmlElement("m:endChr")
    end_character.set(qn("m:val"), end)
    properties.extend([begin_character, end_character])
    body_node = OxmlElement("m:e")
    append_math_items(body_node, body)
    node.extend([properties, body_node])
    return node


def build_office_math_formula(formula: str):
    """把本文已知公式构造成 Word 原生 Office Math 节点。"""
    r = math_run
    sub = math_subscript
    sup = math_superscript
    frac = math_fraction
    root = math_radical
    hat = math_accent
    delim = math_delimiter
    op = lambda value: r(value, "p")
    bold = lambda value: r(value, "bi")

    formulas = {
        "ΔV ≈ −E · L": [r("Δ"), r("V"), op("≈−"), bold("E"), op("·"), bold("L")],
        "J = σE": [bold("J"), op("="), r("σ"), bold("E")],
        "φ(r) = I / (4πσr)": [r("φ"), delim([r("r")]), op("="), frac([r("I")], [op("4"), r("πσr")])],
        "p = I · d": [bold("p"), op("="), r("I"), op("·"), bold("d")],
        "φ(r) = (p · r) / (4πσr³)": [
            r("φ"), delim([bold("r")]), op("="),
            frac([bold("p"), op("·"), bold("r")], [op("4"), r("πσ"), sup([r("r")], [op("3")])]),
        ],
        "E(r) = [3r̂(p · r̂) − p] / (4πσr³)": [
            bold("E"), delim([bold("r")]), op("="),
            frac(
                [op("3"), hat([bold("r")]), delim([bold("p"), op("·"), hat([bold("r")])]), op("−"), bold("p")],
                [op("4"), r("πσ"), sup([r("r")], [op("3")])],
            ),
        ],
        "|E| = √(Eₓ² + Eᵧ² + E_z²)": [
            op("|"), bold("E"), op("|="),
            root([
                sup([sub([r("E")], [r("x")])], [op("2")]), op("+"),
                sup([sub([r("E")], [r("y")])], [op("2")]), op("+"),
                sup([sub([r("E")], [r("z")])], [op("2")]),
            ]),
        ],
        "I_est ≈ j_corr · S_wet · η_damage": [
            sub([r("I")], [op("est")]), op("≈"), sub([r("j")], [op("corr")]), op("·"),
            sub([r("S")], [op("wet")]), op("·"), sub([r("η")], [op("damage")]),
        ],
        "p_static ≈ I_est · d_eff": [
            sub([r("p")], [op("static")]), op("≈"), sub([r("I")], [op("est")]), op("·"),
            sub([r("d")], [op("eff")]),
        ],
        "R(t) = R₀[1 + m cos(2πf_st + φ)]": [
            r("R"), delim([r("t")]), op("="), sub([r("R")], [op("0")]),
            delim([op("1+"), r("m"), op("cos"), delim([op("2"), r("π"), sub([r("f")], [r("s")]), r("t"), op("+"), r("φ")])], "[", "]"),
        ],
        "f_s = n / 60": [sub([r("f")], [r("s")]), op("="), frac([r("n")], [op("60")])],
        "p_shaft(t) = p₁cos(2πf_st+φ₁) + p₂cos(4πf_st+φ₂) + p₃cos(6πf_st+φ₃) + …": [
            sub([r("p")], [op("shaft")]), delim([r("t")]), op("="),
            sub([r("p")], [op("1")]), op("cos"), delim([op("2"), r("π"), sub([r("f")], [r("s")]), r("t"), op("+"), sub([r("φ")], [op("1")])]),
            op("+"), sub([r("p")], [op("2")]), op("cos"), delim([op("4"), r("π"), sub([r("f")], [r("s")]), r("t"), op("+"), sub([r("φ")], [op("2")])]),
            op("+"), sub([r("p")], [op("3")]), op("cos"), delim([op("6"), r("π"), sub([r("f")], [r("s")]), r("t"), op("+"), sub([r("φ")], [op("3")])]), op("+⋯"),
        ],
        "r_T(t) = r_T0 + v̂vt": [
            sub([bold("r")], [r("T")]), delim([r("t")]), op("="), sub([bold("r")], [r("T"), op("0")]),
            op("+"), hat([bold("v")]), r("v"), r("t"),
        ],
        "R(t) = r_S − r_T(t)": [
            bold("R"), delim([r("t")]), op("="), sub([bold("r")], [r("S")]), op("−"),
            sub([bold("r")], [r("T")]), delim([r("t")]),
        ],
        "R(t) = √[b² + (vt)²]": [
            r("R"), delim([r("t")]), op("="), root([
                sup([r("b")], [op("2")]), op("+"), sup([delim([r("v"), r("t")])], [op("2")]),
            ]),
        ],
        "A(t) ∝ 1 / [b² + (vt)²]³ᐟ²": [
            r("A"), delim([r("t")]), op("∝"),
            frac([op("1")], [sup([delim([sup([r("b")], [op("2")]), op("+"), sup([delim([r("v"), r("t")])], [op("2")])], "[", "]")], [frac([op("3")], [op("2")])])]),
        ],
        "E_total(t) = E_static(t) + E_shaft(t)": [
            sub([bold("E")], [op("total")]), delim([r("t")]), op("="),
            sub([bold("E")], [op("static")]), delim([r("t")]), op("+"),
            sub([bold("E")], [op("shaft")]), delim([r("t")]),
        ],
        "f_sample > 2f_max": [
            sub([r("f")], [op("sample")]), op(">2"), sub([r("f")], [op("max")]),
        ],
        "频率分辨率 Δf ≈ 1 / T": [r("Δ"), r("f"), op("≈"), frac([op("1")], [r("T")])],
        "t_CPA = 300 / 8 = 37.5 s": [
            sub([r("t")], [op("CPA")]), op("="), frac([op("300")], [op("8")]), op("=37.5"), op(" s"),
        ],
        "b = √(100² + 25²) ≈ 103.1 m": [
            r("b"), op("="), root([sup([op("100")], [op("2")]), op("+"), sup([op("25")], [op("2")])]), op("≈103.1"), op(" m"),
        ],
        "(317.2 / 103.1)³ ≈ 29": [
            sup([delim([frac([op("317.2")], [op("103.1")])])], [op("3")]), op("≈29"),
        ],
    }

    if formula not in formulas:
        raise KeyError(f"尚未定义公式的 Office Math 结构：{formula}")
    math = OxmlElement("m:oMath")
    append_math_items(math, formulas[formula])
    return math


def set_run_font(
    run,
    *,
    name: str = BODY_FONT,
    chinese_name: str = CHINESE_FONT,
    size: float = 11,
    color: str = COLOR_TEXT,
    bold: bool = False,
    italic: bool = False,
) -> None:
    """为中英文设置明确字体，避免不同软件打开时发生字体漂移。"""
    run.font.name = name
    properties = run._element.get_or_add_rPr()
    properties.rFonts.set(qn("w:ascii"), name)
    properties.rFonts.set(qn("w:hAnsi"), name)
    properties.rFonts.set(qn("w:eastAsia"), chinese_name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_paragraph_format(
    paragraph,
    *,
    before: float = 0,
    after: float = 6,
    line_spacing: float = 1.25,
    keep_with_next: bool = False,
) -> None:
    """统一设置段落节奏。"""
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.keep_with_next = keep_with_next


def configure_style(
    document: Document,
    style_name: str,
    *,
    size: float,
    color: str,
    bold: bool,
    before: float,
    after: float,
) -> None:
    """把预设值写入真实 Word 样式。"""
    style = document.styles[style_name]
    style.font.name = BODY_FONT
    properties = style._element.get_or_add_rPr()
    properties.rFonts.set(qn("w:ascii"), BODY_FONT)
    properties.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    properties.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    style.paragraph_format.line_spacing = 1.25
    style.paragraph_format.keep_with_next = True


def add_page_field(paragraph) -> None:
    """在页脚插入动态页码。"""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=9, color=COLOR_MUTED)


def add_numbering_definitions(document: Document) -> tuple[int, int]:
    """创建符合紧凑参考手册预设的真实项目符号和数字编号。"""
    numbering = document.part.numbering_part.element

    def create_definition(num_id: int, abstract_id: int, bullet: bool) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)

        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        number_format = OxmlElement("w:numFmt")
        number_format.set(qn("w:val"), "bullet" if bullet else "decimal")
        level.append(number_format)
        level_text = OxmlElement("w:lvlText")
        level_text.set(qn("w:val"), "•" if bullet else "%1.")
        level.append(level_text)

        paragraph_properties = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        paragraph_properties.append(tabs)
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), "540")
        indent.set(qn("w:hanging"), "270")
        paragraph_properties.append(indent)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        paragraph_properties.append(spacing)
        level.append(paragraph_properties)
        numbering.append(abstract)

        number = OxmlElement("w:num")
        number.set(qn("w:numId"), str(num_id))
        reference = OxmlElement("w:abstractNumId")
        reference.set(qn("w:val"), str(abstract_id))
        number.append(reference)
        numbering.append(number)

    create_definition(91, 91, True)
    create_definition(92, 92, False)
    return 91, 92


def apply_numbering(paragraph, num_id: int) -> None:
    """把段落关联到真实编号定义。"""
    properties = paragraph._p.get_or_add_pPr()
    number_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    number_properties.extend([level, number])
    properties.append(number_properties)


def set_cell_margins(cell) -> None:
    """设置表格单元格内边距，保证文字不贴边。"""
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", 90), ("bottom", 90), ("start", 120), ("end", 120)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    """设置单元格底色。"""
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_table_borders(table, color: str = COLOR_BORDER) -> None:
    """设置轻量表格边框。"""
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_repeat_header(row) -> None:
    """让跨页表格自动重复表头。"""
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    """使用固定 DXA 宽度，保证 Word 和其他阅读器中的表格几何一致。"""
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError("表格列宽总和必须为 9360 DXA")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr

    width = properties.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        properties.insert(0, width)
    width.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    width.set(qn("w:type"), "dxa")

    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for column_width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(column_width))
        grid.append(column)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            column_width = widths_dxa[index]
            properties = cell._tc.get_or_add_tcPr()
            cell_width = properties.first_child_found_in("w:tcW")
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                properties.append(cell_width)
            cell_width.set(qn("w:w"), str(column_width))
            cell_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def configure_document(document: Document) -> tuple[int, int]:
    """配置页面、正文、标题、页眉页脚和编号。"""
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    properties = normal._element.get_or_add_rPr()
    properties.rFonts.set(qn("w:ascii"), BODY_FONT)
    properties.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    properties.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(COLOR_TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.25

    configure_style(document, "Heading 1", size=16, color=COLOR_BLUE, bold=True, before=18, after=10)
    configure_style(document, "Heading 2", size=13, color=COLOR_BLUE, bold=True, before=14, after=7)
    configure_style(document, "Heading 3", size=12, color=COLOR_DARK_BLUE, bold=True, before=10, after=5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_format(header, after=0, line_spacing=1.0)
    run = header.add_run("水下目标静电场与轴频电场｜通俗原理说明")
    set_run_font(run, size=8.5, color=COLOR_MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_format(footer, after=0, line_spacing=1.0)
    run = footer.add_run("第 ")
    set_run_font(run, size=9, color=COLOR_MUTED)
    add_page_field(footer)
    run = footer.add_run(" 页")
    set_run_font(run, size=9, color=COLOR_MUTED)
    return add_numbering_definitions(document)


class GuideWriter:
    """封装常用排版操作，使正文内容更容易维护。"""

    def __init__(self, document: Document, bullet_id: int, number_id: int) -> None:
        self.document = document
        self.bullet_id = bullet_id
        self.number_id = number_id

    def body(self, text: str, *, bold_lead: str | None = None) -> None:
        """添加正文段落，可选加粗开头标签。"""
        paragraph = self.document.add_paragraph()
        set_paragraph_format(paragraph)
        if bold_lead and text.startswith(bold_lead):
            lead = paragraph.add_run(bold_lead)
            set_run_font(lead, bold=True, color=COLOR_DARK_BLUE)
            text = text[len(bold_lead) :]
        run = paragraph.add_run(text)
        set_run_font(run)

    def heading1(self, text: str) -> None:
        """添加章标题。"""
        self.document.add_paragraph(text, style="Heading 1")

    def heading2(self, text: str) -> None:
        """添加二级标题。"""
        self.document.add_paragraph(text, style="Heading 2")

    def heading3(self, text: str) -> None:
        """添加三级标题。"""
        self.document.add_paragraph(text, style="Heading 3")

    def bullet(self, text: str) -> None:
        """添加真实项目符号。"""
        paragraph = self.document.add_paragraph()
        apply_numbering(paragraph, self.bullet_id)
        run = paragraph.add_run(text)
        set_run_font(run)

    def number(self, text: str) -> None:
        """添加真实数字编号。"""
        paragraph = self.document.add_paragraph()
        apply_numbering(paragraph, self.number_id)
        run = paragraph.add_run(text)
        set_run_font(run)

    def equation(self, formula: str, explanation: str | None = None) -> None:
        """添加 LaTeX 结构对应的 Word 原生公式和通俗解释。"""
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(paragraph, before=3, after=3, line_spacing=1.0, keep_with_next=bool(explanation))
        properties = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F4F6F9")
        properties.append(shading)
        paragraph._p.append(build_office_math_formula(formula))
        if explanation:
            note = self.document.add_paragraph()
            note.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_format(note, after=7, line_spacing=1.15)
            run = note.add_run(explanation)
            set_run_font(run, size=9.5, color=COLOR_MUTED)

    def callout(self, title: str, text: str, kind: str = "info") -> None:
        """添加重点、提醒或结论卡片。"""
        fill = {
            "info": COLOR_INFO_FILL,
            "key": COLOR_KEY_FILL,
            "caution": COLOR_CAUTION_FILL,
        }[kind]
        accent = {
            "info": COLOR_BLUE,
            "key": COLOR_TEAL,
            "caution": COLOR_GOLD,
        }[kind]
        # 提示卡采用带底色和边框的普通段落，避免把非表格内容伪装成表格。
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.08)
        paragraph.paragraph_format.right_indent = Inches(0.08)
        set_paragraph_format(paragraph, before=3, after=8, line_spacing=1.2)
        properties = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        properties.append(shading)
        borders = OxmlElement("w:pBdr")
        for edge in ("top", "left", "bottom", "right"):
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "10")
            border.set(qn("w:space"), "6")
            border.set(qn("w:color"), accent)
            borders.append(border)
        properties.append(borders)
        lead = paragraph.add_run(f"{title}：")
        set_run_font(lead, size=10.5, color=accent, bold=True)
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5, color=COLOR_TEXT)

    def table(self, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
        """添加固定列宽的对照表。"""
        table = self.document.add_table(rows=len(rows) + 1, cols=len(headers))
        set_table_geometry(table, widths)
        set_table_borders(table)
        set_repeat_header(table.rows[0])
        for column, header in enumerate(headers):
            cell = table.cell(0, column)
            set_cell_shading(cell, COLOR_TABLE_HEADER)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_format(paragraph, after=0, line_spacing=1.1)
            run = paragraph.add_run(header)
            set_run_font(run, size=9.5, color=COLOR_DARK_BLUE, bold=True)

        for row_index, row in enumerate(rows, start=1):
            for column, value in enumerate(row):
                cell = table.cell(row_index, column)
                if row_index % 2 == 0:
                    set_cell_shading(cell, COLOR_TABLE_ALT)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                paragraph = cell.paragraphs[0]
                if column < len(row) - 1 and len(value) < 18:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_paragraph_format(paragraph, after=0, line_spacing=1.15)
                run = paragraph.add_run(value)
                set_run_font(run, size=9.3)
        spacer = self.document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)

    def figure(self, image_path: Path, caption: str, reading_tip: str, width: float = 6.35) -> None:
        """添加内嵌图片、标题和读图提示。"""
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(paragraph, before=4, after=3, line_spacing=1.0, keep_with_next=True)
        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Inches(width))
        # 为图片写入替代文字，便于无障碍阅读器说明图片内容。
        for doc_properties in run._r.xpath(".//wp:docPr"):
            doc_properties.set("title", caption)
            doc_properties.set("descr", reading_tip)

        caption_paragraph = self.document.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(caption_paragraph, after=3, line_spacing=1.15, keep_with_next=True)
        run = caption_paragraph.add_run(caption)
        set_run_font(run, size=9.5, color=COLOR_DARK_BLUE, bold=True)

        tip = self.document.add_paragraph()
        tip.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(tip, after=8, line_spacing=1.15)
        lead = tip.add_run("读图提示：")
        set_run_font(lead, size=9, color=COLOR_TEAL, bold=True)
        run = tip.add_run(reading_tip)
        set_run_font(run, size=9, color=COLOR_MUTED)


def add_cover(document: Document) -> None:
    """采用 editorial_cover 形式添加封面。"""
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(84)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(kicker, after=16)
    run = kicker.add_run("技术原理与仿真入门 · 详细通俗版")
    set_run_font(run, size=11, color=COLOR_GOLD, bold=True)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(title, after=10, line_spacing=1.1)
    run = title.add_run("水下目标静电场与轴频电场\n仿真原理说明")
    set_run_font(run, size=26, color=COLOR_NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(subtitle, after=28)
    run = subtitle.add_run("静电场形成、轴频调制、空间传播与时序信号分析")
    set_run_font(run, size=14, color=COLOR_BLUE)

    audience = document.add_paragraph()
    audience.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(audience, after=72)
    run = audience.add_run("面向非专业读者 · 公式逐项解释 · 含图示、算例和结果判读")
    set_run_font(run, size=10.5, color=COLOR_MUTED)

    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(date, after=3)
    run = date.add_run("2026年8月")
    set_run_font(run, size=11, color=COLOR_NAVY, bold=True)

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(note, after=0)
    run = note.add_run("本文只讲物理原理和仿真思路，不涉及程序代码")
    set_run_font(run, size=9.5, color=COLOR_MUTED, italic=True)
    document.add_page_break()


def build_document(assets_dir: Path, output_path: Path) -> None:
    """构建优化后的完整 Word 文档。"""
    document = Document()
    bullet_id, number_id = configure_document(document)
    writer = GuideWriter(document, bullet_id, number_id)
    add_cover(document)

    writer.heading1("全文核心内容概览")
    writer.callout(
        "一句话理解",
        "舰船在海水中可以看成一节非常微弱、结构很复杂的“电池”：腐蚀和防腐系统让电流在船体与海水之间形成闭合回路，这部分构成缓慢变化的静电场；轴转动又让回路电阻周期变化，于是在静电场上叠加有固定节奏的轴频电场。",
        "key",
    )
    writer.body("整套仿真其实只回答四个连续问题：")
    writer.number("源从哪里来？——船体、螺旋桨、涂层破损区和阴极保护装置之间存在电化学电位差。")
    writer.number("电场怎样传到远处？——电流通过能够导电的海水扩散，可用点源或电流偶极子近似。")
    writer.number("为什么会变成一段波形？——轴旋转造成周期调制，目标运动又不断改变距离和方向。")
    writer.number("传感器最后输出什么？——某个方向的带符号电场分量，或三个分量合成后的电场大小。")
    writer.table(
        ["信号", "变化速度", "主要来源", "看图时抓住什么"],
        [
            ["静电场", "慢", "腐蚀与阴极保护电流", "目标靠近时整体增强、远离时减弱"],
            ["轴频电场", "快且周期性", "轴系接触电阻和电流通路随转动变化", "波形周期由轴转速决定"],
            ["总电场", "慢包络上叠加快振荡", "静电场矢量与轴频电场矢量相加", "既有通过过程，也有轴频细纹"],
        ],
        [1500, 1300, 3000, 3560],
    )
    writer.figure(
        assets_dir / "image1.png",
        "图1  舰船水下静电场与轴频电场的产生机理",
        "先看绿色电流通路，它代表较稳定的腐蚀/防腐电流；再看橙色轴系，它给这条通路增加周期性变化。传感器同时接收到二者叠加后的结果。",
    )

    writer.heading1("1. 水下电场传感器的测量对象与输出量")
    writer.heading2("1.1 电位、电压和电场不是同一个量")
    writer.body("电位可以理解为某一点所处的“电压水平”；电压是两个位置电位的差；电场则表示电位在空间中变化得有多快、朝哪个方向变化。")
    writer.body("可以把电位想成地形图上的“海拔”，把电压想成两地的“高度差”，把电场想成坡度和下坡方向。只知道一个地点有多高，不能判断周围坡有多陡；同样，只知道一个点的电位，也不能得到电场，必须比较相邻位置的电位差。")
    writer.body("实际水下电场传感器通常用相隔一段距离的电极测量微小电压差，再除以电极间距，换算成某个方向上的电场分量。若两电极间距矢量为 L，在电场变化不太剧烈时，可以近似写成：")
    writer.equation("ΔV ≈ −E · L", "ΔV 是两电极电压差；E 是电场矢量；L 同时包含电极间距和测量方向。负号表示电位沿电场方向降低。")
    writer.body("例如，两只电极沿 x 方向相隔 1 m，测得电压差的绝对值为 2 μV，那么该方向的平均电场分量约为 2 μV/m。若电极间距改为 0.5 m，而电压差仍为 2 μV，则换算结果约为 4 μV/m。这个例子说明，前端若接收的是电压差，还必须同时知道电极间距和朝向，才能还原电场分量。")
    writer.callout("直观例子", "两个电极东西向摆放，主要测到 Eₓ；南北向摆放，主要测到 Eᵧ；竖直摆放，主要测到 E_z。三组正交电极可以得到三分量电场。", "info")
    writer.body("仿真常把结果写成 μV/m（微伏每米）。1 μV/m 等于 10⁻⁶ V/m，说明目标电场通常非常微弱。")

    writer.heading2("1.2 这里的“静电场”不是摩擦起电")
    writer.body("日常生活中的静电常指衣物摩擦后电荷积累。舰船水下“静电场”不是这个意思。海水能导电，船体相关电流会持续或缓慢变化，因此更准确的名称是“腐蚀相关准静态电场”。行业中为了简便，通常仍称它为静电场。")
    writer.body("所谓“准静态”，是说在我们关心的几秒到几分钟时间尺度上，源强和空间分布变化得比较慢，可以把每一个采样时刻都近似看成一个已经稳定下来的导电问题。目标仍在运动、电流仍在流动，只是不用像高频无线电那样重点考虑电磁波传播和辐射效应。")
    writer.callout("容易混淆", "“静态”描述的是变化较慢，不代表电流为零；恰恰是持续流过海水的微弱电流形成了这个场。", "caution")

    writer.heading2("1.3 水下目标电场仿真的三层结构")
    writer.table(
        ["层次", "要回答的问题", "典型参数"],
        [
            ["源", "目标为什么产生电流，源有多强、朝哪边", "偶极矩、腐蚀电流密度、涂层破损率"],
            ["介质", "电流怎样在海水中扩散", "电导率、海水分层、海面和海床边界"],
            ["几何与时间", "目标和传感器怎样相对运动", "位置、航速、航向、传感器坐标、采样率"],
        ],
        [1500, 4300, 3560],
    )

    writer.heading1("2. 舰船水下静电场的形成机理")
    writer.heading2("2.1 海水导电的基本原理")
    writer.body("海水中有大量可以移动的正、负离子。电场一出现，离子就会定向移动，形成电流。低频条件下最常用的关系是连续介质欧姆定律：")
    writer.body("在金属导线里，电流主要由电子移动形成；在海水里，主要是带正电和带负电的离子向相反方向移动。盐溶解后产生的离子越多，海水通常越容易导电。这里的“导电”不是说海水中有一根看得见的电线，而是整个水体都可以为微弱电流提供分散的通路。")
    writer.equation("J = σE", "J：电流密度（A/m²）；σ：海水电导率（S/m）；E：电场强度（V/m）。")
    writer.body("这个公式可以按“流量 = 通行能力 × 推动力”来理解：J 表示单位面积通过多少电流，σ 表示海水有多容易让电流通过，E 表示推动离子移动的电场。若要求同样大小的 J，σ 越大，所需的 E 越小；反过来，在同样的 E 下，σ 越大，能够形成的 J 越大。")
    writer.body("若把电流源强暂时看成固定，电导率越高，同样电流越容易通过，维持它所需的电场反而越小。但真实腐蚀电流也会随盐度、温度、流速和氧含量改变，所以不能简单说“盐度越高，电场一定越大”或“一定越小”。")
    writer.body("海水电导率由盐度、温度和压力共同决定。在实用盐度 35、15 ℃、海压 0 dbar 附近，参考电导率约为 4.2914 S/m [1]。有现场实测电导率时，应优先使用实测值。")

    writer.heading2("2.2 舰船水下电化学电流回路")
    writer.body("船体钢材、铜合金螺旋桨、轴系、牺牲阳极和外加电流阴极保护装置的材料与电化学电位不同。它们同时浸在海水中，就可能形成类似电池的电位差。")
    writer.body("“像电池”并不是说船上真的装了一节专门向海水放电的电池，而是不同金属、不同表面状态和保护装置之间天然存在电化学电位差。这个电位差提供了推动电荷运动的条件。涂层相当于绝缘保护层；一旦局部破损，裸露金属与海水接触，电化学反应和电流交换就会明显增强。")
    writer.body("电流不会凭空消失：它从某些表面进入海水，在海水中扩散，再从另一些表面返回目标。涂层完好的区域电流较小，涂层破损、裸露金属、螺旋桨和保护装置附近往往更重要。")
    writer.body("把完整回路拆开看，可以分成四步：第一，目标不同部位形成电位差；第二，电荷在船体金属内部通过电子传递；第三，电流在金属—海水界面通过电化学反应进入或离开海水；第四，海水中的正、负离子定向移动，把分散的流出区和回流区连接起来。传感器测到的正是第四步在周围水体中建立的电位差和电场。")
    writer.callout("先记住闭环", "有“流出端”就一定有“回流端”。远处传感器看到的不是孤立电流源，而是源和汇共同形成的空间电场。", "key")

    writer.heading2("2.3 静态电场源强的主要影响因素")
    writer.body("仿真中不能只用“吨位大或小”决定电场强弱。吨位和外形主要提供几何尺度；真正决定源强的，是有多少电流穿过金属—海水界面、这些流出区与回流区相隔多远，以及它们在空间中朝什么方向分布。")
    writer.table(
        ["因素", "变化会影响什么", "通俗理解"],
        [
            ["涂层破损率", "有效裸露面积", "破损越多，可参与电化学反应的面积通常越大"],
            ["腐蚀电流密度", "单位面积流出的电流", "综合反映材料、氧含量、流速和保护状态"],
            ["阴极保护系统", "电流大小与分布", "保护腐蚀的同时，也会改变水下电场"],
            ["目标尺寸", "浸水面积和源—汇尺度", "大目标可能有更大作用面积，但绝不是只按吨位简单放大"],
            ["海水环境", "源反应和传播条件", "盐度、温度、压力和流速都会产生影响"],
        ],
        [1800, 3000, 4560],
    )

    writer.heading1("3. 静电场的等效电流偶极子建模方法")
    writer.heading2("3.1 点电流源模型与空间扩散")
    writer.body("想象海水中有一个很小的电极，持续向外送出电流 I。若海水无限、均匀，电流会向四周扩散，经过的球面面积越来越大，所以越远越稀。该点的电位为：")
    writer.body("为什么会“越远越稀”？距离点源 r 的球面面积是 4πr²，而穿过每个球面的总电流仍然是 I。于是平均电流密度约为 I/(4πr²)：距离变成 2 倍，球面面积变成 4 倍，单位面积分到的电流只剩 1/4。再由 J=σE 可知，点源电场也按 1/r² 衰减。")
    writer.equation("φ(r) = I / (4πσr)", "φ：电位；I：电流；σ：电导率；r：到点源的距离。点源电位按 1/r 衰减。")
    writer.body("电位为什么按 1/r 而不是 1/r² 衰减？因为电场是电位随距离的变化率。把按 1/r² 变化的电场从远处沿距离累加回来，得到的电位正比于 1/r。也可以反过来理解：对 1/r 的电位求空间变化率，会得到 1/r² 的电场。")

    writer.heading2("3.2 远场电流偶极子等效")
    writer.body("真实目标既有电流流出，也有电流流回。把一正一负、大小相同、距离较近的点源放在一起，就是电流偶极子。距离足够远时，传感器分不清船体上每个细小电流区，只能看到它们合成后的主要方向和强度。")
    writer.body("偶极子的关键是“两个相反点源发生部分抵消”。单独看流出端和回流端，它们各自产生约按 1/r² 衰减的电场；在远处，两者距离相近、大小相反，大部分共同成分互相抵消，只留下由源—汇间距造成的微小差别。因此偶极子远场比单点源衰减得更快，最终表现为约 1/r³。")
    writer.equation("p = I · d", "p：电流偶极矩（A·m）；I：等效电流；d：等效源—汇间距。p 的方向从等效回流端指向等效流出端。")
    writer.body("偶极矩 p 把三个信息压缩到一个矢量里：它的大小同时包含等效电流 I 和源—汇间距 d，它的方向表示主要电流分布方向。比如 2 A 的等效电流跨越 5 m，与 1 A 的等效电流跨越 10 m，在这个简化模型中都对应 10 A·m 的偶极矩；但真实近场分布仍可能不同。")
    writer.figure(
        assets_dir / "image2.png",
        "图2  等效电流偶极子的场线与相对强弱分布",
        "同一距离但方向不同，电场也可能不同；这就是为什么传感器在目标前方、侧面和下方会得到不同结果。",
        width=6.2,
    )

    writer.heading2("3.3 偶极子电场公式及参数含义")
    writer.body("下面两条公式不是凭空设定的。第一条电位公式来自“正点源电位减去负点源电位”的远距离近似；第二条电场公式来自 E=−∇φ，也就是对电位在 x、y、z 三个方向的变化求导。阅读时不必先掌握完整微积分，可以按“源强、方向、距离、介质”四部分拆开。")
    writer.equation("φ(r) = (p · r) / (4πσr³)", "点积 p·r 把“源方向”和“观测方向”的关系带进了电位公式。")
    writer.equation("E(r) = [3r̂(p · r̂) − p] / (4πσr³)", "r̂ 是目标指向传感器的单位方向。方括号决定方向，分母决定总体衰减速度。")
    writer.number("先确定相对位置 r：从等效源中心指向传感器。它的长度 r 是距离，单位方向 r̂ 只表示朝向。")
    writer.number("再确定偶极矩 p：目标坐标系中的偶极方向要根据航向、俯仰等姿态转换到全局坐标系。")
    writer.number("计算 p·r̂：该点积等于 |p|cosθ，θ 是偶极方向与观测方向的夹角，用来衡量传感器位于目标前后还是侧面。")
    writer.number("计算方括号中的矢量：它决定 Eₓ、Eᵧ、E_z 的方向和正负。")
    writer.number("最后除以 4πσr³：这一步体现海水导电能力和距离衰减。")
    writer.table(
        ["公式部分", "代表什么", "它变大时通常怎样"],
        [
            ["p", "等效偶极矩", "若其他条件不变，电场近似同比增大"],
            ["r³", "距离的三次方", "距离增大，场强快速变小"],
            ["σ", "海水电导率", "若偶极矩固定，电场近似按 1/σ 变小"],
            ["p·r̂", "源方向与观测方向的投影", "决定前方、侧面、下方的方向差异"],
            ["4π", "三维空间扩散的几何常数", "是球面对称扩散留下的系数"],
        ],
        [1800, 3300, 4260],
    )
    writer.callout("最重要的数量级", "理想偶极远场大约按 1/r³ 衰减。距离变为 2 倍，场强约变为 1/8；距离变为 3 倍，约变为 1/27。方向因子会改变具体数值，但不会改变“远处快速变弱”的总体规律。", "key")
    writer.figure(
        assets_dir / "image3.png",
        "图3  理想点源与理想偶极源的距离衰减对比",
        "蓝色偶极曲线比橙色点源曲线下降得更快。纵轴是对数坐标，每下降一格代表再缩小 10 倍。",
        width=6.3,
    )

    writer.heading2("3.4 电场分量的方向与正负号")
    writer.body("电场是三维矢量，通常写成 E = (Eₓ, Eᵧ, E_z)。每个分量都可以为正或负：正负只表示相对于坐标轴的方向，不表示“场强有问题”。")
    writer.body("例如规定 x 轴指向东，目标从传感器西侧驶向东侧。目标在西侧时，某个电场分量可能指向东，记为正；通过传感器附近后，该分量可能改为指向西，记为负。曲线穿过零点往往意味着方向发生转换，不等同于总电场在该时刻完全消失，因为另外两个分量可能仍然存在。")
    writer.body("电场大小不带方向，计算为：")
    writer.equation("|E| = √(Eₓ² + Eᵧ² + E_z²)", "分量可正可负；模值一定大于或等于 0。两者用途不同。")

    writer.heading2("3.5 缺少实测数据时的静态源强估算")
    writer.body("最简估算可以分两步：先估算参与反应的总电流，再乘以等效源—汇间距。")
    writer.equation("I_est ≈ j_corr · S_wet · η_damage", "j_corr：等效腐蚀电流密度；S_wet：浸水面积；η_damage：涂层等效破损比例。")
    writer.equation("p_static ≈ I_est · d_eff", "d_eff：等效源—汇间距。该结果是用于仿真的等效参数，不是直接测得的船体真实分布。")
    writer.body("计算顺序可以这样理解：先用船长、船宽和吃水粗略估算浸水面积 S_wet；再用破损比例 η_damage 估计真正裸露或有效参与反应的面积；乘以等效腐蚀电流密度 j_corr 得到总电流数量级；最后乘以等效源—汇间距 d_eff 得到偶极矩。每一步都有不确定性，所以最终结果更适合给出数量级或参数范围，而不是声称精确到某个小数。")
    writer.body("吨位本身没有直接出现在偶极场公式中。它通常只用于帮助判断目标尺度，并间接影响湿表面积、源—汇间距和可能的保护系统规模。两艘吨位相近的舰船，如果涂层状态、材料配置或阴极保护电流不同，静态电场仍可能相差很大。")
    writer.callout("工程提醒", "吨位、长度、宽度和吃水只能帮助估算浸水面积与源间距，不能唯一决定真实偶极矩。需要定量预测时，应使用实测电位/电场反演或高保真电化学模型标定 p_static。", "caution")

    writer.heading2("3.6 偶极子模型的适用条件与升级方法")
    writer.bullet("传感器非常靠近船体，已能分辨不同电流区域。")
    writer.bullet("浅海条件下，海面和海床边界明显改变电流路径。")
    writer.bullet("海水电导率随深度或水平位置变化明显。")
    writer.bullet("需要研究涂层局部破损、牺牲阳极或阴极保护电极的细节。")
    writer.body("这些情况下可升级为多偶极子、点源阵列、边界元或有限元模型 [2][3][7]。")

    writer.heading1("4. 轴频电场的产生机理与周期特征")
    writer.heading2("4.1 轴系周期电阻调制模型")
    writer.body("船体腐蚀和防腐系统先建立了一个基础电流回路。轴、轴承、接地刷和滑环等部位的接触状态会随转角变化，使回路的等效电阻呈周期变化。电阻一变，回路电流也跟着变。")
    writer.body("可以把它想成一个旋转的、接触不完全均匀的旋钮：轴每转一圈，某些接触位置会重复出现，回路电阻也大致重复一次。若轴表面、轴承接触或接地装置并不完全对称，那么不同转角对应的电阻就不相同。这个随转角重复的电阻变化，就是轴频电场的“节拍来源”。")
    writer.equation("R(t) = R₀[1 + m cos(2πf_st + φ)]", "R₀：平均电阻；m：变化比例；f_s：轴频；φ：初相位。")
    writer.body("如果驱动电位差在一个转动周期内近似不变，I(t)=U/R(t)。当 m 较小时，可以把电流近似理解为“平均值上叠加一个同频率的小波动”。所以轴频场不是凭空多出来的场，而是原有电流回路被轴转动调制后的结果 [4][5]。")
    writer.body("因果链条可以逐步写成：轴转角变化 → 接触状态变化 → 等效电阻 R(t) 变化 → 回路电流 I(t) 变化 → 等效偶极矩 p(t) 变化 → 传感器处电场随时间周期变化。在电压近似不变时，电阻稍微增大，电流会稍微减小；电阻稍微减小，电流会稍微增大，因此电流波动与电阻波动大体反向。")
    writer.callout("静电场与轴频电场的关系", "轴频电场通常不是一套与静电场完全独立的源，而是轴系把原有腐蚀/防腐电流回路做了周期调制。仿真时可把它拆成“平均静态部分 + 周期变化部分”，便于分别观察和标定。", "key")

    writer.heading2("4.2 轴转速与轴频的换算")
    writer.equation("f_s = n / 60", "n 用 r/min 表示，除以 60 后得到每秒转数，也就是 Hz。")
    writer.body("例如轴转速为 120 r/min：")
    writer.number("120 ÷ 60 = 2，因此基波频率为 2 Hz。")
    writer.number("一个周期为 1 ÷ 2 = 0.5 s，也就是每半秒重复一次。")
    writer.number("如果信号包含二次和三次谐波，还会在 4 Hz 和 6 Hz 出现成分。")
    writer.callout("轴频不等于叶频", "轴频是轴每转一圈重复一次。若螺旋桨有 Z 个叶片，叶片通过频率通常为 f_b = Zf_s。两者可能同时出现在实测信号中，但物理含义不同。", "caution")

    writer.heading2("4.3 轴频谐波的形成原因")
    writer.body("如果接触电阻随转角变化是一条完美正弦曲线，频谱主要只有基波。现实中的接触不均匀、偏心、多处接触和材料差异会让波形出现尖、扁、偏斜等形状。任何稳定重复但不是纯正弦的波形，都可以拆成基波和若干整数倍频率，这些整数倍成分就是谐波。")
    writer.body("傅里叶分解可以用“用多组规则正弦波拼出一条复杂周期曲线”来理解。基波负责最基本的一圈一次变化；二次谐波每圈重复两次，三次谐波每圈重复三次。谐波不是凭空产生了新的转轴，而是用不同倍频的正弦成分描述同一根轴的不规则周期波形。")
    writer.equation("p_shaft(t) = p₁cos(2πf_st+φ₁) + p₂cos(4πf_st+φ₂) + p₃cos(6πf_st+φ₃) + …", "p₁、p₂、p₃ 分别表示基波、二次谐波和三次谐波的等效偶极矩。")
    writer.figure(
        assets_dir / "image4.png",
        "图4  2 Hz 轴频电场的时域波形与频谱示例",
        "上图每 0.5 s 重复一次；下图最高峰位于 2 Hz，4 Hz 和 6 Hz 的小峰说明波形不是完美正弦。",
        width=6.25,
    )

    writer.heading2("4.4 轴频振荡与距离包络的叠加")
    writer.body("目标不动时，轴频场看起来像稳定的周期波。目标经过固定传感器时，传播距离不断变化：远处振幅很小，靠近时振幅变大，远离后再次变小。于是时域曲线呈现“慢慢鼓起的包络中装着快速振荡”。")
    writer.body("在简化表达中，可以把轴频分量理解为“距离和方向决定的慢变系数”乘以“由轴转动决定的快变周期项”。前者在几秒到几十秒内变化，后者可能每秒重复多次。频谱主要从快变项识别轴频，目标通过时的整体强弱则主要由慢变包络控制。")
    writer.callout("两条时间尺度", "慢变化来自目标通过传感器的几何过程；快变化来自轴转动。把二者分开理解，复杂波形就容易读懂。", "key")

    writer.heading2("4.5 多轴目标的频率、幅值与相位")
    writer.body("多轴舰船可能同时存在多个轴频源。若各轴转速不同，频谱中可能出现多组峰；若转速相同但相位不同，某些方向上的电场可能相加，也可能部分抵消。单轴模型适合入门和基线仿真，多轴定量分析需要分别描述每根轴的幅值、方向、转速和相位 [5]。")
    writer.body("相位表示周期波在起始时刻走到哪一步。两根轴的频率和方向相同、相位也相同时，峰和峰容易叠加；若相差约半个周期，一根轴处于正峰时另一根可能处于负峰，合成信号就会减弱。若两轴转速略有不同，合成幅度还可能出现缓慢起伏，形成类似“时强时弱”的拍频现象。")

    writer.heading1("5. 目标运动与传感器时序信号的关系")
    writer.heading2("5.1 目标位置与相对位置的逐时刻更新")
    writer.body("空间模型计算的是“某一个位置上的电场”，而传感器记录的是“电场随时间变化的曲线”。连接两者的办法，是在每个采样时刻重新计算目标位置，再把新的相对距离和方向代入同一个电场公式。这样，一串空间计算结果就按时间顺序组成了时序信号。")
    writer.body("假设目标做匀速直线运动，初始中心位置为 r_T0，航速为 v，运动方向单位矢量为 v̂，则：")
    writer.equation("r_T(t) = r_T0 + v̂vt", "初始位置加上“方向 × 速度 × 时间”，得到当前目标位置。")
    writer.body("传感器位置 r_S 固定。真正进入电场公式的是目标指向传感器的相对矢量：")
    writer.equation("R(t) = r_S − r_T(t)", "距离为 |R(t)|，方向为 R(t)/|R(t)|。减法顺序会影响方向和分量正负号。")
    writer.body("使用三维坐标前必须先统一约定，例如 x 向东、y 向北、z 向上，水下深度用负值表示。目标位置、传感器位置、航向和偶极方向都必须服从同一套约定。坐标方向若混用，距离大小可能仍然正确，但 Eₓ、Eᵧ、E_z 的方向和正负会出错。")

    writer.heading2("5.2 最近通过距离与信号包络")
    writer.body("在最简单的平直航迹中，把最近点设为 t=0，目标沿航迹方向运动，传感器到航迹的最短距离为 b，则目标到传感器的距离是：")
    writer.equation("R(t) = √[b² + (vt)²]", "b 是最近通过距离；v 是航速。t=0 时距离最小，为 b。")
    writer.body("若只看 1/R³ 的距离包络，可近似写成：")
    writer.equation("A(t) ∝ 1 / [b² + (vt)²]³ᐟ²", "目标靠近时包络升高，过最近点后对称下降；真实分量还会受方向因子影响。")
    writer.body("航速提高通常会让同一空间分布被更快扫过，因此曲线在时间轴上变窄；在航迹和源强完全相同的理想模型中，最近点的空间场值并不会仅因航速提高而自动变大。")
    writer.body("这里的 b 是三维最短距离，不一定只是海图上的水平距离。若传感器与目标航迹在水平方向相隔 100 m、深度方向相隔 25 m，则 b 应由这两个正交距离共同合成。忽略深度差会把最近距离算小，并因 1/r³ 衰减而明显高估峰值。")
    writer.figure(
        assets_dir / "image5.png",
        "图5  目标通过固定传感器时的几何关系与时序包络",
        "右图蓝线是慢变化包络，橙色是轴频振荡。中间竖线对应最近点；分量曲线还可能在附近发生正负翻转。",
        width=6.35,
    )

    writer.heading2("5.3 目标姿态对三分量电场的影响")
    writer.body("偶极矩有方向。航向和俯仰改变后，同一个偶极矩在全局 x、y、z 方向上的投影会改变。因此，目标即使距离传感器相同，朝向不同也可能产生不同的三分量结果。")
    writer.body("通常先在目标自身坐标系中定义偶极矩，例如纵向偶极沿船艏方向、横向偶极沿右舷方向、垂向偶极沿上方。仿真时再根据航向、俯仰和横滚把它旋转到全局坐标系。这个旋转改变的是方向投影，不会凭空改变偶极矩自身的大小。")

    writer.heading2("5.4 静电场与轴频电场必须按矢量相加")
    writer.equation("E_total(t) = E_static(t) + E_shaft(t)", "这是三个方向逐分量相加，不是把两个非负模值直接相加。")
    writer.body("正确做法是分别计算 Eₓ、Eᵧ、E_z，再合成总矢量。若只把 |E_static| 和 |E_shaft| 相加，会丢失方向和抵消关系。")
    writer.body("例如静态场为 (3, 0, 0) μV/m，某时刻轴频场为 (−1, 2, 0) μV/m，那么总场是 (2, 2, 0) μV/m，模值约为 2.83 μV/m；若错误地把两个模值直接相加，会得到 3+√5≈5.24 μV/m，明显高估。")

    writer.heading2("5.5 频谱分析中的电场分量选择")
    writer.body("|E| 总是非负，相当于把负半周期翻到正方向，可能人为产生 2 倍频等成分。因此做 FFT 时，应优先使用带正负号的 Eₓ、Eᵧ 或 E_z；模值更适合展示总体强弱。")
    writer.body("最简单的例子是纯余弦信号：它每个周期有一次正峰和一次负峰；取绝对值后，负峰被翻成正峰，于是每个原周期出现两个峰，看起来像频率翻倍。这个倍频来自数学处理，不一定对应真实物理谐波。")
    writer.callout("实用规则", "看总体强弱用 |E|；看方向变化和做频谱用 Eₓ、Eᵧ、E_z。", "key")

    writer.heading1("6. 水下目标电场仿真的完整流程")
    writer.heading2("6.1 六类输入参数及其作用")
    writer.body("输入参数不是简单地全部代入同一条公式。它们分别负责定义源强、传播介质、空间几何、周期变化和数字采样。把参数按作用分类，可以避免把“影响频率的参数”和“影响幅值的参数”混在一起。")
    writer.table(
        ["参数组", "典型内容", "在模型中的作用"],
        [
            ["目标几何", "类型、吨位、长度、宽度、吃水", "估算浸水面积与等效源尺度"],
            ["运动姿态", "初始位置、航速、航向、俯仰", "更新目标位置和偶极方向"],
            ["轴系状态", "轴转速、初相位、谐波比例", "确定周期波形与频谱峰"],
            ["海水环境", "盐度、温度、压力或实测电导率", "确定导电介质参数"],
            ["等效源", "静态偶极矩、腐蚀电流密度、破损率", "决定绝对场强尺度"],
            ["观测设置", "传感器坐标、时长、采样率", "决定时序范围和可分析频率"],
        ],
        [1500, 3700, 4160],
    )
    writer.callout("最值得标定的参数", "目标尺寸是已知外形，真正控制绝对场强的往往是静态偶极矩、轴频调制比例和谐波比例。没有实测标定时，输出更适合做趋势和算法测试。", "caution")

    writer.heading2("6.2 单个采样时刻的七步计算")
    writer.body("假设采样率为 20 Hz，相邻两个采样时刻相差 0.05 s。仿真会从 t=0 开始，重复执行下面七步，直到达到设定时长。每次循环只负责计算当前时刻的一组结果。")
    writer.number("确定海水电导率：优先使用实测值，否则根据盐度、温度和压力估算。")
    writer.number("确定静态等效偶极矩：使用标定值，或根据腐蚀参数做初步估算。")
    writer.number("根据轴转速、初相位和谐波比例，计算当前轴频偶极矩。")
    writer.number("更新目标位置和姿态，计算目标到传感器的相对矢量。")
    writer.number("分别计算静态场和轴频场的 x、y、z 分量。")
    writer.number("逐分量相加，得到综合电场，并计算需要展示的模值。")
    writer.number("保存时间、位置、距离、频率和电场结果，进入下一采样时刻。")
    writer.body("循环结束后得到的核心结果应至少包括：时间 t，静态场三分量，轴频场三分量，总场三分量和总场模值。为了前端绘图、问题定位和结果复核，最好同时保存目标位置、目标—传感器距离、轴相位与轴频基波频率。")
    writer.figure(
        assets_dir / "image6.png",
        "图6  水下目标电场仿真的完整流程",
        "流程中最容易忽略的是“更新相对位置”和“矢量相加”；这两步决定了时序包络、方向与正负号。",
        width=6.35,
    )

    writer.heading2("6.3 采样率选择原则")
    writer.equation("f_sample > 2f_max", "奈奎斯特最低条件：采样率必须大于最高关注频率的两倍。")
    writer.body("如果轴频为 2 Hz，并希望观察到三次谐波 6 Hz，那么采样率必须大于 12 Hz。实际为了让波形更平滑、频谱更稳定，可选 20～50 Hz。")
    writer.body("采样率只决定数字记录是否足够细，不会改变真实物理轴频。采样率太低时，高频会折叠成错误的低频，这叫混叠。")

    writer.heading2("6.4 仿真时长选择原则")
    writer.body("仿真时长应覆盖目标从“足够远”靠近传感器、经过最近点、再离开的全过程。时长越长，频谱频率间隔越细；但数据量和计算量也会增加。")
    writer.equation("频率分辨率 Δf ≈ 1 / T", "T 是记录时长。例如记录 75 s，理想频率间隔约为 0.0133 Hz。")

    writer.heading1("7. 从输入到结果的完整算例")
    writer.body("下面用一组假设参数走完整个思考过程。数值只用于说明方法，不代表某艘真实舰船。")
    writer.table(
        ["参数", "示例值", "解释"],
        [
            ["目标尺寸", "5000 t；100 m × 15 m × 5 m", "用于估算浸水面积和等效尺度"],
            ["初始位置", "(−300, 0, −5) m", "仿真开始时目标在传感器西侧"],
            ["航速与航向", "8 m/s；沿 x 轴正方向", "目标做匀速直线运动"],
            ["轴转速", "120 r/min", "轴频基波为 2 Hz"],
            ["海水环境", "盐度 35，18 ℃，30 dbar", "用于估算海水电导率"],
            ["传感器位置", "(0, 100, −30) m", "距航迹水平 100 m，深度差 25 m"],
            ["时长与采样率", "75 s；20 Hz", "覆盖经过过程并观察到 6 Hz 三次谐波"],
        ],
        [1800, 3100, 4460],
    )

    writer.heading2("7.1 轴频、谐波与采样点计算")
    writer.number("轴频：120 ÷ 60 = 2 Hz。")
    writer.number("二次、三次谐波：4 Hz 和 6 Hz。")
    writer.number("采样率 20 Hz 的奈奎斯特频率为 10 Hz，高于 6 Hz，因此可保留前三次谐波。")
    writer.number("采样点数：75 × 20 + 1 = 1501 点，包含开始时刻。")

    writer.heading2("7.2 最近通过时刻与距离计算")
    writer.body("目标初始 x = −300 m，以 8 m/s 向 x 正方向运动。到达 x = 0 需要：")
    writer.equation("t_CPA = 300 / 8 = 37.5 s", "CPA 表示最近通过点。75 s 仿真恰好把最近点放在记录中间。")
    writer.body("最近点处，水平横向距离为 100 m，垂向差为 25 m，因此最近距离约为：")
    writer.equation("b = √(100² + 25²) ≈ 103.1 m", "仿真开始时距离约为 √(300²+100²+25²) ≈ 317.2 m。")

    writer.heading2("7.3 距离包络变化估算")
    writer.body("只看 1/r³ 距离因子，不考虑方向变化，最近点与开始时刻的理想场强比约为：")
    writer.equation("(317.2 / 103.1)³ ≈ 29", "这不是最终精确倍数，因为偶极方向因子也会随位置变化，但它能说明为什么靠近时信号会明显鼓起。")

    writer.heading2("7.4 仿真曲线的预期特征")
    writer.bullet("静电场分量形成一个随靠近而增强、随远离而减弱的慢变化过程。")
    writer.bullet("轴频分量以 2 Hz 为主，在慢变化包络内快速振荡，并可能包含 4 Hz、6 Hz 谐波。")
    writer.bullet("x、y、z 分量形状不一定相同，某些分量可能在最近点附近穿过零点。")
    writer.bullet("总电场模值通常在最近点附近较大，但峰值不一定精确落在最近点，因为方向因子和静、轴频相位也在变化。")

    writer.heading1("8. 仿真参数对计算结果的影响")
    writer.table(
        ["参数变化", "最直接的结果", "容易误解的地方"],
        [
            ["静态偶极矩增大", "静电场近似同比增大", "它比吨位更直接控制绝对场强"],
            ["距离增大", "理想偶极远场约按 1/r³ 迅速减弱", "方向因子仍会造成同距离不同场强"],
            ["电导率增大", "在偶极矩固定时，场约按 1/σ 减小", "真实腐蚀电流也可能随环境改变"],
            ["航速增大", "通过过程在时间轴上变窄", "理想相同航迹的空间峰值不一定增加"],
            ["轴转速增大", "轴频峰向更高频率移动", "转速本身不等于振幅一定增大"],
            ["轴频调制比例增大", "轴频振幅增大", "不会改变基波频率"],
            ["谐波比例增大", "波形更不像正弦，倍频峰更明显", "二次谐波不是第二根轴"],
            ["采样率增大", "记录更细、混叠风险下降", "不会改变真实物理场强"],
            ["吨位或尺寸增大", "可能增大估算的湿面积和源尺度", "不能只凭吨位推断真实场强"],
        ],
        [2300, 3300, 3760],
    )

    writer.heading1("9. 仿真结果的合理性检查与参数标定")
    writer.heading2("9.1 无实测数据时的五项基本检查")
    writer.number("单位检查：位置用 m、速度用 m/s、转速用 r/min、电场输出用 μV/m。")
    writer.number("频率检查：120 r/min 必须对应 2 Hz，不能直接把 120 当作 Hz。")
    writer.number("距离检查：目标靠近时总体增强，远离时总体减弱；极近距离应触发模型适用性警告。")
    writer.number("符号检查：分量可以为负，模值不能为负；FFT 不应优先使用模值。")
    writer.number("缩放检查：把距离加倍后，远场结果应大致减少到原来的 1/8 数量级。")

    writer.heading2("9.2 参数敏感性检查")
    writer.body("每次只改变一个参数，例如把偶极矩、距离或调制比例变为两倍，观察结果是否符合预期关系。敏感性检查比只看一条曲线更容易发现公式、单位和坐标错误。")

    writer.heading2("9.3 基于实测数据的参数标定")
    writer.number("先用远离目标的时段估计背景和噪声。")
    writer.number("根据航迹和传感器位置确定每个时刻的相对几何。")
    writer.number("用低频或平滑包络反演静态等效偶极矩及方向。")
    writer.number("在分量信号频谱中识别轴频及谐波，拟合调制比例和相位。")
    writer.number("用另一段数据验证参数，避免只对单次记录拟合得很好。")

    writer.heading1("10. 模型适用范围与局限性")
    writer.table(
        ["适合使用", "不能直接宣称"],
        [
            ["生成算法验证和前端联调用的合成时序", "仅凭吨位和尺寸精确预测某艘实船绝对场强"],
            ["研究距离、方向、速度、轴频和采样率的影响", "替代船模试验、海试或高保真电化学仿真"],
            ["验证 FFT、特征提取和目标通过检测流程", "精确描述浅海边界、海床和分层海水传播"],
            ["作为多偶极子、边界元和有限元模型的基线", "在传感器紧贴船体时仍保持远场偶极精度"],
        ],
        [4680, 4680],
    )
    writer.callout("升级路线", "趋势分析用单偶极子；需要描述不同船体区域时用多偶极子或点源阵列；需要处理真实船体、海面、海床和局部电极时，再使用边界元或有限元。", "info")

    writer.heading1("11. 常见误区快速澄清")
    writer.table(
        ["常见说法", "正确理解"],
        [
            ["静电场就是船体带了很多静电", "这里主要是导电海水中的腐蚀/防腐电流形成的准静态场"],
            ["吨位越大，电场一定按比例越大", "吨位只间接影响几何估算，真实源强还取决于材料、涂层和保护系统"],
            ["轴频是螺旋桨叶片数乘转速", "那通常是叶频；轴频是轴每秒转数"],
            ["总场等于两个场强模值相加", "必须先按 x、y、z 分量做矢量相加"],
            ["曲线出现负值说明算错了", "分量的负值只表示方向反向"],
            ["采样率越高，物理场强越大", "采样率只影响数字记录，不改变真实物理场"],
            ["频谱直接对 |E| 做就可以", "模值整流会制造额外倍频，应优先分析带符号分量"],
        ],
        [3600, 5760],
    )

    writer.heading1("12. 常用术语小词典")
    writer.table(
        ["术语", "通俗解释"],
        [
            ["电位 φ", "某一点的电压水平；两个位置的电位差可以推动电流"],
            ["电场 E", "电位在空间中变化的快慢和方向，常用 V/m 或 μV/m"],
            ["电导率 σ", "海水让电流通过的能力，单位 S/m"],
            ["电流密度 J", "单位面积上通过的电流，单位 A/m²"],
            ["电流偶极矩 p", "用“等效电流 × 源汇间距”概括复杂电流分布的矢量"],
            ["静电场", "由持续或缓慢变化的腐蚀/防腐电流形成的准静态场"],
            ["轴频电场", "轴系周期变化调制原有电流回路后形成的周期电场"],
            ["基波", "与轴转一圈对应的最低主要频率"],
            ["谐波", "基波整数倍频率处的周期成分"],
            ["包络", "快速振荡外面那条缓慢变化的幅度轮廓"],
            ["最近通过点 CPA", "目标到传感器距离最小的位置或时刻"],
            ["时域 / 频域", "分别观察信号随时间变化，以及由哪些频率成分组成"],
        ],
        [2200, 7160],
    )

    writer.heading1("13. 核心知识自测")
    writer.number("为什么舰船水下静电场不是普通的摩擦静电？")
    writer.number("为什么远处可以用一个电流偶极子近似复杂船体？")
    writer.number("为什么理想偶极远场对距离特别敏感？")
    writer.number("120 r/min 的轴频是多少，二次和三次谐波又是多少？")
    writer.number("为什么目标通过传感器时会出现“慢包络 + 快振荡”？")
    writer.number("为什么做 FFT 时应使用 Eₓ、Eᵧ 或 E_z，而不是优先使用 |E|？")
    writer.callout("如果这六题都能用自己的话解释", "说明静电场、轴频电场、目标运动和信号处理之间的知识链条已经连起来了。", "key")

    writer.heading1("参考文献")
    references = [
        "McDougall, T. J., Barker, P. M. Notes on the GSW function gsw_C_from_SP: conductivity from Practical Salinity, in-situ temperature and pressure. TEOS-10/GSW Oceanographic Toolbox, 2011. https://www.teos-10.org/pubs/gsw/pdf/C_from_SP.pdf",
        "Sun, Q., Jiang, R.-X., Yu, P. Inversion of the Equivalent Electric Dipole Moment of Ship’s Corrosion-Related Static Electric Field in Frequency Domain. Mathematical Problems in Engineering, 2020, Article 3486082. https://doi.org/10.1155/2020/3486082",
        "Wang, X., Xu, Q., Zhang, J. Simulating Underwater Electric Field Signal of Ship Using the Boundary Element Method. Progress In Electromagnetics Research M, 2018, 76: 43-54. https://doi.org/10.2528/PIERM18092706",
        "Zhang, A., Zhang, H., Cheng, Z. Shaft-ground Active Compensation Technology of the Ship Shaft-rate Electric Field. Chinese Journal of Ship Research, 2016, 11(2): 121-126. https://doi.org/10.3969/j.issn.1673-3185.2016.02.017",
        "Wang, X., Wang, S., Hu, Y., Tong, Y. Mixed Electric Field of Multi-Shaft Ship Based on Oxygen Mass Transfer Process under Turbulent Conditions. Electronics, 2022, 11(22): 3684. https://doi.org/10.3390/electronics11223684",
        "Yang, P., Yang, J., Jiang, R. Underwater Potential Characteristics of Static Electric Field Related to Ship Corrosion. Journal of Unmanned Undersea Systems, 2023, 31(4): 545-551. https://doi.org/10.11993/j.issn.2096-3920.2023-0056",
        "Bian, Q., Liu, D. H., Wang, X. J. Research on Underwater Static Electric Field for Warship Based on Hybrid Model. Applied Mechanics and Materials, 2015, 713-715: 925-929. https://doi.org/10.4028/www.scientific.net/AMM.713-715.925",
    ]
    for reference in references:
        writer.number(reference)
    writer.callout(
        "说明",
        "文中示意图与归一化曲线用于解释原理，不代表特定舰船实测幅值。公式采用均匀导电介质和等效源近似；工程分析应以实测数据、具体边界条件和原始文献为准。",
        "info",
    )

    # 清理可识别的本机用户信息，并填写通用文档元数据。
    document.core_properties.title = "水下目标静电场与轴频电场仿真原理说明（通俗优化版）"
    document.core_properties.subject = "水下目标静电场、轴频电场及运动目标时序仿真原理"
    document.core_properties.author = "仿真工程文档"
    document.core_properties.keywords = "静电场, 轴频电场, 电流偶极子, 水下目标, 仿真原理"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    """解析命令行参数并生成 Word 文档。"""
    if len(sys.argv) != 3:
        raise SystemExit("用法：build_electric_principles_optimized_docx.py <图片目录> <输出DOCX>")
    build_document(Path(sys.argv[1]), Path(sys.argv[2]))


if __name__ == "__main__":
    main()
