"""把舰艇磁场原理文档中的纯文本公式转换为 Word 原生数学公式。"""

from __future__ import annotations

from collections.abc import Callable, Iterable
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


# 输入文件保持不变，转换结果另存为新的“公式排版版”文档。
PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_DOCX = PROJECT_ROOT / "docs" / "舰艇磁场特性仿真原理讲解（纯原理版）.docx"
OUTPUT_DOCX = PROJECT_ROOT / "docs" / "舰艇磁场特性仿真原理讲解（LaTeX公式排版版）.docx"

# 公式和解释文字沿用原文档的配色与字体。
FORMULA_COLOR = "0B2545"
COVER_FORMULA_COLOR = "2E74B5"
EXPLANATION_COLOR = "607786"
MATH_FONT = "Cambria Math"
BODY_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"


def _append_expression(parent, expression) -> None:
    """把一个公式节点或一组公式节点追加到指定的 OMML 容器。"""

    if expression is None:
        return
    if isinstance(expression, (list, tuple)):
        for child in expression:
            _append_expression(parent, child)
        return
    parent.append(expression)


def math_run(
    text: str,
    *,
    plain: bool = False,
    bold: bool = False,
    color: str = FORMULA_COLOR,
    size_half_points: int = 24,
):
    """创建一个数学文本节点；plain=True 用于函数名和说明性下标。"""

    run = OxmlElement("m:r")
    if plain or bold:
        math_properties = OxmlElement("m:rPr")
        style = OxmlElement("m:sty")
        # 普通函数名使用直立体，矢量变量使用 LaTeX 常见的粗斜体效果。
        style.set(qn("m:val"), "bi" if bold else "p")
        math_properties.append(style)
        run.append(math_properties)

    word_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), MATH_FONT)
    fonts.set(qn("w:hAnsi"), MATH_FONT)
    fonts.set(qn("w:eastAsia"), MATH_FONT)
    word_properties.append(fonts)
    color_element = OxmlElement("w:color")
    color_element.set(qn("w:val"), color)
    word_properties.append(color_element)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(size_half_points))
    word_properties.append(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), str(size_half_points))
    word_properties.append(size_cs)
    if bold:
        word_properties.append(OxmlElement("w:b"))
    run.append(word_properties)

    text_element = OxmlElement("m:t")
    text_element.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text_element.text = text
    run.append(text_element)
    return run


def sequence(*parts):
    """把多个公式片段按从左到右的顺序组合。"""

    result = []
    for part in parts:
        if isinstance(part, (list, tuple)):
            result.extend(part)
        else:
            result.append(part)
    return result


def subscript(base, sub):
    """创建下标结构。"""

    element = OxmlElement("m:sSub")
    base_container = OxmlElement("m:e")
    sub_container = OxmlElement("m:sub")
    _append_expression(base_container, base)
    _append_expression(sub_container, sub)
    element.extend([base_container, sub_container])
    return element


def superscript(base, sup):
    """创建上标结构。"""

    element = OxmlElement("m:sSup")
    base_container = OxmlElement("m:e")
    sup_container = OxmlElement("m:sup")
    _append_expression(base_container, base)
    _append_expression(sup_container, sup)
    element.extend([base_container, sup_container])
    return element


def sub_superscript(base, sub, sup):
    """创建同时包含上下标的结构。"""

    element = OxmlElement("m:sSubSup")
    base_container = OxmlElement("m:e")
    sub_container = OxmlElement("m:sub")
    sup_container = OxmlElement("m:sup")
    _append_expression(base_container, base)
    _append_expression(sub_container, sub)
    _append_expression(sup_container, sup)
    element.extend([base_container, sub_container, sup_container])
    return element


def fraction(numerator, denominator):
    """创建上下堆叠的标准分式。"""

    element = OxmlElement("m:f")
    numerator_container = OxmlElement("m:num")
    denominator_container = OxmlElement("m:den")
    _append_expression(numerator_container, numerator)
    _append_expression(denominator_container, denominator)
    element.extend([numerator_container, denominator_container])
    return element


def radical(radicand):
    """创建省略根指数的平方根结构。"""

    element = OxmlElement("m:rad")
    properties = OxmlElement("m:radPr")
    degree_hidden = OxmlElement("m:degHide")
    degree_hidden.set(qn("m:val"), "1")
    properties.append(degree_hidden)
    degree = OxmlElement("m:deg")
    radicand_container = OxmlElement("m:e")
    _append_expression(radicand_container, radicand)
    element.extend([properties, degree, radicand_container])
    return element


def delimiter(content, begin: str = "(", end: str = ")"):
    """创建会随内容高度自动伸展的括号。"""

    element = OxmlElement("m:d")
    properties = OxmlElement("m:dPr")
    begin_character = OxmlElement("m:begChr")
    begin_character.set(qn("m:val"), begin)
    end_character = OxmlElement("m:endChr")
    end_character.set(qn("m:val"), end)
    properties.extend([begin_character, end_character])
    content_container = OxmlElement("m:e")
    _append_expression(content_container, content)
    element.extend([properties, content_container])
    return element


def integral(lower, upper, integrand):
    """创建带上下限的积分结构。"""

    element = OxmlElement("m:nary")
    properties = OxmlElement("m:naryPr")
    character = OxmlElement("m:chr")
    character.set(qn("m:val"), "∫")
    limit_location = OxmlElement("m:limLoc")
    limit_location.set(qn("m:val"), "subSup")
    hide_sub = OxmlElement("m:subHide")
    hide_sub.set(qn("m:val"), "0")
    hide_sup = OxmlElement("m:supHide")
    hide_sup.set(qn("m:val"), "0")
    properties.extend([character, limit_location, hide_sub, hide_sup])
    sub_container = OxmlElement("m:sub")
    sup_container = OxmlElement("m:sup")
    integrand_container = OxmlElement("m:e")
    _append_expression(sub_container, lower)
    _append_expression(sup_container, upper)
    _append_expression(integrand_container, integrand)
    element.extend([properties, sub_container, sup_container, integrand_container])
    return element


def accent_hat(content):
    """创建矢量单位方向所用的帽形重音符号。"""

    element = OxmlElement("m:acc")
    properties = OxmlElement("m:accPr")
    character = OxmlElement("m:chr")
    character.set(qn("m:val"), "̂")
    properties.append(character)
    content_container = OxmlElement("m:e")
    _append_expression(content_container, content)
    element.extend([properties, content_container])
    return element


def _variable(name: str, *, bold: bool = False):
    """创建单字母变量，便于公式构造代码保持简洁。"""

    return math_run(name, bold=bold)


def _plain(text: str):
    """创建直立体数学文本，例如函数名和说明性下标。"""

    return math_run(text, plain=True)


def _formula_anomaly_sum():
    """综合异常磁场公式。"""

    return sequence(
        subscript(_variable("B", bold=True), _plain("anom")),
        math_run(" = "),
        subscript(_variable("B", bold=True), _plain("static")),
        math_run(" + "),
        subscript(_variable("B", bold=True), _plain("induced")),
    )


def _formula_geomagnetic_components():
    """总地磁强度、磁偏角和磁倾角到 ENU 三分量的换算式。"""

    def trig(name: str, variable: str):
        return sequence(_plain(name), delimiter(_variable(variable)))

    return sequence(
        subscript(_variable("B"), _variable("E")),
        math_run(" = "),
        _variable("F"),
        math_run(" "),
        trig("cos", "I"),
        math_run(" "),
        trig("sin", "D"),
        math_run(",     "),
        subscript(_variable("B"), _variable("N")),
        math_run(" = "),
        _variable("F"),
        math_run(" "),
        trig("cos", "I"),
        math_run(" "),
        trig("cos", "D"),
        math_run(",     "),
        subscript(_variable("B"), _variable("U")),
        math_run(" = -"),
        _variable("F"),
        math_run(" "),
        trig("sin", "I"),
    )


def _formula_half_axes():
    """舰艇长宽高到椭球半轴的换算式。"""

    return sequence(
        _variable("a"), math_run(" = "), fraction(_variable("L"), math_run("2")),
        math_run(",     "),
        _variable("b"), math_run(" = "), fraction(_variable("W"), math_run("2")),
        math_run(",     "),
        _variable("c"), math_run(" = "), fraction(_variable("H"), math_run("2")),
    )


def _formula_volume():
    """等效椭球体积和磁性材料体积公式。"""

    return sequence(
        subscript(_variable("V"), _variable("e")), math_run(" = "),
        fraction(sequence(math_run("4π"), _variable("a"), _variable("b"), _variable("c")), math_run("3")),
        math_run(" = "),
        fraction(sequence(math_run("π"), _variable("L"), _variable("W"), _variable("H")), math_run("6")),
        math_run(",     "),
        subscript(_variable("V"), _variable("m")), math_run(" = η"),
        subscript(_variable("V"), _variable("e")),
    )


def _formula_remanent_moment():
    """剩磁强度转换为目标坐标系和全局坐标系磁矩的公式。"""

    body_moment = sequence(subscript(_variable("m", bold=True), _variable("r")), _plain("(body)"))
    global_moment = sequence(subscript(_variable("m", bold=True), _variable("r")), _plain("(global)"))
    return sequence(
        body_moment, math_run(" = "), subscript(_variable("V"), _variable("m")),
        subscript(_variable("M", bold=True), _variable("r")), math_run(",     "),
        global_moment, math_run(" = "), _variable("R", bold=True), math_run(" "), deepcopy(body_moment),
    )


def _formula_demagnetizing_factor():
    """三轴椭球第 i 个主轴方向的退磁因子积分式。"""

    ai_squared = superscript(subscript(_variable("a"), _variable("i")), math_run("2"))
    a_squared = superscript(_variable("a"), math_run("2"))
    b_squared = superscript(_variable("b"), math_run("2"))
    c_squared = superscript(_variable("c"), math_run("2"))
    first_factor = delimiter(sequence(_variable("s"), math_run(" + "), ai_squared))
    square_root = radical(
        sequence(
            delimiter(sequence(_variable("s"), math_run(" + "), a_squared)),
            delimiter(sequence(_variable("s"), math_run(" + "), b_squared)),
            delimiter(sequence(_variable("s"), math_run(" + "), c_squared)),
        )
    )
    integrand = fraction(_plain("d" + "s"), sequence(first_factor, square_root))
    return sequence(
        subscript(_variable("N"), _variable("i")), math_run(" = "),
        fraction(sequence(_variable("a"), _variable("b"), _variable("c")), math_run("2")),
        math_run(" "), integral(math_run("0"), math_run("∞"), integrand),
    )


def _formula_demagnetizing_sum():
    """三个主轴退磁因子的归一化关系。"""

    return sequence(
        subscript(_variable("N"), _variable("x")), math_run(" + "),
        subscript(_variable("N"), _variable("y")), math_run(" + "),
        subscript(_variable("N"), _variable("z")), math_run(" = 1"),
    )


def _formula_susceptibility():
    """相对磁导率、磁化率和地磁场强度的关系。"""

    return sequence(
        math_run("χ = "), subscript(math_run("μ"), _variable("r")), math_run(" - 1,     "),
        subscript(_variable("H", bold=True), _variable("g")), math_run(" = "),
        fraction(subscript(_variable("B", bold=True), _variable("g")), subscript(math_run("μ"), math_run("0"))),
    )


def _formula_induced_moment():
    """第 k 个主轴方向上的感应磁矩。"""

    return sequence(
        subscript(_variable("m"), math_run("i,k")), math_run(" = "),
        fraction(
            sequence(
                subscript(_variable("V"), _variable("m")), math_run(" χ "),
                subscript(_variable("H"), math_run("g,k")),
            ),
            sequence(math_run("1 + "), subscript(_variable("N"), _variable("k")), math_run(" χ")),
        ),
    )


def _formula_dipole_field():
    """磁偶极子在空间位置 r 处产生的磁场。"""

    r_hat_1 = accent_hat(_variable("r", bold=True))
    r_hat_2 = accent_hat(_variable("r", bold=True))
    bracket_content = sequence(
        math_run("3"), delimiter(sequence(_variable("m", bold=True), math_run(" · "), r_hat_1)),
        r_hat_2, math_run(" - "), _variable("m", bold=True),
    )
    return sequence(
        _variable("B", bold=True), delimiter(_variable("r", bold=True)), math_run(" = "),
        fraction(
            subscript(math_run("μ"), math_run("0")),
            sequence(math_run("4π"), superscript(_variable("r"), math_run("3"))),
        ),
        math_run(" "), delimiter(bracket_content, "[", "]"),
    )


def _formula_field_components():
    """静磁场和感应磁场分别由对应磁矩计算。"""

    return sequence(
        subscript(_variable("B", bold=True), _plain("static")), math_run(" = "),
        _plain("Dipole"), delimiter(subscript(_variable("m", bold=True), _variable("r"))),
        math_run(",     "),
        subscript(_variable("B", bold=True), _plain("induced")), math_run(" = "),
        _plain("Dipole"), delimiter(subscript(_variable("m", bold=True), _variable("i"))),
    )


def _formula_motion():
    """匀速直线运动目标的中心位置公式。"""

    body_displacement = superscript(
        delimiter(sequence(_variable("v"), _variable("t"), math_run(", 0, 0")), "[", "]"),
        _plain("T"),
    )
    return sequence(
        subscript(_variable("r", bold=True), _variable("c")), delimiter(_variable("t")), math_run(" = "),
        subscript(_variable("r", bold=True), math_run("0")), math_run(" + "),
        _variable("R", bold=True), math_run(" "), body_displacement,
    )


def _formula_relative_displacement():
    """固定传感器与目标中心之间的相对位移。"""

    return sequence(
        _variable("d", bold=True), delimiter(_variable("t")), math_run(" = "),
        subscript(_variable("r", bold=True), _plain("sensor")), math_run(" - "),
        subscript(_variable("r", bold=True), _variable("c")), delimiter(_variable("t")),
    )


def _formula_sampling():
    """时序仿真的样本数和采样时刻公式。"""

    floor_expression = delimiter(
        sequence(_variable("T"), math_run(" "), subscript(_variable("f"), _variable("s"))),
        "⌊",
        "⌋",
    )
    return sequence(
        _variable("N"), math_run(" = "), floor_expression, math_run(" + 1,     "),
        subscript(_variable("t"), _variable("k")), math_run(" = "),
        subscript(_variable("t"), _plain("start")), math_run(" + "),
        fraction(_variable("k"), subscript(_variable("f"), _variable("s"))),
    )


# 以文档中的原始公式文本为键，确保转换位置准确且不会误改普通正文。
FORMULA_BUILDERS: dict[str, Callable[[], Iterable]] = {
    "B_anom = B_static + B_induced": _formula_anomaly_sum,
    "B_E = F cos(I) sin(D)    B_N = F cos(I) cos(D)    B_U = -F sin(I)": _formula_geomagnetic_components,
    "a = L/2,    b = W/2,    c = H/2": _formula_half_axes,
    "V_e = 4πabc/3 = πLWH/6,    V_m = ηV_e": _formula_volume,
    "m_r(body) = V_m M_r,    m_r(global) = R m_r(body)": _formula_remanent_moment,
    "N_i = (abc/2) ∫₀∞ ds / [(s+a_i²) √((s+a²)(s+b²)(s+c²))]": _formula_demagnetizing_factor,
    "N_x + N_y + N_z = 1": _formula_demagnetizing_sum,
    "χ = μ_r - 1,    H_g = B_g / μ₀": _formula_susceptibility,
    "m_i,k = V_m χ H_g,k / (1 + N_k χ)": _formula_induced_moment,
    "B(r) = μ₀/(4πr³) · [3(m·r̂)r̂ - m]": _formula_dipole_field,
    "B_static = Dipole(m_r),    B_induced = Dipole(m_i)": _formula_field_components,
    "r_c(t) = r₀ + R · [v t, 0, 0]ᵀ": _formula_motion,
    "d(t) = r_sensor - r_c(t)": _formula_relative_displacement,
    "N = floor(T f_s) + 1,    t_k = t_start + k/f_s": _formula_sampling,
}


def _remove_paragraph_content(paragraph) -> None:
    """保留段落属性，只清除原有纯文本公式与解释文字。"""

    paragraph_element = paragraph._p
    for child in list(paragraph_element):
        if child.tag != qn("w:pPr"):
            paragraph_element.remove(child)


def _set_explanation_font(run) -> None:
    """恢复公式解释文字原有的字号、字体和颜色。"""

    run.font.name = BODY_FONT
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(EXPLANATION_COLOR)
    properties = run._element.get_or_add_rPr()
    fonts = properties.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), BODY_FONT)
    fonts.set(qn("w:hAnsi"), BODY_FONT)
    fonts.set(qn("w:eastAsia"), CJK_FONT)


def _make_math_root(expression, *, cover: bool = False):
    """创建完整的 Word 数学对象，并按封面或正文应用字号和颜色。"""

    color = COVER_FORMULA_COLOR if cover else FORMULA_COLOR
    size_half_points = 30 if cover else 24
    root = OxmlElement("m:oMath")
    # 构造函数先生成通用样式节点，此处统一覆盖颜色和字号，避免逐式重复设置。
    for node in expression:
        root.append(node)
    for math_text_run in root.iter(qn("m:r")):
        word_properties = math_text_run.find(qn("w:rPr"))
        if word_properties is None:
            continue
        color_element = word_properties.find(qn("w:color"))
        if color_element is not None:
            color_element.set(qn("w:val"), color)
        for size_tag in ("w:sz", "w:szCs"):
            size_element = word_properties.find(qn(size_tag))
            if size_element is not None:
                size_element.set(qn("w:val"), str(size_half_points))
    return root


def convert_document() -> tuple[Path, int]:
    """转换封面公式和所有公式块，返回输出路径与转换数量。"""

    if not SOURCE_DOCX.exists():
        raise FileNotFoundError(f"找不到源文档：{SOURCE_DOCX}")

    document = Document(SOURCE_DOCX)
    converted_count = 0
    for paragraph in document.paragraphs:
        full_text = paragraph.text
        first_line, separator, explanation = full_text.partition("\n")
        builder = FORMULA_BUILDERS.get(first_line)
        if builder is None:
            continue

        is_formula_block = paragraph.style is not None and paragraph.style.name == "Equation Block"
        is_cover_formula = not is_formula_block and full_text == first_line
        # 只转换公式块以及封面的独立公式，避免碰到正文中偶然出现的相同字符串。
        if not is_formula_block and not is_cover_formula:
            continue

        _remove_paragraph_content(paragraph)
        paragraph._p.append(_make_math_root(builder(), cover=is_cover_formula))
        if separator and explanation:
            explanation_run = paragraph.add_run()
            explanation_run.add_break()
            explanation_run.add_text(explanation)
            _set_explanation_font(explanation_run)
        converted_count += 1

    # 当前文档应包含封面 1 处和正文公式块 14 处；数量不符时立即停止，避免交付漏改文档。
    expected_count = 15
    if converted_count != expected_count:
        raise RuntimeError(f"应转换 {expected_count} 处公式，实际转换 {converted_count} 处")

    document.core_properties.title = "舰艇磁场特性仿真原理讲解（LaTeX公式排版版）"
    document.core_properties.comments = (
        "在纯原理版基础上，将封面和正文中的独立公式转换为 Word 原生数学公式；"
        "原有正文、图示、流程图和参考文献保持不变。"
    )
    document.save(OUTPUT_DOCX)
    return OUTPUT_DOCX, converted_count


def validate_output(path: Path, expected_count: int) -> None:
    """校验输出文件能够打开，并包含预期数量的 Word 原生公式对象。"""

    document = Document(path)
    native_math_count = len(document.element.body.findall(".//" + qn("m:oMath")))
    if native_math_count != expected_count:
        raise RuntimeError(
            f"原生公式数量校验失败：应为 {expected_count}，实际为 {native_math_count}"
        )
    remaining_text_formulas = [
        paragraph.text.partition("\n")[0]
        for paragraph in document.paragraphs
        if paragraph.style is not None
        and paragraph.style.name == "Equation Block"
        and paragraph.text.partition("\n")[0] in FORMULA_BUILDERS
    ]
    if remaining_text_formulas:
        raise RuntimeError(f"仍有未转换的纯文本公式：{remaining_text_formulas}")


if __name__ == "__main__":
    output_path, count = convert_document()
    validate_output(output_path, count)
    print(f"已生成：{output_path}")
    print(f"已转换 Word 原生公式：{count} 处")
