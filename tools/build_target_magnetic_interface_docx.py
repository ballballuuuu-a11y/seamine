"""生成《舰艇磁场特性仿真数据结构与函数说明》Word 文档。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

from build_minesweeping_interface_docx import (
    BLUE,
    CAUTION,
    DARK_BLUE,
    MUTED,
    NAVY,
    add_body,
    add_bullet,
    add_callout,
    add_code_block,
    add_heading,
    add_numbered,
    add_numbering_definitions,
    add_page_break,
    add_page_field,
    add_table,
    configure_page,
    configure_styles,
    set_paragraph_border,
    set_run_font,
)


# 文档输出路径。
PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DOCX = PROJECT_ROOT / "docs" / "舰艇磁场特性仿真数据结构与函数说明.docx"


def configure_header_footer(document: Document) -> None:
    """配置舰艇磁场接口说明的运行页眉和页脚。"""

    section = document.sections[0]
    section.first_page_header.paragraphs[0].text = ""

    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(3)
    header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = header.add_run("舰艇磁场特性仿真")
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    header.add_run("\t")
    right = header.add_run("数据结构与函数说明")
    set_run_font(right, size=8.5, color=MUTED)
    set_paragraph_border(header, side="bottom", color="D8E2E8", size=5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(footer)


def build_cover(document: Document) -> None:
    """创建简洁的技术接口说明首页。"""

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(4)
    kicker_run = kicker.add_run("调用接口参考")
    set_run_font(kicker_run, size=10.5, color=CAUTION, bold=True)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("舰艇磁场特性仿真")
    set_run_font(title_run, size=28, color=NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle_run = subtitle.add_run("数据结构、成员含义与函数输入输出说明")
    set_run_font(subtitle_run, size=14, color=DARK_BLUE)

    metadata = [
        ("对应类", "TargetMagneticFieldModel"),
        ("接口文件", "TargetMagneticFieldModel.h"),
        ("实现文件", "TargetMagneticFieldModel.cpp"),
        ("坐标约定", "ENU：x 向东、y 向北、z 向上"),
        ("主要用途", "单点磁场、运动目标时序、二维/三维空间网格"),
        ("更新时间", "2026-07-30"),
    ]
    for label, value in metadata:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        label_run = paragraph.add_run(f"{label}：")
        set_run_font(label_run, size=10.5, color=NAVY, bold=True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=10.5, color=NAVY)

    rule = document.add_paragraph()
    rule.paragraph_format.space_before = Pt(9)
    rule.paragraph_format.space_after = Pt(18)
    set_paragraph_border(rule, side="bottom", color=BLUE, size=12)

    add_callout(
        document,
        "一句话理解",
        "先用 TargetParameter 描述舰艇，再调用 calculate、simulate 或 simulateGrid；所有结果都放在 MagneticFieldSample 中。",
    )
    add_page_break(document)


def build_document() -> Path:
    """构建完整的数据结构与函数接口说明。"""

    document = Document()
    document.core_properties.title = "舰艇磁场特性仿真数据结构与函数说明"
    document.core_properties.subject = "TargetMagneticFieldModel 数据结构、成员与函数接口"
    document.core_properties.author = "seamine 项目组"
    document.core_properties.last_modified_by = "seamine 项目组"
    document.core_properties.keywords = "舰艇磁场, 数据结构, 函数接口, 三分量, C++"
    document.core_properties.comments = "按当前工程接口整理的通俗调用说明。"
    configure_styles(document)
    configure_page(document)
    configure_header_footer(document)
    bullet_id, decimal_id = add_numbering_definitions(document)
    build_cover(document)

    add_heading(document, "1. 先了解整体调用方式", level=1)
    add_body(
        document,
        "该模型把舰艇或潜艇等效成一个同时具有剩磁和地磁感应磁矩的三轴椭球，再用磁偶极子公式计算目标外部的 Bx、By、Bz 三分量。调用人员不需要直接处理退磁积分和坐标旋转，只需要正确填写参数并选择合适的计算函数。"
    )
    add_heading(document, "1.1 最常用的调用步骤", level=2)
    for text in [
        "创建 TargetParameter，填写尺寸、位置、航速、姿态、地磁和磁性参数。",
        "使用带参数的构造函数创建模型，或对默认构造的模型调用 setTarget。",
        "单点计算调用 calculate；固定传感器时序调用 simulate；空间分布调用 simulateGrid。",
        "从 MagneticFieldSample 中读取静磁场、感应磁场和综合异常磁场三分量及模值。",
    ]:
        add_numbered(document, text, decimal_id)

    add_heading(document, "1.2 三个必须记住的口径", level=2)
    add_callout(
        document,
        "配置状态",
        "默认构造函数只创建空模型，不能立即计算；必须先调用 setTarget，否则计算函数会抛出 std::logic_error。",
        accent=CAUTION,
    )
    add_callout(
        document,
        "时间口径",
        "calculate(observationPosition) 和 simulateGrid(grid) 都使用零时刻目标位置。需要运动时序时应调用 calculate(observationPosition, time) 或 simulate。",
        accent=CAUTION,
    )
    add_callout(
        document,
        "磁场口径",
        "totalFieldVector 只等于 staticFieldVector + inducedFieldVector，不包含 geomagneticFieldNt。若需要绝对磁场读数，应由调用方再加背景地磁。",
        accent=CAUTION,
    )

    add_heading(document, "2. 坐标与单位约定", level=1)
    add_table(
        document,
        ["项目", "约定", "通俗解释"],
        [
            ["全局坐标", "x 东、y 北、z 上", "水下位置通常使用负 z"],
            ["目标坐标", "x 沿舰首、y 横向、z 垂向", "剩磁强度在目标自身坐标系填写"],
            ["位置和尺寸", "m", "长度、宽度、高度、位置、距离统一使用米"],
            ["航速", "m/s", "目标沿自身纵轴 x 正方向匀速运动"],
            ["姿态角", "°", "航向、俯仰和横滚均使用度"],
            ["背景地磁", "nT", "按照东、北、上三分量填写"],
            ["剩磁强度", "A/m", "描述单位磁性体积的剩磁强弱和方向"],
            ["输出磁场", "nT", "三分量和模值统一输出为纳特"],
        ],
        [1800, 2600, 4960],
        font_size=8.8,
    )
    add_page_break(document)

    add_heading(document, "3. 数据结构总览", level=1)
    add_table(
        document,
        ["数据结构", "主要作用", "什么时候使用"],
        [
            ["TargetType", "区分水面舰艇和潜艇", "填写 TargetParameter.type"],
            ["Vector3", "保存三维位置或三分量矢量", "位置、地磁、剩磁和磁场结果"],
            ["TargetParameter", "完整描述一个待仿真的目标", "构造模型或调用 setTarget"],
            ["MagneticFieldSample", "保存一个点、一个时刻的磁场结果", "calculate、simulate、simulateGrid 的输出"],
            ["GridParameter", "描述规则二维或三维计算网格", "调用 simulateGrid"],
        ],
        [2500, 4000, 2860],
        font_size=8.8,
    )

    add_heading(document, "4. TargetType：目标类型", level=1)
    add_table(
        document,
        ["枚举值", "含义", "当前计算差异"],
        [
            ["SurfaceShip", "水面舰艇", "当前主要用于业务分类，不改变磁场公式"],
            ["Submarine", "潜艇", "当前主要用于业务分类，不改变磁场公式"],
        ],
        [2600, 3300, 3460],
        font_size=9.0,
    )
    add_callout(
        document,
        "容易误解",
        "仅修改 type 不会自动改变尺寸、深度、磁导率或剩磁；调用方仍需填写与实际目标相符的完整参数。",
    )

    add_heading(document, "5. Vector3：三维坐标或三分量", level=1)
    add_table(
        document,
        ["成员", "类型", "含义", "单位"],
        [
            ["x", "double", "东向坐标或东向分量", "由所在字段决定"],
            ["y", "double", "北向坐标或北向分量", "由所在字段决定"],
            ["z", "double", "上向坐标或上向分量", "由所在字段决定"],
        ],
        [1700, 1700, 4100, 1860],
        font_size=8.9,
    )
    add_body(
        document,
        "同一个 Vector3 可能表示位置、地磁、剩磁或计算结果，因此必须结合字段名称判断单位。例如 center 的单位是 m，geomagneticFieldNt 的单位是 nT，remanentMagnetizationAm 的单位是 A/m。"
    )
    add_page_break(document)

    add_heading(document, "6. TargetParameter：目标输入参数", level=1)
    add_body(
        document,
        "TargetParameter 是最重要的输入结构。setTarget 会先校验全部成员，再计算并缓存静态磁矩和感应磁矩；如果校验失败，原来已经可用的模型配置不会被破坏。"
    )
    add_table(
        document,
        ["成员", "类型 / 默认值", "单位", "通俗含义与影响"],
        [
            ["type", "TargetType / SurfaceShip", "无", "目标业务类型；当前不改变物理公式"],
            ["length", "double / 0", "m", "目标长度；影响等效体积和纵向退磁因子，必须>0"],
            ["width", "double / 0", "m", "目标宽度；影响等效体积和横向退磁因子，必须>0"],
            ["height", "double / 0", "m", "目标高度；影响等效体积和垂向退磁因子，必须>0"],
            ["center", "Vector3 / {0,0,0}", "m", "零时刻目标几何中心的全局坐标"],
            ["velocity", "double / 0", "m/s", "沿目标纵轴匀速运动的速度，必须≥0"],
            ["headingDegrees", "double / 0", "°", "航向角；全局 x 东向为 0°，逆时针为正"],
            ["pitchDegrees", "double / 0", "°", "俯仰角；抬头为正"],
            ["rollDegrees", "double / 0", "°", "横滚角；绕目标纵轴旋转"],
            ["geomagneticFieldNt", "Vector3 / {0,30000,-40000}", "nT", "背景地磁东、北、上分量；只影响感应磁矩"],
            ["relativePermeability", "double / 200", "无", "等效相对磁导率，表示材料容易被磁化的程度，必须≥1"],
            ["magneticMaterialRatio", "double / 0.04", "0～1", "等效铁磁材料占椭球体积的比例，范围必须为 (0,1]"],
            ["remanentMagnetizationAm", "Vector3 / {8,0,0}", "A/m", "目标坐标系剩磁强度；决定静态磁矩的大小和方向"],
            ["minimumDistance", "double / 1", "m", "目标中心额外最小观测距离；模型同时禁止观测点进入等效舰体椭球"],
        ],
        [2600, 2500, 1000, 3260],
        font_size=8.15,
    )
    add_heading(document, "6.1 参数之间的关系", level=2)
    for text in [
        "长度、宽度和高度共同决定等效椭球体积，不是只看长度。",
        "magneticMaterialRatio 同时影响静态磁矩和感应磁矩。",
        "geomagneticFieldNt 显式设置为 {0,0,0} 时，感应磁场为零，但静磁场仍可能存在。",
        "remanentMagnetizationAm 显式设置为 {0,0,0} 时，静磁场为零，但地磁感应磁场仍可能存在。",
        "航向、俯仰和横滚会旋转磁矩方向，也会改变目标运动方向。",
    ]:
        add_bullet(document, text, bullet_id)

    add_heading(document, "7. MagneticFieldSample：单点计算结果", level=1)
    add_table(
        document,
        ["成员", "类型 / 单位", "含义"],
        [
            ["observationPosition", "Vector3 / m", "传入的观测点或传感器全局坐标"],
            ["targetPosition", "Vector3 / m", "当前时刻计算得到的目标中心位置"],
            ["staticFieldVector", "Vector3 / nT", "目标剩磁产生的 Bx、By、Bz"],
            ["inducedFieldVector", "Vector3 / nT", "背景地磁使目标磁化后产生的 Bx、By、Bz"],
            ["totalFieldVector", "Vector3 / nT", "静磁场与感应磁场的三分量矢量和，不含背景地磁"],
            ["staticField", "double / nT", "staticFieldVector 的模值，始终非负"],
            ["inducedField", "double / nT", "inducedFieldVector 的模值，始终非负"],
            ["totalField", "double / nT", "totalFieldVector 的模值，始终非负"],
            ["distance", "double / m", "当前目标中心到观测点的直线距离"],
            ["time", "double / s", "当前样本对应的仿真时间"],
        ],
        [2800, 2300, 4260],
        font_size=8.45,
    )
    add_callout(
        document,
        "分量与模值",
        "三分量保留方向和正负号；模值只表示强弱。做波形、过零点或敏感轴分析时应使用三分量，做热力图和峰值统计时通常使用模值。",
    )

    add_heading(document, "8. GridParameter：规则空间网格", level=1)
    add_table(
        document,
        ["成员", "类型 / 默认值", "含义与约束"],
        [
            ["minimum", "Vector3 / {0,0,0}", "网格最小坐标，三个分量都必须≤maximum 对应分量"],
            ["maximum", "Vector3 / {0,0,0}", "网格最大坐标"],
            ["xCount", "size_t / 1", "x 方向采样点数，必须>0"],
            ["yCount", "size_t / 1", "y 方向采样点数，必须>0"],
            ["zCount", "size_t / 1", "z 方向采样点数，必须>0"],
        ],
        [2400, 2400, 4560],
        font_size=8.7,
    )
    add_body(
        document,
        "网格的起点和终点都包含在结果中。某个方向的点数为 1 时，该方向只取 minimum 对应坐标。结果按 z、y、x 三层循环展开，也就是 x 变化最快、z 变化最慢。"
    )
    add_callout(
        document,
        "数量限制",
        "单次网格最多生成 1,000,000 个点。总点数为 xCount×yCount×zCount，超过限制会抛出 std::length_error。",
        accent=CAUTION,
    )
    add_page_break(document)

    add_heading(document, "9. 公共函数总览", level=1)
    add_table(
        document,
        ["函数", "主要作用", "返回值"],
        [
            ["TargetMagneticFieldModel()", "创建尚未配置参数的空模型", "模型对象"],
            ["TargetMagneticFieldModel(param)", "创建并立即配置可计算模型", "模型对象"],
            ["setTarget(param)", "校验并更新目标参数，同时重建等效磁矩", "void"],
            ["target()", "读取当前生效的目标参数", "const TargetParameter&"],
            ["calculate(position)", "计算零时刻单个观测点", "MagneticFieldSample"],
            ["calculate(position, time)", "计算指定时刻单个观测点", "MagneticFieldSample"],
            ["simulate(position, start, duration, rate)", "生成固定观测点的运动磁场时序", "vector<MagneticFieldSample>"],
            ["simulateGrid(grid)", "生成零时刻规则空间磁场分布", "vector<MagneticFieldSample>"],
        ],
        [3850, 3670, 1840],
        font_size=8.35,
    )

    add_heading(document, "10. 构造和配置函数", level=1)
    add_heading(document, "10.1 默认构造函数", level=2)
    add_code_block(document, "TargetMagneticFieldModel(); // 创建空模型，必须先调用 setTarget 才能计算。")
    add_table(
        document,
        ["项目", "说明"],
        [
            ["作用", "创建一个尚未配置目标参数的模型对象"],
            ["输入", "无"],
            ["输出", "TargetMagneticFieldModel 对象"],
            ["异常", "构造时不抛出；若未配置就计算，会抛出 std::logic_error"],
        ],
        [1800, 7560],
        font_size=8.8,
    )

    add_heading(document, "10.2 带参数构造函数", level=2)
    add_code_block(
        document,
        "explicit TargetMagneticFieldModel(const TargetParameter& param); // 校验参数并创建可计算模型。",
    )
    add_table(
        document,
        ["项目", "说明"],
        [
            ["作用", "用完整参数创建模型，并立即计算静态、感应等效磁矩"],
            ["输入", "param：完整 TargetParameter"],
            ["输出", "已配置、可立即调用 calculate/simulate/simulateGrid 的模型"],
            ["异常", "参数非法时抛出 std::invalid_argument"],
        ],
        [1800, 7560],
        font_size=8.8,
    )

    add_heading(document, "10.3 setTarget", level=2)
    add_code_block(
        document,
        "void setTarget(const TargetParameter& param); // 更新参数并重建缓存磁矩。",
    )
    add_table(
        document,
        ["项目", "说明"],
        [
            ["作用", "校验新参数，保存后重建静态磁矩和感应磁矩"],
            ["输入", "param：新的完整目标参数"],
            ["输出", "无；成功后新参数立即生效"],
            ["异常", "参数非法时抛出 std::invalid_argument，原配置保持不变"],
            ["性能提示", "该函数包含退磁因子数值计算，不应在每个采样点重复调用"],
        ],
        [1800, 7560],
        font_size=8.8,
    )

    add_heading(document, "10.4 target", level=2)
    add_code_block(
        document,
        "const TargetParameter& target() const; // 只读返回当前生效参数。",
    )
    add_table(
        document,
        ["项目", "说明"],
        [
            ["作用", "让其他模块读取模型当前使用的参数"],
            ["输入", "无"],
            ["输出", "当前 TargetParameter 的常量引用，调用方不能通过该引用修改模型"],
            ["异常", "模型尚未配置时抛出 std::logic_error"],
        ],
        [1800, 7560],
        font_size=8.8,
    )
    add_page_break(document)

    add_heading(document, "11. 单点计算函数", level=1)
    add_heading(document, "11.1 calculate(position)", level=2)
    add_code_block(
        document,
        "MagneticFieldSample calculate(const Vector3& observationPosition) const; // 计算零时刻。",
    )
    add_table(
        document,
        ["项目", "说明"],
        [
            ["作用", "计算零时刻某个空间位置的静磁场、感应磁场和综合异常磁场"],
            ["输入", "observationPosition：观测点全局坐标，单位 m"],
            ["输出", "一个 MagneticFieldSample，time 固定为 0"],
            ["异常", "未配置时抛出 logic_error；坐标非法时抛出 invalid_argument；距离过近时抛出 domain_error"],
        ],
        [1800, 7560],
        font_size=8.7,
    )

    add_heading(document, "11.2 calculate(position, time)", level=2)
    add_code_block(
        document,
        """MagneticFieldSample calculate(
    const Vector3& observationPosition,
    double timeSeconds) const; // 计算指定时刻。""",
    )
    add_table(
        document,
        ["项目", "说明"],
        [
            ["作用", "先计算目标在指定时刻的位置，再计算观测点磁场"],
            ["输入1", "observationPosition：固定观测点全局坐标，单位 m"],
            ["输入2", "timeSeconds：相对零时刻的时间，单位 s，可以为负"],
            ["目标运动", "目标沿自身纵轴以 velocity 匀速运动，姿态保持不变"],
            ["输出", "一个 MagneticFieldSample，包含当前 targetPosition、distance 和 time"],
            ["异常", "未配置、坐标/时间非有限、运动位置溢出或观测距离过近时抛出相应异常"],
        ],
        [1800, 7560],
        font_size=8.7,
    )
    add_body(
        document,
        "当前目标中心位置可理解为：零时刻 center，加上“目标纵轴方向 × velocity × timeSeconds”。因此航向和俯仰会影响全局运动方向，横滚通常不改变纵轴方向。"
    )

    add_heading(document, "12. 时间序列函数 simulate", level=1)
    add_code_block(
        document,
        """std::vector<MagneticFieldSample> simulate(
    const Vector3& observationPosition,
    double startTimeSeconds,
    double durationSeconds,
    double sampleRateHz) const; // 在固定观测点生成时序。""",
    )
    add_table(
        document,
        ["参数", "输入/输出", "含义"],
        [
            ["observationPosition", "输入 / Vector3 / m", "固定传感器或观测点坐标"],
            ["startTimeSeconds", "输入 / double / s", "时序开始时刻，可以为负"],
            ["durationSeconds", "输入 / double / s", "持续时间，必须≥0"],
            ["sampleRateHz", "输入 / double / Hz", "每秒采样次数，必须>0"],
            ["返回值", "输出 / vector<MagneticFieldSample>", "按时间先后排列的完整磁场样本"],
        ],
        [2600, 2800, 3960],
        font_size=8.7,
    )
    add_body(
        document,
        "采样点数量为 N=floor(durationSeconds×sampleRateHz)+1。第 k 个样本时刻为 startTimeSeconds+k/sampleRateHz。结果包含起点，并保留持续时间范围内最后一个完整采样点。"
    )
    add_callout(
        document,
        "数量限制",
        "单次时间序列最多生成 1,000,000 个样本。持续时间与采样率乘积过大时会抛出 std::length_error。",
        accent=CAUTION,
    )

    add_heading(document, "13. 空间网格函数 simulateGrid", level=1)
    add_code_block(
        document,
        "std::vector<MagneticFieldSample> simulateGrid(const GridParameter& grid) const; // 生成零时刻规则网格。",
    )
    add_table(
        document,
        ["项目", "说明"],
        [
            ["作用", "计算零时刻二维平面或三维空间中的规则磁场分布"],
            ["输入", "grid：网格边界与三个方向采样点数"],
            ["输出", "按 z、y、x 顺序展开的 MagneticFieldSample 数组"],
            ["时间", "所有样本 time=0；当前接口不能直接生成运动时刻网格"],
            ["异常", "网格非法时抛出 invalid_argument；总点数过大时抛出 length_error；某点过近时抛出 domain_error"],
        ],
        [1800, 7560],
        font_size=8.7,
    )
    add_callout(
        document,
        "二维网格",
        "需要计算固定深度平面时，可令 minimum.z=maximum.z，并设置 zCount=1。",
    )
    add_page_break(document)

    add_heading(document, "14. 内部关键函数", level=1)
    add_body(
        document,
        "以下函数是模型内部实现细节，其他业务功能不应直接调用，但了解它们有助于理解参数何时校验、磁矩何时重建以及异常从哪里产生。"
    )
    add_table(
        document,
        ["内部函数", "作用", "何时调用"],
        [
            ["validateTarget", "检查尺寸、位置、航速、姿态、地磁、剩磁、磁导率、材料比例和最小距离", "构造函数和 setTarget"],
            ["validateGrid", "检查网格边界顺序及各方向点数", "simulateGrid"],
            ["rebuildDipoleMoments", "根据等效椭球体积、退磁因子、剩磁和地磁重建两个磁矩缓存", "setTarget 成功后"],
            ["ensureConfigured", "确认模型已经配置", "所有读取或计算接口"],
            ["calculateDipoleField", "根据磁矩、目标到观测点位移和距离计算三分量磁场", "calculate"],
            ["calculateDemagnetizingFactors", "根据长宽高计算三个主轴退磁因子", "rebuildDipoleMoments"],
            ["bodyToGlobal / globalToBody", "在目标坐标系与东-北-上坐标系之间旋转矢量", "参数配置和运动计算"],
            ["gridCoordinate", "按点数在最小值和最大值之间生成网格坐标", "simulateGrid"],
        ],
        [2950, 4350, 2060],
        font_size=8.4,
    )

    add_heading(document, "15. 异常类型与处理建议", level=1)
    add_table(
        document,
        ["异常", "常见原因", "调用方建议"],
        [
            ["std::invalid_argument", "参数非有限、尺寸≤0、航速<0、磁导率<1、材料比例越界、时间或采样率非法", "在界面层提示具体参数错误并阻止计算"],
            ["std::logic_error", "默认构造后没有调用 setTarget，或在未配置状态读取 target", "检查初始化顺序"],
            ["std::domain_error", "观测点进入等效舰体椭球，或到目标中心距离小于 minimumDistance", "调整观测点、网格范围或最小距离；不要简单忽略"],
            ["std::length_error", "时间样本或网格点超过 1,000,000", "降低采样率、缩小网格或分批调用"],
        ],
        [2500, 4300, 2560],
        font_size=8.5,
    )

    add_heading(document, "16. 完整调用示例", level=1)
    add_code_block(
        document,
        """using Model = TargetMagneticFieldModel;

Model::TargetParameter target;
target.type = Model::TargetType::SurfaceShip;           // 设置为水面舰艇。
target.length = 100.0;                                  // 舰长，单位 m。
target.width = 15.0;                                    // 舰宽，单位 m。
target.height = 10.0;                                   // 舰高，单位 m。
target.center = {-300.0, 40.0, -5.0};                  // 零时刻目标中心位置。
target.velocity = 8.0;                                  // 沿舰首方向航速，单位 m/s。
target.headingDegrees = 0.0;                            // 舰首朝东。
target.pitchDegrees = 0.0;                              // 无俯仰。
target.rollDegrees = 0.0;                               // 无横滚。
target.geomagneticFieldNt = {0.0, 30000.0, -40000.0};  // 东、北、上地磁分量。
target.relativePermeability = 200.0;                    // 等效相对磁导率。
target.magneticMaterialRatio = 0.035;                   // 等效铁磁材料体积占比。
target.remanentMagnetizationAm = {8.0, 0.0, 0.0};      // 目标坐标系剩磁强度。
target.minimumDistance = 1.0;                           // 最小允许观测距离。

const Model model(target);                              // 创建可立即计算的模型。

const Model::Vector3 sensor{0.0, 0.0, -30.0};          // 固定磁传感器位置。
const auto instant = model.calculate(sensor, 20.0);     // 计算第 20 s 的三分量磁场。

const auto bx = instant.totalFieldVector.x;              // 综合异常磁场东向分量。
const auto by = instant.totalFieldVector.y;              // 综合异常磁场北向分量。
const auto bz = instant.totalFieldVector.z;              // 综合异常磁场上向分量。
const auto magnitude = instant.totalField;               // 综合异常磁场模值。

const auto series = model.simulate(
    sensor,
    0.0,                                                 // 时序起始时刻。
    75.0,                                                // 持续时间，单位 s。
    20.0);                                               // 采样率，单位 Hz。

Model::GridParameter grid;
grid.minimum = {-500.0, -300.0, -30.0};                 // 网格最小坐标。
grid.maximum = {500.0, 300.0, -30.0};                   // 网格最大坐标。
grid.xCount = 101;                                       // x 方向采样点数。
grid.yCount = 61;                                        // y 方向采样点数。
grid.zCount = 1;                                         // 固定深度二维平面。

const auto spatial = model.simulateGrid(grid);           // 生成零时刻空间分布。""",
    )

    add_heading(document, "17. 常见问题", level=1)
    add_table(
        document,
        ["问题", "通俗回答"],
        [
            ["不设置地磁会怎样？", "如果完全不赋值，会使用默认 {0,30000,-40000} nT；若显式设为零，感应磁场为零。"],
            ["totalField 是绝对磁场吗？", "不是。它只包含目标静磁场和感应异常磁场，不包含背景地磁。"],
            ["为什么模值不能代替三分量？", "模值没有方向和正负号，无法表示过零、反向和传感器敏感轴响应。"],
            ["simulateGrid 能计算运动时刻吗？", "当前不能，它固定调用零时刻 calculate；运动单点应使用带 time 的 calculate。"],
            ["为什么离目标太近会报错？", "模型把整艘舰艇等效成一个点磁偶极子；观测点进入由长、宽、高和姿态确定的等效椭球，或距中心小于最小距离时，这一近似不再可靠。"],
            ["修改 target() 返回值能更新模型吗？", "不能。target() 返回常量引用；必须调用 setTarget 才会重新校验并重建磁矩。"],
            ["能否并行调用 calculate？", "模型配置完成后计算函数为 const，通常适合只读并行；不要并发调用 setTarget。"],
        ],
        [3000, 6360],
        font_size=8.6,
    )

    add_heading(document, "18. 调用方检查清单", level=1)
    for text in [
        "确认使用东-北-上坐标，水下 z 使用负值。",
        "确认尺寸、位置和距离使用 m，地磁和输出使用 nT，剩磁强度使用 A/m。",
        "默认构造后先调用 setTarget。",
        "根据业务选择三分量或模值，明确综合异常场是否需要再叠加背景地磁。",
        "时序计算前检查持续时间、采样率和最大样本数量。",
        "空间网格计算前检查边界顺序、点数和 z、y、x 展开顺序。",
        "避免观测点进入 minimumDistance 以内。",
        "参数变化后必须通过 setTarget 生效，不能只修改原 TargetParameter 变量。",
    ]:
        add_bullet(document, text, bullet_id)

    add_callout(
        document,
        "总结",
        "TargetParameter 负责“描述目标”，GridParameter 负责“描述空间范围”，MagneticFieldSample 负责“承载结果”；calculate、simulate 和 simulateGrid 分别对应单点、时序和空间三类调用。",
    )

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


def main() -> None:
    """生成并保存接口说明 Word 文档。"""

    output = build_document()
    print(f"已生成 Word 文档：{output}")


if __name__ == "__main__":
    main()
