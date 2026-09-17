"""把电场接口 Markdown 说明转换成排版统一的 Word 文档。"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# “compact_reference_guide”版式参数，数值与文档技能中的预设保持一致。
PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11.0
MARGIN_IN = 1.0
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_TOP_DXA = 80
CELL_MARGIN_BOTTOM_DXA = 80
CELL_MARGIN_START_DXA = 120
CELL_MARGIN_END_DXA = 120

BODY_FONT = "Calibri"
CHINESE_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"
COLOR_BLUE = "2E74B5"
COLOR_DARK_BLUE = "1F4D78"
COLOR_NAVY = "203748"
COLOR_MUTED = "667085"
COLOR_TABLE_HEADER = "E8EEF5"
COLOR_TABLE_ALT = "F7F9FC"
COLOR_CODE_FILL = "F3F5F7"
COLOR_BORDER = "B8C4D1"


def set_run_font(
    run,
    *,
    name: str = BODY_FONT,
    chinese_name: str = CHINESE_FONT,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    """设置中英文字体、字号、颜色和字形。"""
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), chinese_name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(
    paragraph,
    *,
    before: float = 0,
    after: float = 6,
    line_spacing: float = 1.25,
) -> None:
    """统一设置段前、段后和行距。"""
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = line_spacing


def configure_style(
    document: Document,
    style_name: str,
    *,
    size: float,
    color: str,
    bold: bool,
    before: float,
    after: float,
    keep_with_next: bool = False,
) -> None:
    """设置指定 Word 样式，避免依赖本机默认模板。"""
    style = document.styles[style_name]
    style.font.name = BODY_FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), BODY_FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), BODY_FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    style.paragraph_format.line_spacing = 1.25
    style.paragraph_format.keep_with_next = keep_with_next


def add_page_field(paragraph) -> None:
    """在页脚中插入动态页码字段。"""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=9, color=COLOR_MUTED)


def set_cell_shading(cell, fill: str) -> None:
    """设置表格单元格底色。"""
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell) -> None:
    """设置表格单元格内边距，避免文字贴边。"""
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (
        ("top", CELL_MARGIN_TOP_DXA),
        ("bottom", CELL_MARGIN_BOTTOM_DXA),
        ("start", CELL_MARGIN_START_DXA),
        ("end", CELL_MARGIN_END_DXA),
    ):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    """设置清晰但不过重的表格边框。"""
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), COLOR_BORDER)


def set_repeat_table_header(row) -> None:
    """让跨页表格在每一页重复显示表头。"""
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    """设置固定 DXA 表格宽度，使表格在不同 Word 环境中保持稳定。"""
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr

    table_width = properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.insert(0, table_width)
    table_width.set(qn("w:w"), str(sum(widths_dxa)))
    table_width.set(qn("w:type"), "dxa")

    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            cell.width = Inches(width / 1440)
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.first_child_found_in("w:tcW")
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.append(cell_width)
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def choose_column_widths(rows: list[list[str]]) -> list[int]:
    """根据内容长度分配列宽，同时保证总宽度严格为 9360 DXA。"""
    column_count = len(rows[0])
    preferred = []
    for column_index in range(column_count):
        lengths = [
            min(45, max(4, len(row[column_index]) if column_index < len(row) else 0))
            for row in rows
        ]
        preferred.append(max(lengths))

    minimum = 900 if column_count >= 4 else 1200
    raw_total = sum(preferred)
    widths = [
        max(minimum, round(CONTENT_WIDTH_DXA * value / raw_total))
        for value in preferred
    ]

    # 如果应用最小列宽后超过页面宽度，就按比例重新压缩。
    if sum(widths) > CONTENT_WIDTH_DXA:
        scale = CONTENT_WIDTH_DXA / sum(widths)
        widths = [max(700, round(value * scale)) for value in widths]

    # 将舍入误差集中到最后一列，保证宽度总和精确。
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def add_inline_runs(paragraph, text: str, *, default_size: float = 11) -> None:
    """解析粗体和行内代码，并添加到同一段落中。"""
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, size=default_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=default_size, bold=True, color=COLOR_DARK_BLUE)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(
                run,
                name=CODE_FONT,
                chinese_name=CHINESE_FONT,
                size=max(9, default_size - 0.5),
                color=COLOR_DARK_BLUE,
            )
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=default_size)


def add_numbering_definitions(document: Document) -> tuple[int, int]:
    """创建真实的项目符号和数字编号定义。"""
    numbering = document.part.numbering_part.element

    def create_numbering(num_id: int, abstract_id: int, kind: str) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)

        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)

        num_format = OxmlElement("w:numFmt")
        num_format.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
        level.append(num_format)

        level_text = OxmlElement("w:lvlText")
        level_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
        level.append(level_text)

        justification = OxmlElement("w:lvlJc")
        justification.set(qn("w:val"), "left")
        level.append(justification)

        paragraph_properties = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        paragraph_properties.append(tabs)
        indentation = OxmlElement("w:ind")
        indentation.set(qn("w:left"), "540")
        indentation.set(qn("w:hanging"), "270")
        paragraph_properties.append(indentation)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        paragraph_properties.append(spacing)
        level.append(paragraph_properties)

        run_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), BODY_FONT)
        fonts.set(qn("w:hAnsi"), BODY_FONT)
        fonts.set(qn("w:eastAsia"), CHINESE_FONT)
        run_properties.append(fonts)
        level.append(run_properties)
        numbering.append(abstract)

        number = OxmlElement("w:num")
        number.set(qn("w:numId"), str(num_id))
        abstract_reference = OxmlElement("w:abstractNumId")
        abstract_reference.set(qn("w:val"), str(abstract_id))
        number.append(abstract_reference)
        numbering.append(number)

    create_numbering(81, 81, "bullet")
    create_numbering(82, 82, "decimal")
    return 81, 82


def apply_numbering(paragraph, num_id: int) -> None:
    """把段落关联到指定的真实编号定义。"""
    properties = paragraph._p.get_or_add_pPr()
    number_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    number_properties.extend([level, number])
    properties.append(number_properties)


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    """把 Markdown 表格转换为具有固定列宽的 Word 表格。"""
    column_count = len(rows[0])
    normalized = [
        row[:column_count] + [""] * max(0, column_count - len(row))
        for row in rows
    ]
    table = document.add_table(rows=len(normalized), cols=column_count)
    widths = choose_column_widths(normalized)
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])

    for row_index, row in enumerate(normalized):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index == 0:
                set_cell_shading(cell, COLOR_TABLE_HEADER)
            elif row_index % 2 == 0:
                set_cell_shading(cell, COLOR_TABLE_ALT)

            paragraph = cell.paragraphs[0]
            set_paragraph_spacing(paragraph, after=0, line_spacing=1.15)
            if column_index in (1, 2) and len(value) < 18:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_runs(paragraph, value, default_size=9.5)
            for run in paragraph.runs:
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(COLOR_DARK_BLUE)

    following = document.add_paragraph()
    following.paragraph_format.space_after = Pt(2)


def add_code_block(document: Document, lines: list[str]) -> None:
    """添加带浅灰底色的代码或公式块。"""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    set_paragraph_spacing(paragraph, before=2, after=8, line_spacing=1.0)
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), COLOR_CODE_FILL)
    properties.append(shading)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(
            run,
            name=CODE_FONT,
            chinese_name=CHINESE_FONT,
            size=8.7,
            color="273142",
        )


def add_cover(document: Document, title: str) -> None:
    """添加技术参考手册封面。"""
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(72)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(kicker, after=14)
    run = kicker.add_run("技术接口参考手册")
    set_run_font(run, size=11, color=COLOR_BLUE, bold=True)

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title_paragraph, after=12, line_spacing=1.15)
    run = title_paragraph.add_run(title)
    set_run_font(run, size=25, color=COLOR_NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(subtitle, after=30)
    run = subtitle.add_run("数据结构、成员含义、函数作用、输入输出与调用示例")
    set_run_font(run, size=13, color=COLOR_MUTED)

    rule = document.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(rule, after=28)
    properties = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), COLOR_BLUE)
    borders.append(bottom)
    properties.append(borders)

    summary = document.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(summary, after=8, line_spacing=1.25)
    run = summary.add_run(
        "面向水面舰船、潜艇等目标的静电场与轴频电场仿真调用说明"
    )
    set_run_font(run, size=10.5, color=COLOR_MUTED)

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(note, after=0)
    run = note.add_run("对应类：TargetElectricFieldModel")
    set_run_font(run, name=CODE_FONT, chinese_name=CHINESE_FONT, size=10, color=COLOR_DARK_BLUE)

    document.add_page_break()


def configure_document(document: Document) -> tuple[int, int]:
    """配置页面、样式、页眉页脚和列表编号。"""
    section = document.sections[0]
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.25

    configure_style(
        document,
        "Heading 1",
        size=16,
        color=COLOR_BLUE,
        bold=True,
        before=18,
        after=10,
        keep_with_next=True,
    )
    configure_style(
        document,
        "Heading 2",
        size=13,
        color=COLOR_BLUE,
        bold=True,
        before=14,
        after=7,
        keep_with_next=True,
    )
    configure_style(
        document,
        "Heading 3",
        size=12,
        color=COLOR_DARK_BLUE,
        bold=True,
        before=10,
        after=5,
        keep_with_next=True,
    )

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(header, after=0, line_spacing=1.0)
    run = header.add_run("静电场与轴频电场仿真｜接口参考")
    set_run_font(run, size=8.5, color=COLOR_MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(footer, after=0, line_spacing=1.0)
    run = footer.add_run("第 ")
    set_run_font(run, size=9, color=COLOR_MUTED)
    add_page_field(footer)
    run = footer.add_run(" 页")
    set_run_font(run, size=9, color=COLOR_MUTED)

    return add_numbering_definitions(document)


def parse_table_line(line: str) -> list[str]:
    """把一行 Markdown 表格文本拆成单元格。"""
    return [part.strip() for part in line.strip().strip("|").split("|")]


def is_separator_row(line: str) -> bool:
    """判断是否为 Markdown 表格的分隔线。"""
    cells = parse_table_line(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def build_document(markdown_path: Path, output_path: Path) -> None:
    """读取 Markdown 并生成最终 Word 文档。"""
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    title = lines[0].removeprefix("# ").strip()

    document = Document()
    bullet_num_id, decimal_num_id = configure_document(document)
    add_cover(document, title)

    index = 1
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line.rstrip())
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_separator_row(lines[index + 1]):
            table_rows = [parse_table_line(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append(parse_table_line(lines[index]))
                index += 1
            add_markdown_table(document, table_rows)
            continue

        if stripped.startswith("#### "):
            paragraph = document.add_paragraph(style="Heading 3")
            add_inline_runs(paragraph, stripped[5:], default_size=12)
        elif stripped.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline_runs(paragraph, stripped[4:], default_size=13)
        elif stripped.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline_runs(paragraph, stripped[3:], default_size=16)
        elif stripped.startswith("# "):
            # 首个一级标题已经用于封面，其余一级标题按章标题处理。
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline_runs(paragraph, stripped[2:], default_size=16)
        elif stripped.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.right_indent = Inches(0.18)
            set_paragraph_spacing(paragraph, before=2, after=10, line_spacing=1.2)
            properties = paragraph._p.get_or_add_pPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "EEF4FA")
            properties.append(shading)
            add_inline_runs(paragraph, stripped[2:], default_size=10.5)
        elif stripped.startswith("- "):
            paragraph = document.add_paragraph()
            apply_numbering(paragraph, bullet_num_id)
            add_inline_runs(paragraph, stripped[2:])
        elif re.match(r"^\d+\.\s+", stripped):
            paragraph = document.add_paragraph()
            apply_numbering(paragraph, decimal_num_id)
            add_inline_runs(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
        else:
            paragraph = document.add_paragraph()
            set_paragraph_spacing(paragraph)
            add_inline_runs(paragraph, stripped)

        index += 1

    if in_code and code_lines:
        add_code_block(document, code_lines)

    # 把文档元数据设为通用值，避免带入本机用户信息。
    document.core_properties.title = title
    document.core_properties.subject = "静电场与轴频电场仿真接口说明"
    document.core_properties.author = "仿真工程文档"
    document.core_properties.keywords = "静电场, 轴频电场, 仿真, 数据结构, 函数接口"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    """处理命令行参数并生成文档。"""
    if len(sys.argv) != 3:
        raise SystemExit(
            "用法：build_electric_field_api_doc.py <输入Markdown> <输出DOCX>"
        )
    build_document(Path(sys.argv[1]), Path(sys.argv[2]))


if __name__ == "__main__":
    main()
