"""生成《舰艇磁场特性仿真原理与讲解指南》Word 文档。"""

from __future__ import annotations

import math
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


# 工程路径、技能辅助脚本和最终输出路径。
PROJECT_ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = PROJECT_ROOT / "docs" / "magnetic_field_word_assets"
OUTPUT_DOCX = PROJECT_ROOT / "docs" / "舰艇磁场特性仿真原理讲解（纯原理版）.docx"
SPATIAL_CSV = PROJECT_ROOT / "build" / "validation_magnetic.csv"
TIME_CSV = PROJECT_ROOT / "build" / "validation_magnetic_time_series.csv"
DOCUMENT_SKILL_ROOT = Path(
    "C:/Users/Administrator/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.727.11326/skills/documents"
)
sys.path.insert(0, str(DOCUMENT_SKILL_ROOT / "scripts"))
from table_geometry import apply_table_geometry  # noqa: E402


# compact_reference_guide 预设的精确设计标记。
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
NAVY = "0B2545"
HEADING_BLUE = "2E74B5"
HEADING_DARK = "1F4D78"
MUTED = "607786"
TABLE_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
LIGHT_BLUE = "EAF1F7"
CAUTION_GOLD = "7A5A00"
WHITE = "FFFFFF"
BODY_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"
MATH_FONT = "Cambria Math"


def rgb(hex_value: str) -> RGBColor:
    """把十六进制颜色转换为 python-docx 颜色。"""

    return RGBColor.from_string(hex_value)


def set_run_font(
    run,
    *,
    name: str = BODY_FONT,
    east_asia: str = CJK_FONT,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    """统一设置中西文字体、字号和强调样式。"""

    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_fill(cell, fill: str) -> None:
    """设置表格单元格底色。"""

    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_paragraph_shading(paragraph, fill: str) -> None:
    """设置段落背景色。"""

    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_paragraph_left_border(paragraph, color: str, size: int = 18, space: int = 8) -> None:
    """为提示段落增加左侧强调线。"""

    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)


def set_paragraph_bottom_border(paragraph, color: str, size: int = 6) -> None:
    """为页眉段落增加轻量底边线。"""

    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)


def keep_with_next(paragraph) -> None:
    """防止标题或图注与后续对象分离。"""

    paragraph.paragraph_format.keep_with_next = True


def keep_together(paragraph) -> None:
    """尽量避免单个段落跨页拆分。"""

    paragraph.paragraph_format.keep_together = True


def set_repeat_table_header(row) -> None:
    """把表格首行标记为跨页重复表头。"""

    properties = row._tr.get_or_add_trPr()
    header = properties.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        properties.append(header)


def configure_styles(document: Document) -> None:
    """按 compact_reference_guide 预设配置文档样式。"""

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(NAVY)
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        "Heading 1": (16, HEADING_BLUE, 18, 10),
        "Heading 2": (13, HEADING_BLUE, 14, 7),
        "Heading 3": (12, HEADING_DARK, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Figure Caption" not in styles:
        caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Figure Caption"]
    caption.font.name = BODY_FONT
    caption.font.size = Pt(9.5)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.keep_together = True

    if "Equation Block" not in styles:
        equation = styles.add_style("Equation Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        equation = styles["Equation Block"]
    equation.font.name = MATH_FONT
    equation.font.size = Pt(12)
    equation.font.color.rgb = rgb(NAVY)
    equation._element.rPr.rFonts.set(qn("w:ascii"), MATH_FONT)
    equation._element.rPr.rFonts.set(qn("w:hAnsi"), MATH_FONT)
    equation._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.space_before = Pt(4)
    equation.paragraph_format.space_after = Pt(7)
    equation.paragraph_format.line_spacing = 1.15
    equation.paragraph_format.keep_together = True

    if "Callout" not in styles:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Callout"]
    callout.font.name = BODY_FONT
    callout.font.size = Pt(10.5)
    callout.font.color.rgb = rgb(NAVY)
    callout._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.10)
    callout.paragraph_format.space_before = Pt(5)
    callout.paragraph_format.space_after = Pt(8)
    callout.paragraph_format.line_spacing = 1.2
    callout.paragraph_format.keep_together = True

    if "Reference" not in styles:
        reference = styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
    else:
        reference = styles["Reference"]
    reference.font.name = BODY_FONT
    reference.font.size = Pt(9.5)
    reference.font.color.rgb = rgb(NAVY)
    reference._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    reference._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    reference._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    reference.paragraph_format.left_indent = Inches(0.28)
    reference.paragraph_format.first_line_indent = Inches(-0.28)
    reference.paragraph_format.space_before = Pt(0)
    reference.paragraph_format.space_after = Pt(5)
    reference.paragraph_format.line_spacing = 1.15
    reference.paragraph_format.keep_together = True


def configure_page(document: Document) -> None:
    """设置 Letter 页面、页边距、页眉和页脚距离。"""

    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        section.different_first_page_header_footer = True


def add_page_field(paragraph) -> None:
    """在页脚中插入 Word 页码域。"""

    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)
    end_run = paragraph.add_run(" 页")
    set_run_font(end_run, size=9, color=MUTED)


def configure_header_footer(document: Document) -> None:
    """设置安静的运行页眉和右对齐页码。"""

    section = document.sections[0]
    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left_run = paragraph.add_run("舰艇磁场特性仿真原理与讲解指南")
    set_run_font(left_run, size=8.5, color=MUTED, bold=True)
    paragraph.add_run("\t")
    right_run = paragraph.add_run("技术讲解版")
    set_run_font(right_run, size=8.5, color=MUTED)
    set_paragraph_bottom_border(paragraph, "D8E2E8", size=5)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.clear()
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(footer_paragraph)


def add_numbering_definitions(document: Document) -> tuple[int, int]:
    """创建符合预设缩进的真实项目符号和编号定义。"""

    numbering = document.part.numbering_part.element
    existing_abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
    ]
    existing_num_ids = [
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
    ]
    next_abstract = max(existing_abstract_ids, default=0) + 1
    next_num = max(existing_num_ids, default=0) + 1

    def create_definition(abstract_id: int, num_id: int, fmt: str, text: str, font_name: str) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        level.append(num_fmt)
        level_text = OxmlElement("w:lvlText")
        level_text.set(qn("w:val"), text)
        level.append(level_text)
        justification = OxmlElement("w:lvlJc")
        justification.set(qn("w:val"), "left")
        level.append(justification)
        paragraph_properties = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        paragraph_properties.append(tabs)
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), "540")
        indent.set(qn("w:hanging"), "270")
        paragraph_properties.append(indent)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        paragraph_properties.append(spacing)
        level.append(paragraph_properties)
        run_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), font_name)
        fonts.set(qn("w:hAnsi"), font_name)
        run_properties.append(fonts)
        level.append(run_properties)
        abstract.append(level)
        numbering.append(abstract)

        number = OxmlElement("w:num")
        number.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        number.append(abstract_ref)
        numbering.append(number)

    create_definition(next_abstract, next_num, "bullet", "•", "Symbol")
    create_definition(next_abstract + 1, next_num + 1, "decimal", "%1.", BODY_FONT)
    return next_num, next_num + 1


def apply_numbering(paragraph, num_id: int) -> None:
    """把段落绑定到指定的真实编号定义。"""

    properties = paragraph._p.get_or_add_pPr()
    num_properties = properties.find(qn("w:numPr"))
    if num_properties is None:
        num_properties = OxmlElement("w:numPr")
        properties.append(num_properties)
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_properties.append(level)
    num_properties.append(number)


def add_bullet(document: Document, text: str, bullet_num_id: int) -> None:
    """添加真实项目符号段落。"""

    paragraph = document.add_paragraph()
    apply_numbering(paragraph, bullet_num_id)
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color=NAVY)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    keep_together(paragraph)


def add_numbered(document: Document, text: str, decimal_num_id: int) -> None:
    """添加真实编号段落。"""

    paragraph = document.add_paragraph()
    apply_numbering(paragraph, decimal_num_id)
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color=NAVY)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.25
    keep_together(paragraph)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    """添加真实标题样式。"""

    paragraph = document.add_heading(text, level=level)
    keep_with_next(paragraph)
    keep_together(paragraph)


def add_body(
    document: Document,
    text: str,
    *,
    bold_lead: str | None = None,
    justified: bool = False,
) -> None:
    """添加正文段落，可选择加粗开头标签。"""

    paragraph = document.add_paragraph()
    if justified:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        lead_run = paragraph.add_run(bold_lead)
        set_run_font(lead_run, size=11, color=NAVY, bold=True)
        body_run = paragraph.add_run(text[len(bold_lead):])
        set_run_font(body_run, size=11, color=NAVY)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11, color=NAVY)
    keep_together(paragraph)


def add_reference(document: Document, number: int, text: str) -> None:
    """添加带悬挂缩进的编号参考文献。"""

    paragraph = document.add_paragraph(style="Reference")
    number_run = paragraph.add_run(f"[{number}] ")
    set_run_font(number_run, size=9.5, color=NAVY, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=9.5, color=NAVY)


def add_callout(document: Document, label: str, text: str, accent: str = HEADING_BLUE) -> None:
    """添加带左侧强调线的讲解提示。"""

    paragraph = document.add_paragraph(style="Callout")
    set_paragraph_shading(paragraph, LIGHT_FILL)
    set_paragraph_left_border(paragraph, accent)
    label_run = paragraph.add_run(f"{label}：")
    set_run_font(label_run, size=10.5, color=accent, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=10.5, color=NAVY)


def add_equation(document: Document, equation: str, explanation: str | None = None) -> None:
    """添加公式块及可选的公式解释。"""

    paragraph = document.add_paragraph(style="Equation Block")
    set_paragraph_shading(paragraph, LIGHT_BLUE)
    set_paragraph_left_border(paragraph, HEADING_BLUE, size=10, space=6)
    run = paragraph.add_run(equation)
    set_run_font(
        run,
        name=MATH_FONT,
        east_asia=CJK_FONT,
        size=12,
        color=NAVY,
        bold=True,
    )
    if explanation:
        explanation_run = paragraph.add_run(f"\n{explanation}")
        set_run_font(explanation_run, size=10, color=MUTED)


def add_figure(
    document: Document,
    filename: str,
    caption: str,
    alt_text: str,
    width_inches: float = 6.25,
) -> None:
    """插入图片、替代文本和图注。"""

    image_path = ASSET_DIR / filename
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(image_path), width=Inches(width_inches))
    inline_shape._inline.docPr.set("descr", alt_text)
    inline_shape._inline.docPr.set("title", caption)
    caption_paragraph = document.add_paragraph(caption, style="Figure Caption")
    keep_together(caption_paragraph)


def add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    widths_dxa: list[int],
    *,
    font_size: float = 9.4,
) -> None:
    """创建具有精确几何、重复表头和统一间距的表格。"""

    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError("表格列宽之和必须为 9360 DXA")
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Table Grid"
    header_row = table.rows[0]
    set_repeat_table_header(header_row)

    for index, header_text in enumerate(headers):
        cell = header_row.cells[index]
        set_cell_fill(cell, TABLE_FILL)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(header_text)
        set_run_font(run, size=font_size, color=NAVY, bold=True)

    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        for column_index, value in enumerate(row_values):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 1:
                set_cell_fill(cell, "F8FAFC")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.12
            if column_index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(value))
            set_run_font(run, size=font_size, color=NAVY)

    apply_table_geometry(
        table,
        widths_dxa,
        table_width_dxa=CONTENT_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa=CELL_MARGINS_DXA,
    )
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_page_break(document: Document) -> None:
    """添加明确分页，保持讲解章节完整。"""

    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def build_cover(document: Document) -> None:
    """创建 editorial_cover 风格封面。"""

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(55)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    run = kicker.add_run("技术原理 · 公式分析 · 仿真示例 · 讲解指南")
    set_run_font(run, size=10.5, color=CAUTION_GOLD, bold=True)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title_run = title.add_run("舰艇磁场特性仿真")
    set_run_font(title_run, size=30, color=NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    subtitle_run = subtitle.add_run("从地磁感应与剩磁，到空间分布和运动时序")
    set_run_font(subtitle_run, size=15, color=HEADING_DARK)

    formula = document.add_paragraph()
    formula.alignment = WD_ALIGN_PARAGRAPH.CENTER
    formula.paragraph_format.space_after = Pt(36)
    formula_run = formula.add_run("B_anom = B_static + B_induced")
    set_run_font(formula_run, name=MATH_FONT, size=15, color=HEADING_BLUE, bold=True)

    summary = document.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.left_indent = Inches(0.55)
    summary.paragraph_format.right_indent = Inches(0.55)
    summary.paragraph_format.space_after = Pt(58)
    summary_run = summary.add_run(
        "面向非专业听众的通俗讲解材料，重点说明物理来源、等效方法、公式含义、图像解读与模型边界。"
    )
    set_run_font(summary_run, size=11, color=MUTED)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(5)
    meta_run = meta.add_run("原理框架：等效磁性椭球 + 单磁偶极子")
    set_run_font(meta_run, size=10, color=MUTED)
    date_paragraph = document.add_paragraph()
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_paragraph.add_run("2026 年 7 月")
    set_run_font(date_run, size=10, color=MUTED)
    add_page_break(document)


def build_document() -> Path:
    """构建完整 Word 文档并返回输出路径。"""

    document = Document()
    document.core_properties.title = "舰艇磁场特性仿真原理讲解（纯原理版）"
    document.core_properties.subject = "舰艇与潜艇目标静磁场、感应磁场和运动时序仿真"
    document.core_properties.author = "seamine 项目组"
    document.core_properties.last_modified_by = "seamine 项目组"
    document.core_properties.keywords = "舰艇磁场, 地磁感应, 剩磁, 磁偶极子, 水雷目标"
    document.core_properties.comments = "仅讲解舰艇磁场特性仿真的物理原理、公式和图像，不包含软件实现说明。"
    configure_styles(document)
    configure_page(document)
    configure_header_footer(document)
    bullet_num_id, decimal_num_id = add_numbering_definitions(document)

    spatial_data = pd.read_csv(SPATIAL_CSV)
    time_data = pd.read_csv(TIME_CSV)
    spatial_peak_row = spatial_data.loc[spatial_data["total_magnitude_nt"].idxmax()]
    time_peak_row = time_data.loc[time_data["total_magnitude_nt"].idxmax()]
    minimum_distance_row = time_data.loc[time_data["distance_m"].idxmin()]

    build_cover(document)

    # 阅读导航和一分钟说明。
    add_heading(document, "阅读导航：先讲什么，再讲公式", level=1)
    add_callout(
        document,
        "一句话说明",
        "舰艇相当于一个会被地磁“磁化”、同时又保留自身剩磁的大磁体；模型把复杂舰体简化为磁性椭球，再计算它在传感器处产生的异常磁场。",
    )
    add_body(
        document,
        "如果只用 1 分钟介绍，先讲“磁场从哪里来”，再讲“为什么舰艇靠近时信号会变强”，最后展示热力图和时间曲线。公式用于回答追问，不必一开始全部展开。"
    )
    add_heading(document, "推荐讲解顺序", level=2)
    for item in [
        "先区分背景地磁、目标静磁场、感应磁场和综合目标异常磁场。",
        "再用等效椭球解释长度、宽度、高度为什么会影响磁化程度。",
        "用磁偶极子公式说明场强随距离约按 1/r³ 快速衰减。",
        "最后展示空间热力图和固定传感器处的时间曲线。",
    ]:
        add_numbered(document, item, decimal_num_id)
    add_heading(document, "本说明覆盖的内容", level=2)
    for item in [
        "模型输入、坐标和单位",
        "静态磁矩、感应磁矩和退磁因子",
        "空间分布与运动时序的计算流程",
        "空间分布图、三分量图和运动时序图的物理含义",
        "模型边界、常见问题和五分钟讲解提纲",
    ]:
        add_bullet(document, item, bullet_num_id)
    add_page_break(document)

    # 第一章：基本概念。
    add_heading(document, "1. 模拟的到底是什么磁场", level=1)
    add_body(
        document,
        "地球本身存在一个背景地磁场。钢制舰艇或潜艇进入该环境后，船体会被磁化并产生感应磁矩；与此同时，材料加工、建造历史和长期服役还会留下剩磁。水雷磁传感器关注的通常不是背景地磁本身，而是目标经过时叠加在背景上的变化量。[5]"
    )
    add_figure(
        document,
        "principle_schematic.png",
        "图 1  舰艇磁场形成机理与坐标关系（本文绘制，原理依据[1][2][5]）",
        "舰艇等效磁性椭球、背景地磁、剩磁磁矩、感应磁矩和固定磁传感器的关系示意图。",
        width_inches=6.3,
    )
    add_equation(
        document,
        "B_anom = B_static + B_induced",
        "综合目标异常磁场 = 目标剩磁产生的静磁场 + 地磁磁化产生的感应磁场。",
    )
    add_callout(
        document,
        "重要区别",
        "本文所说的“综合目标异常磁场”不包含背景地磁本身。若要表示磁传感器的绝对读数，应计算 B_abs = B_g + B_anom。",
        accent=CAUTION_GOLD,
    )
    add_page_break(document)

    # 第二章：坐标和输入参数。
    add_heading(document, "2. 坐标、单位与输入参数", level=1)
    add_body(
        document,
        "模型采用东-北-上（ENU）坐标：x 向东、y 向北、z 向上，水下位置的 z 为负数。目标坐标系中 x 轴沿舰首方向，航向、俯仰和横滚决定目标坐标系如何旋转到全局坐标系。地磁强度、偏角、倾角以及 X、Y、Z 分量的定义参见 NOAA/WMM 和 IGRF 资料。[3][4]"
    )
    add_table(
        document,
        ["参数", "单位", "通俗含义", "主要影响"],
        [
            ["L、W、H", "m", "目标长、宽、高", "体积、退磁因子和磁矩大小"],
            ["r₀", "m", "零时刻目标中心位置", "目标与传感器的相对距离"],
            ["v", "m/s", "沿舰首方向的航速", "时间曲线宽度和通过时刻"],
            ["ψ、θ、φ", "°", "航向、俯仰、横滚", "磁矩方向和运动方向"],
            ["B_g", "nT", "东、北、上地磁三分量", "感应磁矩大小和方向"],
            ["μ_r", "无量纲", "材料被磁化的容易程度", "感应磁矩强弱"],
            ["η", "0～1", "等效铁磁材料体积占比", "静态和感应磁矩均近似成比例变化"],
            ["M_r", "A/m", "目标自身剩磁强度三分量", "静磁场大小和方向"],
            ["r_min", "m", "偶极模型最小适用距离", "避免进入单偶极模型不适用的近场区域"],
        ],
        [2100, 900, 2850, 3510],
        font_size=8.8,
    )
    add_heading(document, "地磁三分量怎样填写", level=2)
    add_equation(
        document,
        "B_E = F cos(I) sin(D)    B_N = F cos(I) cos(D)    B_U = -F sin(I)",
        "F 为总地磁强度，D 为从正北向东为正的磁偏角，I 为向下为正的磁倾角。",
    )
    add_callout(
        document,
        "单位提醒",
        "输入单位必须是 nT。若仪器给出 μT，需要乘以 1000；若给出北-东-下分量 X、Y、Z，应转换为 {东 Y，北 X，上 -Z}。",
    )
    add_heading(document, "核心术语先解释", level=2)
    add_table(
        document,
        ["术语", "通俗释义", "在本文中的作用"],
        [
            ["背景地磁场", "地球在目标所在位置原本存在的磁场", "是舰体产生感应磁化的外部条件"],
            ["剩磁", "外部磁场撤去后，铁磁材料仍保留的磁性", "形成目标静态磁矩和静磁场"],
            ["感应磁化", "舰体在当前地磁作用下被临时磁化", "形成随地磁方向变化的感应磁场"],
            ["磁矩", "描述一个磁体强弱和主方向的矢量", "连接舰体磁化状态与外部磁场分布"],
            ["退磁因子", "描述物体形状对自身磁化产生反作用的系数", "决定不同主轴方向被磁化的难易"],
            ["目标异常磁场", "目标引起的、叠加在背景地磁上的变化量", "是水雷或磁传感器重点关注的信号"],
        ],
        [1900, 4100, 3360],
        font_size=9.0,
    )
    add_page_break(document)

    # 第三章：等效椭球和静磁场。
    add_heading(document, "3. 为什么把舰艇等效成磁性椭球", level=1)
    add_body(
        document,
        "真实舰体包含外壳、骨架、轴系和设备，磁性结构十分复杂。为了快速仿真，当前模型用长、宽、高相同的三轴椭球代表整体形状。这个简化保留了最重要的三件事：总体尺寸、三个主轴方向，以及不同方向上的退磁差异。椭球退磁因子的经典推导与主轴性质见 Osborn。[1]"
    )
    add_equation(
        document,
        "a = L/2,    b = W/2,    c = H/2",
        "a、b、c 是等效椭球三个半轴。",
    )
    add_equation(
        document,
        "V_e = 4πabc/3 = πLWH/6,    V_m = ηV_e",
        "V_e 为几何体积，η 为磁性材料等效体积占比，V_m 为参与磁化的等效磁性体积。",
    )
    example_volume = math.pi / 6.0 * 100.0 * 15.0 * 10.0
    magnetic_volume = example_volume * 0.035
    add_callout(
        document,
        "示例",
        f"长 100 m、宽 15 m、高 10 m 的等效椭球体积约为 {example_volume:,.1f} m³；当 η=0.035 时，等效磁性体积约为 {magnetic_volume:,.1f} m³。",
    )
    add_heading(document, "剩磁如何变成静态磁矩", level=2)
    add_equation(
        document,
        "m_r(body) = V_m M_r,    m_r(global) = R m_r(body)",
        "先在目标坐标系中用体积乘剩磁强度，再通过姿态旋转矩阵 R 转到全局坐标系。",
    )
    add_body(
        document,
        "剩磁强度是带方向的三分量。即使两艘舰艇尺寸相同，只要磁化历史或航向不同，传感器观测到的 Bx、By、Bz 也可能明显不同。"
    )
    add_page_break(document)

    # 第四章：感应磁场和退磁因子。
    add_heading(document, "4. 地磁如何在舰体中产生感应磁场", level=1)
    add_body(
        document,
        "可以把船体想象成一块放进地磁场中的软磁材料。地磁越强、材料越容易被磁化、参与磁化的材料越多，感应磁矩通常越大。但船体自身会产生一个反向的退磁场，因此不能简单地只用“磁导率乘地磁场”。"
    )
    add_heading(document, "4.1 退磁因子", level=2)
    add_equation(
        document,
        "N_i = (abc/2) ∫₀∞ ds / [(s+a_i²) √((s+a²)(s+b²)(s+c²))]",
        "i 分别对应纵向、横向、垂向；一般可通过数值积分求得。[1]",
    )
    add_equation(
        document,
        "N_x + N_y + N_z = 1",
        "细长目标的纵向退磁因子通常较小，因此更容易沿纵向被磁化。",
    )
    add_heading(document, "4.2 感应磁矩", level=2)
    add_equation(
        document,
        "χ = μ_r - 1,    H_g = B_g / μ₀",
        "χ 是磁化率；地磁 B 需先换算为磁场强度 H。",
    )
    add_equation(
        document,
        "m_i,k = V_m χ H_g,k / (1 + N_k χ)",
        "分母中的退磁项会限制高磁导率材料的有效磁化程度。",
    )
    add_body(
        document,
        "计算前，模型先把全局地磁三分量旋转到目标坐标系；三个主轴分别计算感应磁矩后，再旋转回东-北-上坐标系。这样航向变化会自然改变最终的感应磁场三分量。"
    )
    add_callout(
        document,
        "讲解比喻",
        "地磁像外部推力，材料磁化率像“容易被推动的程度”，退磁因子像舰体自身的反作用。细长方向的反作用较弱，所以纵向磁化更明显。",
    )
    add_page_break(document)

    # 第五章：外部磁场公式与衰减。
    add_heading(document, "5. 等效磁矩怎样变成传感器处的磁场", level=1)
    add_body(
        document,
        "得到静态磁矩和感应磁矩后，模型把它们分别视为位于目标中心的磁偶极子。对于目标外部、且距离目标中心足够远的观测点，磁偶极子公式能够快速给出三分量磁场。[2]"
    )
    add_equation(
        document,
        "B(r) = μ₀/(4πr³) · [3(m·r̂)r̂ - m]",
        "m 为磁矩，r 为目标中心到传感器的位移，r̂ 为其单位方向。",
    )
    add_equation(
        document,
        "B_static = Dipole(m_r),    B_induced = Dipole(m_i)",
        "静磁场与感应磁场分别计算，最后按矢量相加。",
    )
    add_figure(
        document,
        "parameter_sensitivity.png",
        "图 2  距离与磁性材料比例的归一化敏感性（本文示例）",
        "左侧为磁偶极场随距离三次方反比衰减，右侧为等效磁矩随磁性材料比例线性变化的趋势图。",
        width_inches=6.3,
    )
    add_callout(
        document,
        "为什么最近点最重要",
        "场强约按 1/r³ 衰减。距离增大到 2 倍时，理想偶极场约降为 1/8；因此传感器深度、横向偏距和航迹误差都会显著改变峰值。",
        accent=CAUTION_GOLD,
    )
    add_page_break(document)

    # 第六章：运动时序。
    add_heading(document, "6. 运动目标磁场时序如何生成", level=1)
    add_body(
        document,
        "当前运动模型采用固定航向、固定姿态和恒定航速。目标磁矩方向保持不变，时间变化主要来自目标与固定磁传感器之间的距离和观察方向不断变化。"
    )
    add_equation(
        document,
        "r_c(t) = r₀ + R · [v t, 0, 0]ᵀ",
        "r₀ 为零时刻中心位置，v 为沿舰首方向的航速。",
    )
    add_equation(
        document,
        "d(t) = r_sensor - r_c(t)",
        "每个时刻重新计算目标到固定传感器的位移和距离。",
    )
    add_equation(
        document,
        "N = floor(T f_s) + 1,    t_k = t_start + k/f_s",
        "T 为持续时间，f_s 为采样率，结果包含起始点和最后一个完整采样点。",
    )
    add_heading(document, "项目中的演示场景", level=2)
    add_table(
        document,
        ["项目", "设置"],
        [
            ["目标初始位置", "(-300, 40, -5) m"],
            ["航速与航向", "8 m/s，航向 0°，沿东向匀速航行"],
            ["固定传感器", "(0, 0, -30) m"],
            ["时长与采样率", "75 s，20 Hz，共 1501 点"],
            ["最近距离", f"{minimum_distance_row['distance_m']:.2f} m，发生在 {minimum_distance_row['time_s']:.2f} s"],
            ["综合场峰值", f"{time_peak_row['total_magnitude_nt']:.2f} nT，发生在 {time_peak_row['time_s']:.2f} s"],
        ],
        [2350, 7010],
        font_size=9.5,
    )
    add_callout(
        document,
        "观察",
        "磁场峰值不一定与几何最近点完全重合，因为磁偶极场还取决于磁矩方向和目标-传感器连线方向。",
    )
    add_page_break(document)

    # 第七章：完整流程。
    add_heading(document, "7. 从参数到图像的完整流程", level=1)
    add_body(
        document,
        "模型只在配置目标时计算一次等效静态磁矩与感应磁矩。之后，无论生成空间网格还是时间序列，都重复使用这两个磁矩，从而避免在每个采样点重新积分退磁因子。"
    )
    add_figure(
        document,
        "simulation_flow.png",
        "图 3  磁场特性仿真的完整原理流程（本文整理）",
        "输入参数经过校验、等效椭球、静态和感应磁矩计算后，分为空间网格与运动时序两条输出路径。",
        width_inches=6.15,
    )
    add_heading(document, "计算结果的三层结构", level=2)
    for item in [
        "分量层：Bx、By、Bz，保留正负号，表示方向。",
        "模值层：|B|，始终非负，表示强度。",
        "业务层：空间分布用于观察影响区域，时间序列用于观察目标通过过程。",
    ]:
        add_bullet(document, item, bullet_num_id)
    add_page_break(document)

    # 第八章：空间结果。
    add_heading(document, "8. 空间分布结果怎样看", level=1)
    add_body(
        document,
        "空间网格模式在零时刻遍历指定的 x-y-z 观测点。当前示例在水下 30 m 平面生成 101×61=6161 个点，适合绘制热力图、等值线和三维曲面。"
    )
    add_figure(
        document,
        "spatial_heatmaps.png",
        "图 4  静磁场、感应磁场和综合异常磁场空间分布（本文仿真示例）",
        "水下三十米平面上的静磁场强度、感应磁场强度和综合目标异常磁场强度热力图。",
        width_inches=6.45,
    )
    add_callout(
        document,
        "图像解读",
        f"本示例综合异常磁场最大值约为 {spatial_peak_row['total_magnitude_nt']:.1f} nT，位置约为 ({spatial_peak_row['x_m']:.0f}, {spatial_peak_row['y_m']:.0f}, {spatial_peak_row['z_m']:.0f}) m。高场区集中在目标附近，并沿目标磁矩方向呈现明显方向性。",
    )
    add_body(
        document,
        "静磁场图主要反映剩磁参数；感应磁场图主要反映地磁、几何退磁因子、相对磁导率和材料比例。综合图是两者矢量相加后的结果，不能简单把两个模值直接相加。"
    )
    add_page_break(document)

    # 第九章：三分量结果。
    add_heading(document, "9. 为什么必须看 Bx、By、Bz 三分量", level=1)
    add_body(
        document,
        "磁场是矢量。模值只告诉我们“有多强”，三分量还告诉我们“朝哪个方向”。水雷磁传感器的敏感轴、目标航向和地磁方向都会使不同分量表现出不同的正负变化。"
    )
    add_figure(
        document,
        "induced_components.png",
        "图 5  地磁感应磁场东、北、上三分量分布（本文仿真示例）",
        "感应磁场 Bx、By、Bz 分量的发散色带热力图，暖色和冷色分别代表正负方向。",
        width_inches=6.45,
    )
    add_heading(document, "分量图的讲解方法", level=2)
    for item in [
        "暖色与冷色代表相反方向，不代表“好”和“坏”。",
        "零值附近的分界线表示该分量发生方向反转。",
        "三分量的峰值位置可能不同，原因是偶极场具有方向性。",
        "做频谱或波形分析时应使用带符号的分量，而不是只使用模值。",
    ]:
        add_bullet(document, item, bullet_num_id)
    add_page_break(document)

    # 第十章：时间结果。
    add_heading(document, "10. 运动时序结果怎样看", level=1)
    add_body(
        document,
        "固定传感器观测到的时间曲线，本质上是舰艇穿过其敏感区域时，对空间磁场分布的一条动态切片。曲线形状同时受航速、横向偏距、深度、磁矩方向和传感器位置影响。"
    )
    add_figure(
        document,
        "time_series.png",
        "图 6  运动目标综合磁场三分量与场强时间曲线（本文仿真示例）",
        "上图为综合目标异常磁场 Bx、By、Bz 随时间变化，下图为静磁场、感应磁场和综合场强随时间变化。",
        width_inches=6.35,
    )
    add_callout(
        document,
        "讲解重点",
        f"演示场景的综合场峰值约为 {time_peak_row['total_magnitude_nt']:.2f} nT，发生在 {time_peak_row['time_s']:.2f} s。分量曲线会过零或改变符号，而模值曲线始终非负。",
    )
    add_page_break(document)

    # 第十一章：只解释常见图像及其物理意义。
    add_heading(document, "11. 常见仿真图像及其物理意义", level=1)
    add_body(
        document,
        "磁场仿真的输出不只是一个数值。不同图像回答的问题不同：空间图回答“哪里强、哪里弱”，分量图回答“磁场朝哪个方向”，时间图回答“目标通过时信号如何变化”。"
    )
    add_table(
        document,
        ["图像类型", "主要观察内容", "适合回答的问题"],
        [
            ["磁场强度热力图", "某一深度平面上的 |B| 大小", "高场区位于哪里？影响范围有多大？"],
            ["三分量分布图", "Bx、By、Bz 的正负和空间变化", "哪个方向的磁场最明显？哪里发生方向反转？"],
            ["等值线图", "相同场强位置形成的轮廓", "给定阈值对应多大的探测区域？"],
            ["三维曲面图", "场强随两个空间坐标的起伏", "峰值、谷值和方向性是否明显？"],
            ["运动时间曲线", "固定传感器处磁场随时间变化", "目标何时接近？峰值何时出现？"],
            ["距离—场强关系图", "目标距离与异常场强的对应变化", "距离衰减是否符合约 1/r³ 的趋势？"],
            ["轨迹叠加图", "目标航迹与磁场区域的相对关系", "偏距、航向和深度怎样改变波形？"],
        ],
        [1800, 3900, 3660],
        font_size=9.0,
    )
    add_heading(document, "图像解读的推荐顺序", level=2)
    for item in [
        "先看综合场强模值，确认目标异常信号的大致范围和峰值。",
        "再看 Bx、By、Bz 三分量，判断信号方向、过零点和分量差异。",
        "把目标航迹或传感器位置叠加到图上，解释峰值与几何关系。",
        "最后对照静磁场和感应磁场，判断剩磁与地磁磁化各自的贡献。",
    ]:
        add_bullet(document, item, bullet_num_id)
    add_callout(
        document,
        "注意",
        "磁场模值始终非负，适合表示强弱；三分量保留正负号，适合分析方向。只看模值可能掩盖分量过零和方向反转。",
    )
    add_page_break(document)

    # 第十二章：参数敏感性。
    add_heading(document, "12. 参数变化会怎样影响结果", level=1)
    add_table(
        document,
        ["参数变化", "主要趋势", "解释时的注意事项"],
        [
            ["长度、宽度、高度增大", "等效体积增大，磁矩通常增大；退磁因子也会改变", "不是单纯按体积同比例变化"],
            ["地磁场增强或方向改变", "感应磁矩增强或旋转", "静磁场不受地磁输入影响"],
            ["相对磁导率增大", "感应磁矩增大后逐渐受退磁项限制", "高磁导率下不会无限线性增长"],
            ["材料比例 η 增大", "静态和感应磁矩近似线性增大", "需要通过结构资料或实测标定"],
            ["剩磁强度增大", "静磁场近似线性增大", "剩磁方向与大小同样重要"],
            ["传感器距离增大", "场强按约 1/r³ 快速下降", "最敏感、最容易造成数量级变化"],
            ["航向或姿态改变", "磁矩和分量方向改变", "模值与三分量都可能变化"],
            ["航速增大", "时间曲线变窄、通过更快", "空间分布本身不因匀速大小改变"],
        ],
        [2050, 3400, 3910],
        font_size=8.7,
    )
    add_heading(document, "建议的标定顺序", level=2)
    for item in [
        "先确认坐标、单位、目标位置、传感器位置和航迹。",
        "再使用目标尺寸和结构资料估计材料比例与相对磁导率。",
        "通过静态或低速实测数据标定剩磁三分量。",
        "通过不同航向实测数据校正感应磁化参数。",
        "最后使用独立航次验证峰值、波形和三分量方向。",
    ]:
        add_numbered(document, item, decimal_num_id)
    add_page_break(document)

    # 第十三章：假设与限制。
    add_heading(document, "13. 当前模型的假设、适用范围与边界", level=1)
    add_heading(document, "适合做什么", level=2)
    for item in [
        "水雷磁引信算法验证和阈值分析",
        "参数敏感性研究与合成数据生成",
        "目标通过传感器时的三分量波形演示",
        "不同尺寸、航向、深度和偏距的快速对比",
    ]:
        add_bullet(document, item, bullet_num_id)
    add_heading(document, "暂时没有模拟什么", level=2)
    for item in [
        "舰体内部复杂分区、多磁极结构和局部近场细节",
        "转弯、变速、姿态随时间变化和复杂航迹",
        "轴系旋转、电机和设备产生的交变磁场与谐波",
        "海底、海水磁导率变化和周边铁磁物体影响",
        "消磁线圈、主动补偿系统和实时磁滞过程",
    ]:
        add_bullet(document, item, bullet_num_id)
    add_callout(
        document,
        "模型边界",
        "当观测点接近或进入舰体尺寸范围时，单偶极子无法描述局部磁场细节。工程级近场预测应使用多偶极子、有限元结果或实测数据标定；舰船磁特征的多源组成和工程应用可参见 Holmes。[5]",
        accent=CAUTION_GOLD,
    )
    add_heading(document, "结果应该怎样表述", level=2)
    add_body(
        document,
        "建议称为“等效磁偶极子模型下的目标异常磁场预测”或“用于算法验证的合成磁场”。避免把未标定结果直接表述为真实舰艇的绝对磁场值。"
    )
    add_page_break(document)

    # 第十四章：讲解提纲。
    add_heading(document, "14. 可直接照着讲的五分钟提纲", level=1)
    talk_track = [
        "第一分钟：地球有背景地磁，钢制舰艇会被磁化，还会保留剩磁。水雷传感器看的是目标经过时引起的异常变化。",
        "第二分钟：把复杂舰艇简化成具有相同长、宽、高的磁性椭球。体积决定磁性材料规模，退磁因子描述不同方向磁化难易。",
        "第三分钟：剩磁形成静态磁矩，地磁通过磁化率和退磁因子形成感应磁矩。两者都旋转到东-北-上坐标。",
        "第四分钟：用磁偶极子公式计算传感器处 Bx、By、Bz。距离增大一倍，场强大约降到八分之一。",
        "第五分钟：空间网格生成热力图，运动采样生成时间曲线。当前结果是目标异常磁场，工程应用还需实测标定。",
    ]
    for item in talk_track:
        add_numbered(document, item, decimal_num_id)
    add_heading(document, "常见追问", level=2)
    add_table(
        document,
        ["问题", "简短回答"],
        [
            ["不设置地磁会怎样？", "模型使用默认地磁 {0, 30000, -40000} nT；显式设为零时，感应磁场为零。"],
            ["综合异常磁场包括地磁背景吗？", "不包括。它是静磁场与感应异常场的矢量和。"],
            ["为什么峰值不一定在最近点？", "偶极场不仅取决于距离，还取决于磁矩与观察方向。"],
            ["为什么要看三分量？", "三分量保留方向信息；模值只反映强度。"],
            ["结果能直接代表真实舰艇吗？", "不能直接代表。需要用实测数据标定材料比例、剩磁和感应参数。"],
        ],
        [2700, 6660],
        font_size=9.3,
    )
    add_page_break(document)

    # 参考文献：仅保留支撑物理原理和地磁定义的外部资料。
    add_heading(document, "参考文献", level=1)
    add_callout(
        document,
        "引用说明",
        "正文中的[1]—[5]用于支撑退磁因子、磁偶极子、地磁定义和舰船磁特征说明。图 1—图 6 均为本文绘制或根据示例参数生成，不直接复制外部文献图片。",
    )
    add_reference(
        document,
        1,
        "Osborn, J. A. Demagnetizing Factors of the General Ellipsoid[J]. "
        "Physical Review, 1945, 67(11–12): 351–357. "
        "DOI: 10.1103/PhysRev.67.351.",
    )
    add_reference(
        document,
        2,
        "Griffiths, D. J. Introduction to Electrodynamics[M]. 4th ed. "
        "Cambridge: Cambridge University Press, 2017. "
        "DOI: 10.1017/9781108333511.",
    )
    add_reference(
        document,
        3,
        "National Centers for Environmental Information (NCEI). "
        "World Magnetic Model (WMM)[EB/OL]. "
        "https://www.ncei.noaa.gov/products/world-magnetic-model "
        "(访问日期：2026-07-28).",
    )
    add_reference(
        document,
        4,
        "Alken, P., Thébault, E., Beggan, C. D., et al. "
        "International Geomagnetic Reference Field: the thirteenth generation[J]. "
        "Earth, Planets and Space, 2021, 73: 49. "
        "DOI: 10.1186/s40623-020-01288-x.",
    )
    add_reference(
        document,
        5,
        "Holmes, J. J. Exploitation of a Ship's Magnetic Field Signatures[M]. "
        "Springer Cham, 2006. DOI: 10.1007/978-3-031-01693-6.",
    )
    add_page_break(document)

    # 附录：纯物理符号和专有名词释义。
    add_heading(document, "附录 A：主要物理符号", level=1)
    add_table(
        document,
        ["符号", "含义", "单位或说明"],
        [
            ["L、W、H", "目标长、宽、高", "m"],
            ["a、b、c", "等效椭球三个半轴", "m，分别为 L/2、W/2、H/2"],
            ["η", "铁磁材料等效体积占比", "0～1"],
            ["μ_r", "相对磁导率", "无量纲"],
            ["χ", "磁化率", "无量纲，χ=μ_r−1"],
            ["N_x、N_y、N_z", "三个主轴方向的退磁因子", "无量纲，三者之和为 1"],
            ["B_g", "背景地磁场矢量", "T、μT 或 nT"],
            ["M_r", "剩磁强度矢量", "A/m"],
            ["m_r", "等效静态磁矩", "A·m²"],
            ["m_i", "等效感应磁矩", "A·m²"],
            ["B_static", "剩磁产生的静态异常磁场", "T、μT 或 nT"],
            ["B_induced", "地磁磁化产生的感应异常磁场", "T、μT 或 nT"],
            ["B_anom", "综合目标异常磁场", "B_static+B_induced"],
            ["r_c(t)", "运动目标中心位置", "m"],
            ["r", "目标中心到观测点的位移或距离", "m"],
        ],
        [1700, 4200, 3460],
        font_size=8.8,
    )
    add_heading(document, "附录 B：专有名词释义", level=1)
    add_table(
        document,
        ["专有名词", "词语释义", "理解要点"],
        [
            ["铁磁材料", "能够被强烈磁化并可能保留磁性的材料，如部分钢铁材料", "舰体外壳、骨架和设备中的铁磁材料是磁特征的重要来源"],
            ["背景地磁场", "目标不存在时，地球在观测位置原本具有的磁场", "决定感应磁化的外部方向和强度"],
            ["磁化", "材料内部微观磁矩在外场作用下趋向排列的过程", "宏观上表现为材料产生磁矩"],
            ["剩磁", "外部磁场减弱或撤去后仍保留的磁化状态", "与材料、加工、建造和服役历史有关"],
            ["感应磁场", "舰体受到当前背景地磁作用而产生的异常磁场", "地磁改变时，其大小和方向也会改变"],
            ["静磁场", "在所研究时间尺度内近似不随时间变化的磁场", "本文主要指剩磁形成的目标异常磁场"],
            ["磁矩", "同时描述磁性强弱和主方向的矢量", "磁矩方向不等于空间中所有位置的磁场方向"],
            ["磁偶极子", "由相邻异名磁极构成的理想化磁源", "远离目标时可用一个磁偶极子近似整体舰艇"],
            ["等效磁性椭球", "用相同长、宽、高的椭球近似复杂舰体磁性结构", "保留总体尺寸、主轴方向和退磁差异"],
            ["相对磁导率", "材料导引磁通能力相对于真空的倍数", "数值越大通常越容易被磁化，但会受到退磁作用限制"],
            ["磁化率", "材料磁化强度对外部磁场强度的响应系数", "线性近似下 χ=μ_r−1"],
            ["退磁场", "被磁化物体自身产生、并倾向于削弱内部磁化的反向磁场", "是形状影响磁化效果的主要原因"],
            ["退磁因子", "把物体形状对退磁作用的影响表示成的无量纲系数", "细长舰体通常沿纵向更容易磁化"],
            ["目标异常磁场", "目标引起的、相对于背景地磁的附加变化量", "通常是磁探测和水雷磁引信关注的对象"],
            ["三分量", "磁场在东、北、上三个互相垂直方向的投影", "保留磁场方向和正负变化信息"],
            ["场强模值", "三个磁场分量平方和开根号得到的总强度", "始终非负，但会丢失方向信息"],
            ["磁偏角", "水平地磁方向相对真北方向的夹角", "用于把总地磁强度转换成东、北分量"],
            ["磁倾角", "地磁方向相对水平面的俯仰角", "用于计算垂向地磁分量"],
            ["近场", "观测距离与目标尺寸相当的区域", "局部结构影响明显，单偶极子近似可能不准确"],
            ["远场", "观测距离显著大于目标尺寸的区域", "整体磁偶极子近似通常更合理"],
            ["运动时序", "固定传感器处磁场随目标运动而形成的时间变化序列", "本质上是对空间磁场分布的一条动态切片"],
        ],
        [1900, 4550, 2910],
        font_size=8.4,
    )
    add_callout(
        document,
        "结论",
        "这套模型的价值在于“计算快、参数含义清楚、能够生成三分量空间和时间数据”；它最适合讲解、算法验证和标定前的方案比较。",
    )

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


def main() -> None:
    """构建并保存最终 Word 文档。"""

    missing_assets = [
        name
        for name in [
            "principle_schematic.png",
            "simulation_flow.png",
            "spatial_heatmaps.png",
            "induced_components.png",
            "time_series.png",
            "parameter_sensitivity.png",
        ]
        if not (ASSET_DIR / name).exists()
    ]
    if missing_assets:
        raise FileNotFoundError(f"缺少文档图片：{missing_assets}")
    if not SPATIAL_CSV.exists() or not TIME_CSV.exists():
        raise FileNotFoundError("缺少仿真 CSV，请先运行验证示例。")

    output = build_document()
    print(f"已生成 Word 文档：{output}")


if __name__ == "__main__":
    main()
