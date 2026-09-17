"""更新水压场仿真原理文档，使其与当前分工况模型和 MATLAB 验证保持一致。"""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


# 统一定义仓库内的输入、备份和验证图片路径，避免脚本依赖当前工作目录。
ROOT_DIR = Path(__file__).resolve().parents[1]
DOC_PATH = ROOT_DIR / "docs" / "水压场仿真原理与讲解指南.docx"
BACKUP_PATH = (
    ROOT_DIR
    / "docs"
    / "水压场仿真原理与讲解指南（2026-09-16模型分支更新前备份）.docx"
)
RESULT_DIR = ROOT_DIR / "build" / "matlab_pressure_regime_2017_compat"
TIME_SERIES_IMAGE = RESULT_DIR / "pressure_regime_time_series.png"
TRANSITION_IMAGE = RESULT_DIR / "pressure_transition_behavior.png"
PRESSURE_FIELD_IMAGE = RESULT_DIR / "surface_ship_pressure_fields.png"
FLOWCHART_IMAGE = (
    ROOT_DIR / "build" / "pressure_principles_assets" / "pressure_regime_flowchart.png"
)


def paragraph_full_text(paragraph: Paragraph) -> str:
    """同时读取普通文字与 OMML 公式文字，用于稳定定位原文段落。"""

    return "".join(
        element.text or ""
        for element in paragraph._p.iter()
        if element.tag.endswith("}t")
    )


def find_paragraph(
    document: Document, prefix: str, style: str | None = None
) -> Paragraph:
    """按段落开头查找唯一锚点，未找到时立即报错，防止误改其他内容。"""

    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph_full_text(paragraph).startswith(prefix)
        and (style is None or paragraph.style.name == style)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"段落锚点数量异常：{prefix!r}，实际找到 {len(matches)} 个")
    return matches[0]


def clear_paragraph(paragraph: Paragraph) -> None:
    """保留段落样式属性，仅移除原有文字、图片和公式内容。"""

    paragraph_properties = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not paragraph_properties:
            paragraph._p.remove(child)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """替换段落文字，并显式设置中文字体以保证跨 Office 版本显示稳定。"""

    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    run.font.name = "微软雅黑"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_math_run(math: OxmlElement, text: str) -> None:
    """向原生 Word 公式中添加一段可编辑的数学文字。"""

    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.set(qn("xml:space"), "preserve")
    math_text.text = text
    math_run.append(math_text)
    math.append(math_run)


def set_equation(paragraph: Paragraph, expression: str) -> None:
    """将段落替换为居中的原生 OMML 公式，避免使用截图或原始 LaTeX。"""

    clear_paragraph(paragraph)
    paragraph.style = "Formula"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math_paragraph = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    add_math_run(math, expression)
    math_paragraph.append(math)
    paragraph._p.append(math_paragraph)


def new_paragraph_before(
    anchor: Paragraph, text: str = "", style: str | None = None
) -> Paragraph:
    """在指定段落前插入新段落，并保持当前文档的段落样式体系。"""

    paragraph_element = OxmlElement("w:p")
    anchor._p.addprevious(paragraph_element)
    paragraph = Paragraph(paragraph_element, anchor._parent)
    if style is not None:
        paragraph.style = style
    if text:
        set_paragraph_text(paragraph, text)
    return paragraph


def previous_paragraph(paragraph: Paragraph) -> Paragraph:
    """返回当前段落前一段，供替换紧邻标题或图注的公式与图片。"""

    previous_element = paragraph._p.getprevious()
    if previous_element is None or previous_element.tag != qn("w:p"):
        raise RuntimeError(f"段落前方不是可编辑段落：{paragraph_full_text(paragraph)!r}")
    return Paragraph(previous_element, paragraph._parent)


def remove_preceding_page_break(paragraph: Paragraph) -> None:
    """移除紧邻段落前的显式分页，避免内容回流后产生空白页。"""

    previous_element = paragraph._p.getprevious()
    if previous_element is None or previous_element.tag != qn("w:p"):
        return
    page_breaks = previous_element.xpath('.//w:br[@w:type="page"]')
    if page_breaks:
        previous_element.getparent().remove(previous_element)


def add_page_break_before(paragraph: Paragraph) -> None:
    """在段落前插入显式分页，用于把一级章节与前一节内容分开。"""

    page_break_paragraph = new_paragraph_before(paragraph)
    page_break_paragraph.add_run().add_break(WD_BREAK.PAGE)


def add_page_break_before_table(table) -> None:
    """在表格前插入分页，防止表头单独留在上一页。"""

    paragraph_element = OxmlElement("w:p")
    table._tbl.addprevious(paragraph_element)
    paragraph = Paragraph(paragraph_element, table._parent)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def set_explicit_numbered_text(paragraph: Paragraph, number: int, text: str) -> None:
    """移除自动编号并写入明确序号，避免不同列表之间发生编号串联。"""

    paragraph_properties = paragraph._p.get_or_add_pPr()
    numbering_properties = paragraph_properties.find(qn("w:numPr"))
    if numbering_properties is not None:
        paragraph_properties.remove(numbering_properties)
    paragraph.style = "Normal"
    set_paragraph_text(paragraph, f"{number}.  {text}")


def add_picture_before(anchor: Paragraph, image_path: Path, width_inches: float) -> Paragraph:
    """在锚点前插入居中的验证图片，并控制宽度以避免越过页面边距。"""

    paragraph = new_paragraph_before(anchor)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    return paragraph


def set_picture(paragraph: Paragraph, image_path: Path, width_inches: float) -> None:
    """用新的仿真结果图替换原图，保留原段落所在位置。"""

    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    """设置表格单元格文字、中文字体和垂直居中格式。"""

    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "微软雅黑"
    run.font.size = Pt(9)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill: str) -> None:
    """为表头设置浅色底纹，延续原文档清晰、克制的表格风格。"""

    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        cell_properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "B7C3D0", size: str = "4") -> None:
    """为表格单元格设置浅灰实线边框，保证打印和低亮度屏幕下仍可辨认。"""

    cell_properties = cell._tc.get_or_add_tcPr()
    borders = cell_properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        cell_properties.append(borders)
    for edge_name in ("top", "left", "bottom", "right"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)


def set_table_data(table, rows: list[list[str]]) -> None:
    """按给定二维数据重建表格行，保留既有表格对象及其样式。"""

    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) < len(rows):
        table.add_row()
    while len(table.columns) < len(rows[0]):
        table.add_column(Inches(1.2))

    for row_index, row_data in enumerate(rows):
        for column_index, value in enumerate(row_data):
            cell = table.cell(row_index, column_index)
            set_cell_text(cell, value, bold=(row_index == 0))
            set_cell_borders(cell)
            if row_index == 0:
                shade_cell(cell, "D9E2F3")


def add_table_before(
    document: Document,
    anchor: Paragraph,
    rows: list[list[str]],
    reference_table_index: int = 0,
):
    """在指定段落前插入表格，并复用现有表格样式。"""

    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = document.tables[reference_table_index].style
    table.autofit = True
    set_table_data(table, rows)
    anchor._p.addprevious(table._tbl)
    return table


def add_formula_row(table, name: str, expression: str) -> None:
    """向公式速查表追加一行，并在公式单元格内使用原生 Word 公式。"""

    row = table.add_row()
    set_cell_text(row.cells[0], name)
    formula_paragraph = row.cells[1].paragraphs[0]
    set_equation(formula_paragraph, expression)


def add_glossary_row(table, term: str, explanation: str, role: str) -> None:
    """向术语表追加新模型概念，便于非流体专业读者查阅。"""

    row = table.add_row()
    set_cell_text(row.cells[0], term)
    set_cell_text(row.cells[1], explanation)
    set_cell_text(row.cells[2], role)


def validate_inputs() -> None:
    """在写入前检查文档和三幅 MATLAB 验证图是否齐全。"""

    required_paths = [
        DOC_PATH,
        TIME_SERIES_IMAGE,
        TRANSITION_IMAGE,
        PRESSURE_FIELD_IMAGE,
        FLOWCHART_IMAGE,
    ]
    missing_paths = [path for path in required_paths if not path.exists()]
    if missing_paths:
        missing_text = "\n".join(str(path) for path in missing_paths)
        raise FileNotFoundError(f"缺少更新所需文件：\n{missing_text}")


def update_document() -> None:
    """执行正文、公式、表格和验证图片的一致性更新。"""

    validate_inputs()

    # 首次运行时保留一份明确命名的修改前备份，便于回溯原模型说明。
    if not BACKUP_PATH.exists():
        shutil.copy2(DOC_PATH, BACKUP_PATH)

    document = Document(DOC_PATH)
    # 固定保存原文档表格引用，避免插入新工况表后序号发生变化。
    original_tables = list(document.tables)

    # 更新封面、副标题和阅读导航，使读者先看到分工况建模的核心变化。
    set_paragraph_text(
        find_paragraph(document, "舰艇 / 潜艇运动目标"),
        "舰艇 / 潜艇运动目标 · 低速刚盖与浅水亚临界兴波分工况模型",
    )
    set_paragraph_text(
        find_paragraph(document, "从“船体推水”"),
        "从“船体推水”这一生活化直觉出发，逐步解释回转体偶极子、低速刚盖、浅水亚临界兴波、平滑过渡、潜艇兴波衰减和工程适用范围。",
    )
    set_paragraph_text(
        find_paragraph(document, "有限水深快速仿真模型"),
        "分工况水压场快速仿真模型\n原理讲解版  ·  2026 年 9 月",
    )
    set_paragraph_text(
        find_paragraph(document, "一句话概括"),
        "一句话概括  舰艇和潜艇不再共用单一边界模型：水面舰艇低速时采用回转体刚盖模型，浅水亚临界较高速时加入兴波压力并在 FrL=0.08～0.12 内平滑过渡；潜艇深潜或低速时忽略自由表面兴波。",
    )
    set_paragraph_text(
        find_paragraph(document, "5  有限水深", style="Normal"),
        "5  自由表面兴波、工况判定与有限水深边界",
    )
    set_paragraph_text(
        find_paragraph(document, "7  真实示例", style="Normal"),
        "7  MATLAB 分工况验证：如何读仿真结果",
    )
    set_paragraph_text(
        find_paragraph(document, "重要定位"),
        "重要定位  这是面向算法验证和合成数据的低阶工程模型。浅水亚临界兴波项用于改善工况趋势，不等同于高保真自由液面 CFD；深水兴波、临界与超临界工况仍需升级模型或使用试验数据。",
    )

    # 更新模型范围，消除“完全不考虑兴波”这一已过时的描述。
    set_paragraph_text(
        find_paragraph(document, "本文采用有限水深势流偶极子近似"),
        "本文采用分工况低阶势流模型：船体本体以运动偶极子表示；低速工况用有限水深刚盖镜像处理上下边界；水面舰艇在浅水亚临界较高速工况下叠加线性深度平均兴波压力，并通过平滑权重避免模型切换产生跳变。",
    )
    set_paragraph_text(
        find_paragraph(document, "静水面与海床的有限水深边界镜像影响"),
        "低速刚盖镜像、浅水亚临界兴波修正，以及潜艇深潜/低速时的自由表面兴波忽略判据。",
    )
    set_paragraph_text(
        find_paragraph(document, "自由表面真实兴波、破波和波浪背景噪声"),
        "深水兴波、浅水临界/超临界波系、破波、真实波面演化以及环境波浪背景噪声。",
    )
    set_paragraph_text(
        find_paragraph(document, "讲解建议"),
        "讲解建议  可以把模型理解为“先用回转体偶极子给出船体推水的主体压力，再由工况判据选择低速边界修正或浅水兴波修正”。这种分层结构保留了快速计算优势，也明确了不同航速、水深和潜深下的适用边界。",
    )

    # 在输入符号表中补充工况选择所需的无量纲量与可调阈值。
    symbol_table = original_tables[0]
    for symbol, unit, meaning in [
        ("FrL", "—", "船长弗劳德数 U/√(gL)，用于判断低速与兴波过渡"),
        ("FrH", "—", "水深弗劳德数 U/√(gH)，用于判断亚临界、临界和超临界"),
        ("H/L", "—", "相对水深；H/L≤0.30 判为浅水"),
        ("w", "—", "兴波模型平滑权重，范围 0～1"),
        ("Ew", "—", "潜艇自由表面兴波估算衰减量 exp(−gh/U²)"),
    ]:
        row = symbol_table.add_row()
        set_cell_text(row.cells[0], symbol)
        set_cell_text(row.cells[1], unit)
        set_cell_text(row.cells[2], meaning)

    # 原“真实自由面/刚盖”对比表改为两条已实现计算分支的差异说明。
    set_table_data(
        original_tables[1],
        [
            ["对比项", "低速刚盖模型", "浅水亚临界兴波模型"],
            ["自由表面", "固定在 z=0，不计算水面起伏", "用线性深度平均长波压力近似"],
            ["主要边界处理", "海面与海床周期镜像", "本体压力、兴波压力与海床修正"],
            ["适用判据", "FrL≤0.08（过渡带外）", "H/L≤0.30、FrH<0.95、FrL≥0.12"],
            ["过渡方式", "权重 1−w 逐渐减小", "权重 w 逐渐增大"],
        ],
    )

    # 重新定义第 5 章为模型选择主章节，并复用原有镜像图说明低速分支。
    set_paragraph_text(
        find_paragraph(document, "5  有限水深：镜像偶极子如何工作", style="Heading 1"),
        "5  自由表面兴波、工况判定与有限水深边界",
    )
    remove_preceding_page_break(
        find_paragraph(document, "5  自由表面兴波、工况判定与有限水深边界", style="Heading 1")
    )
    set_paragraph_text(
        find_paragraph(document, "无限深海中只需计算真实目标"),
        "当前模型先根据目标类型、航速、船长、海深和潜深选择工况，再计算压力分量。水面舰艇与潜艇采用不同判据；低速模型、浅水亚临界兴波模型及其过渡结果都会在输出中明确标识。下图所示镜像偶极子是低速刚盖分支的边界处理方式。",
    )
    set_paragraph_text(
        find_paragraph(document, "图 3  自由表面和海床的镜像偶极子示意"),
        "图 3  低速刚盖分支中的自由表面与海床镜像偶极子",
    )
    set_paragraph_text(find_paragraph(document, "5.1  三个边界处理名词"), "5.1  三个工况判据")
    set_paragraph_text(
        find_paragraph(document, "镜像法（method of images）"),
        "船长弗劳德数 FrL：以目标长度 L 为特征尺度，衡量航速相对重力波速度的大小，用于低速/兴波模型切换。",
    )
    set_paragraph_text(
        find_paragraph(document, "刚盖近似（rigid-lid approximation）"),
        "水深弗劳德数 FrH：以水深 H 为特征尺度；FrH<0.95 为当前浅水亚临界适用域，0.95～1.05 为临界缓冲区，FrH>1.05 为超临界。",
    )
    set_paragraph_text(
        find_paragraph(document, "无穿透边界（no-penetration boundary）"),
        "相对水深 H/L：当 H/L≤0.30 时判为浅水；大于 0.30 时属于当前尚未实现的深水兴波分支。",
    )
    set_paragraph_text(
        find_paragraph(document, "5.1.1  刚盖近似是什么"),
        "5.1.1  水面舰艇：低速、平滑过渡与浅水亚临界",
    )
    set_paragraph_text(
        find_paragraph(document, "真实自由水面会随流体运动上下起伏"),
        "水面舰艇首先计算 FrL、FrH 和 H/L。FrL≤0.08 时采用低速回转体刚盖模型；在浅水且亚临界条件满足时，0.08<FrL<0.12 为平滑过渡带，FrL≥0.12 采用浅水亚临界兴波模型。",
    )
    set_equation(
        find_paragraph(document, "w(z=0)"),
        "Frₗ = U / √(gL),    Frₕ = U / √(gH),    λ = H / L",
    )
    set_paragraph_text(
        find_paragraph(document, "直观理解：w 是垂向流速"),
        "直观理解：FrL 决定是否需要从低速模型转入兴波模型，FrH 决定浅水波是否仍处于亚临界，λ=H/L 决定是否属于浅水。三个判据必须联合使用。",
    )
    set_paragraph_text(
        find_paragraph(document, "刚盖近似取消了水面形状"),
        "过渡带中心取 FrL=0.10、半宽取 0.02，因此边界为 0.08 和 0.12。先把 FrL 线性归一化为 s，再使用 smoothstep 权重 w=s²(3−2s)；该函数在两端一阶导数为零，可避免压力随航速出现突跳。",
    )
    set_equation(
        previous_paragraph(find_paragraph(document, "何时适用  刚盖近似")),
        "s = clamp[(Frₗ − 0.08) / 0.04, 0, 1],    w = s²(3 − 2s)",
    )
    set_paragraph_text(
        find_paragraph(document, "何时适用  刚盖近似"),
        "适用域  H/L≤0.30 且 FrH<0.95 时可进入浅水亚临界或过渡模型。若 H/L>0.30 且 FrL>0.10，则需要深水兴波模型；若 0.95≤FrH≤1.05 或 FrH>1.05，则分别属于当前不支持的临界或超临界工况。",
    )
    set_paragraph_text(
        find_paragraph(document, "5.1.2  为什么采用刚盖近似"),
        "5.1.2  潜艇：深潜或低速时忽略水面兴波",
    )
    set_paragraph_text(
        find_paragraph(document, "真实自由水面需要同时求解水面形状"),
        "潜艇自由表面兴波随艇体中心潜深 h 增大而迅速衰减。当前模型用 Ew=exp(−gh/U²) 估算兴波传到自由面的相对量；U=0 时直接令 Ew=0。",
    )
    set_paragraph_text(find_paragraph(document, "采用刚盖近似的主要理由"), "无兴波潜艇模型满足以下任一条件即可启用：")
    set_paragraph_text(find_paragraph(document, "比完全忽略海面更合理"), "低速条件：FrL≤0.10。")
    set_paragraph_text(find_paragraph(document, "比真实自由水面更简单"), "深潜条件：Ew≤0.01，即自由表面兴波估算量不超过 1%。")
    set_paragraph_text(find_paragraph(document, "可以使用镜像法"), "低速潜艇采用刚盖有限水深压力，保留海面和海床无穿透边界的镜像影响。")
    set_paragraph_text(find_paragraph(document, "计算快且稳定"), "较高速但深潜、Ew 仍很小时，采用本体偶极子加第一海床镜像，不再叠加刚盖自由面镜像。")
    set_paragraph_text(find_paragraph(document, "适合批量仿真"), "近水面且 FrL>0.10、Ew>0.01 时，潜艇自由表面兴波不能忽略；当前实现将其标为不支持，不用无兴波模型强行外推。")
    set_paragraph_text(find_paragraph(document, "与整体模型精度一致"), "该指数衰减是工程判据，不是完整的潜艇兴波解；阈值 0.01 可按试验或高保真计算结果标定。")
    set_equation(find_paragraph(document, "无限水域（忽略海面）"), "Ew = exp(−g h / U²),    h = −zc")
    set_paragraph_text(
        find_paragraph(document, "直观理解：当前模型选择中间层级"),
        "直观理解：潜艇越深，兴波到达水面的能量越弱；潜艇越快，深度衰减越不充分。因此“深潜”和“低速”是两条相互独立的无兴波理由。",
    )
    set_paragraph_text(
        find_paragraph(document, "选择原则"),
        "选择原则  只有模型适用域明确支持的工况才输出计算值；深水水面舰艇兴波、浅水临界/超临界以及近水面高速潜艇兴波均显式报为不支持，避免把简化模型用于错误场景。",
    )

    # 在镜像边界章节前加入一张工况选择表，便于快速核对代码分支。
    regime_anchor = find_paragraph(document, "5.2  为什么要生成镜像目标")
    add_table_before(
        document,
        regime_anchor,
        [
            ["目标/工况", "判据", "采用模型"],
            ["水面舰艇低速", "FrL≤0.08；或深水且 FrL≤0.10", "回转体偶极子 + 刚盖周期镜像"],
            ["水面舰艇过渡", "H/L≤0.30，FrH<0.95，0.08<FrL<0.12", "低速边界与浅水兴波按 smoothstep 混合"],
            ["水面舰艇浅水亚临界", "H/L≤0.30，FrH<0.95，FrL≥0.12", "本体偶极子 + 浅水亚临界兴波/海床修正"],
            ["潜艇无兴波", "FrL≤0.10 或 Ew≤0.01", "低速刚盖；或深潜时本体 + 第一海床镜像"],
            ["当前不支持", "深水兴波、临界/超临界、近水面高速潜艇", "停止计算并报告工况"],
        ],
        reference_table_index=1,
    )

    # 明确镜像模型只属于低速分支，并保留原文的镜像族公式与截断说明。
    set_paragraph_text(find_paragraph(document, "5.2  为什么要生成镜像目标"), "5.2  低速刚盖模型：为什么要生成镜像目标")
    set_paragraph_text(
        find_paragraph(document, "如果只计算真实舰艇的偶极子压力"),
        "在低速无兴波分支中，如果只计算真实目标的偶极子压力，就相当于目标处在没有海面和海床的无限水域。刚盖模型在水层外放置虚拟偶极子并叠加其压力，使静水面和海床处近似满足无穿透条件。",
    )
    set_paragraph_text(
        find_paragraph(document, "不要把镜像法误解为压力波反射"),
        "不要把镜像法误解为压力波反射  镜像项只处理低速理想势流的无穿透边界。浅水亚临界分支改用深度平均兴波修正；两者在过渡带内由权重连续混合。",
    )
    set_paragraph_text(
        find_paragraph(document, "海深 H 会改变所有镜像点的位置"),
        "在低速刚盖分支中，海深 H 会改变所有镜像点的位置；在浅水兴波分支中，H 还同时进入 FrH、亚临界伸缩因子 β 和压力尺度 1/H。因此海深影响不是单一比例，必须结合所处工况解释。",
    )

    # 在第 6 章前补充浅水亚临界兴波公式、核半径和压力分量合成关系。
    chapter6_anchor = find_paragraph(
        document, "6  目标运动、时间采样与空间网格", style="Heading 1"
    )
    remove_preceding_page_break(chapter6_anchor)
    new_paragraph_before(chapter6_anchor, "5.5  浅水亚临界兴波压力与平滑合成", "Heading 2")
    new_paragraph_before(
        chapter6_anchor,
        "当 H/L≤0.30 且 FrH<0.95 时，模型采用线性深度平均二维偶极子近似。亚临界伸缩因子 β=√(1−FrH²) 将纵向坐标放大；当 FrH 接近 1 时 β 变小，因此临界附近不再使用本公式。",
        "Normal",
    )
    equation = new_paragraph_before(chapter6_anchor, style="Formula")
    set_equation(equation, "β = √(1 − Frₕ²),    xβ = x / β,    rc = max(rmin, B/2)")
    new_paragraph_before(
        chapter6_anchor,
        "为避免二维点偶极子在中心附近产生数值奇异，水平核半径取 rc=max(rmin,B/2)，并用 rβ²=max(xβ²+y²,rc²) 限制最小尺度。该处理属于正则化，不应被解释为真实船体近场细节。",
        "Normal",
    )
    equation = new_paragraph_before(chapter6_anchor, style="Formula")
    set_equation(equation, "pwave = [ρ q U² / (Hβ)] · (xβ² − y²) / (rβ²)²")
    new_paragraph_before(
        chapter6_anchor,
        "式中 q 为单位航速偶极子强度。该兴波压力在纵向为正、横向为负，按水平距离平方量级衰减；它描述的是浅水长波趋势，不直接给出真实波面高度、开尔文波系或破波。",
        "Normal",
    )
    equation = new_paragraph_before(chapter6_anchor, style="Formula")
    set_equation(equation, "pdyn = pbody + (1 − w) plow,boundary + w pshallow,boundary")
    new_paragraph_before(
        chapter6_anchor,
        "压力输出按本体、自由表面/兴波和海床修正分量组织。过渡带内低速边界分量乘以 (1−w)，浅水兴波边界分量乘以 w；由于 w 及其端点斜率连续，合成压力不会在 0.08 和 0.12 处跳变。",
        "Normal",
    )
    add_page_break_before(chapter6_anchor)

    # 输出说明新增工况、无量纲数、过渡权重和压力分量，便于前端追溯计算来源。
    set_paragraph_text(
        find_paragraph(document, "给定观测位置和时刻"),
        "给定观测位置和时刻，模型返回目标位置、距离、纵横向相对位置、观测深度、FrL、FrH、潜艇兴波衰减量、过渡权重、实际工况，以及静水压力、本体动压、自由表面/兴波修正、海床修正、总动态压力和总表压。",
    )
    flowchart_caption = find_paragraph(document, "图 4  从输入参数到压力结果的原理流程")
    set_picture(previous_paragraph(flowchart_caption), FLOWCHART_IMAGE, 5.75)
    set_paragraph_text(flowchart_caption, "图 4  从无量纲判据到分工况压力合成的仿真流程")

    # 用 MATLAB 2017 兼容仿真结果替换已不再落在支持域内的旧单一示例。
    set_paragraph_text(
        find_paragraph(document, "7  真实示例：如何读压力曲线", style="Heading 1"),
        "7  MATLAB 分工况验证：如何读仿真结果",
    )
    set_paragraph_text(
        find_paragraph(document, "工程示例使用水面舰艇"),
        "配套 MATLAB 流程在相同几何与观测布置下构造四个代表工况：水面舰艇低速回转体、平滑过渡、浅水亚临界，以及深潜潜艇无兴波。脚本采用 MATLAB 2017 可用的语法和绘图接口，并导出时间序列、工况摘要和空间压力图。",
    )
    old_result_image = previous_paragraph(find_paragraph(document, "图 5  示例工况的动态压力"))
    set_picture(old_result_image, TIME_SERIES_IMAGE, 6.25)
    set_paragraph_text(
        find_paragraph(document, "图 5  示例工况的动态压力、距离和航迹"),
        "图 5  四类支持工况的动态压力时间序列（MATLAB 仿真）",
    )
    set_table_data(
        original_tables[3],
        [
            ["工况", "关键判据", "动态压力绝对峰值"],
            ["水面舰艇低速回转体", "FrL=0.06，w=0", "9.61 Pa"],
            ["水面舰艇平滑过渡", "FrL=0.10，w=0.50", "30.11 Pa"],
            ["水面舰艇浅水亚临界", "FrL=0.18，FrH=0.329", "111.12 Pa"],
            ["深潜潜艇无兴波", "FrL=0.286，Ew≤0.01", "19.91 Pa"],
        ],
    )
    set_paragraph_text(find_paragraph(document, "7.1  为什么总表压图不明显"), "7.1  如何解读验证结果")
    set_paragraph_text(
        find_paragraph(document, "示例静水压力约为"),
        "四条曲线均在目标最近通过附近出现最显著变化。低速到浅水亚临界的峰值随航速和兴波权重增大，但这不意味着所有工况都只按 U² 缩放：一旦进入过渡或浅水分支，FrH、β、H/L 以及边界分量混合都会共同改变幅值和空间形态。",
    )
    set_paragraph_text(
        find_paragraph(document, "绘图建议"),
        "验证定位  这些结果证明工况选择、平滑权重和结果导出流程能够运行，并用于观察模型趋势；它们不是对真实舰型试验数据的精度验证。工程使用前仍应以实测或高保真 CFD 对阈值、衰减容差和幅值进行标定。",
    )

    # 在参数分析章节前增加空间压力图，使不同分支的形态差异可直接观察。
    chapter8_anchor = find_paragraph(document, "8  参数影响与量纲分析", style="Heading 1")
    remove_preceding_page_break(chapter8_anchor)
    add_picture_before(chapter8_anchor, PRESSURE_FIELD_IMAGE, 6.25)
    new_paragraph_before(
        chapter8_anchor,
        "图 6  水面舰艇低速、过渡与浅水亚临界的空间动态压力分布",
        "Figure Caption",
    )
    new_paragraph_before(
        chapter8_anchor,
        "从左到右，压力场由低速刚盖偶极子主导逐步过渡到浅水深度平均兴波结构。比较空间图时应保持相同坐标范围和色标原则，并屏蔽点偶极子近场核心区。",
        "Normal",
    )
    add_page_break_before(chapter8_anchor)

    # 用过渡权重曲线替换原有单一比例示意，并修正参数趋势表中的过时结论。
    old_scaling_image = previous_paragraph(find_paragraph(document, "图 6  航速、排水体积和距离"))
    set_picture(old_scaling_image, TRANSITION_IMAGE, 6.10)
    set_paragraph_text(
        find_paragraph(document, "图 6  航速、排水体积和距离的主要标度规律"),
        "图 7  FrL=0.08～0.12 过渡带的 smoothstep 权重及压力连续变化",
    )
    parameter_table = original_tables[4]
    for row in parameter_table.rows[1:]:
        if row.cells[0].text.strip() == "航速 U":
            set_cell_text(row.cells[1], "本体项含 U²，同时改变 FrL、FrH、β 与模型权重")
            set_cell_text(row.cells[2], "通常增强；跨工况时不再是单一 U² 比例")
        if row.cells[0].text.strip() == "海深 H":
            set_cell_text(row.cells[1], "改变镜像位置、H/L、FrH、β 和浅水压力尺度")
            set_cell_text(row.cells[2], "同时影响工况选择与幅值，必须联合分析")
    add_page_break_before_table(parameter_table)

    # 在结果解读中强调分量与工况字段，避免只比较总动压峰值。
    set_paragraph_text(
        find_paragraph(document, "先看基线"),
        "先看工况与基线：确认 regime、FrL、FrH、H/L、Ew 和 w，再查看静水压力基线；不支持工况不应继续解释数值结果。",
    )
    set_paragraph_text(
        find_paragraph(document, "再看峰值"),
        "再看压力分量：分别检查本体、自由表面/兴波和海床修正，再记录总动态压力的最大正峰与最大负峰。",
    )
    set_paragraph_text(
        find_paragraph(document, "比较工况时保持几何关系一致"),
        "比较工况时保持几何关系一致，并标明是否跨越 0.08～0.12 过渡带、H/L=0.30 浅水边界或 FrH 临界区，避免把模型切换误判为单纯航速效应。",
    )

    # 修订适用范围、升级路径和问答，使文档与新实现的支持/不支持状态一致。
    set_paragraph_text(
        find_paragraph(document, "做参数趋势研究、教学演示和不同工况对比"),
        "做低速、过渡、浅水亚临界和深潜潜艇无兴波工况的参数趋势研究、教学演示与对比。",
    )
    set_paragraph_text(
        find_paragraph(document, "分析螺旋桨叶频"),
        "分析螺旋桨叶频、尾流湍流、空化、深水航行波、浅水临界/超临界波系或近水面高速潜艇兴波。",
    )
    set_paragraph_text(
        find_paragraph(document, "若形状差异仍明显"),
        "若形状或波系差异仍明显，再升级为分布源、线性自由表面面元法或带自由液面的 CFD；临界与超临界工况不应通过调节本模型系数来替代。",
    )

    # 章节内容回流后移除第 10 章前分页，避免第 9 章末尾出现大面积空白。
    remove_preceding_page_break(
        find_paragraph(document, "10  适用范围、标定建议与常见问答", style="Heading 1")
    )
    set_paragraph_text(
        find_paragraph(document, "问：海越浅压力一定越大吗"),
        "问：H/L=0.30 是物理突变点吗？",
    )
    set_paragraph_text(
        find_paragraph(document, "答：不一定。海深同时改变多项镜像"),
        "答：不是。它是当前模型的浅水适用域阈值，用于选择算法分支；真实流体不会在 H/L=0.30 处突然改变。阈值附近应结合试验或更高保真模型检查。",
    )
    set_paragraph_text(
        find_paragraph(document, "答：可以。将目标中心设置到实际潜深"),
        "答：可以，但只在 FrL≤0.10 或 Ew≤0.01 时采用无兴波潜艇模型。近水面且较高速、Ew>0.01 的潜艇自由表面兴波尚未实现，程序会明确报告不支持。",
    )

    # 自动编号在新增列表后可能跨章节续编，改为显式编号以确保每组都从 1 开始。
    calibration_items = [
        ("收集目标尺寸", "收集目标尺寸、排水量、航速、航迹和传感器几何关系。"),
        ("用排水量反推", "用排水量反推或修正 blockCoefficient。"),
        ("对齐实测与模拟", "对齐实测与模拟的时间轴，先校正最近通过时刻和航迹。"),
        ("比较动态压力峰值", "比较动态压力峰值、峰宽和正负结构，建立幅值修正系数。"),
        ("用另一组工况验证", "用另一组工况验证修正系数，避免只对单次试验过拟合。"),
        ("若形状或波系差异", "若形状或波系差异仍明显，再升级为分布源、线性自由表面面元法或带自由液面的 CFD；临界与超临界工况不应通过调节本模型系数来替代。"),
    ]
    for item_number, (prefix, text) in enumerate(calibration_items, start=1):
        set_explicit_numbered_text(find_paragraph(document, prefix), item_number, text)

    outline_items = [
        ("先说现象", "先说现象：船体运动会推水，固定水下点会出现随时间变化的压力。"),
        ("再说简化", "再说简化：把复杂船体等效成由体积和形状决定的势流偶极子。"),
        ("指出三条规律", "指出三条规律：压力约随航速平方、体积线性变化，并按距离三次方衰减。"),
        ("解释海深", "解释海深：先用 FrL、FrH 和 H/L 判定低速、过渡或浅水亚临界分支。"),
        ("展示曲线", "展示曲线：说明平滑权重、压力分量和不同工况的空间形态。"),
        ("最后强调边界", "最后强调边界：深水兴波、临界/超临界和近水面高速潜艇仍需更高保真模型。"),
    ]
    for item_number, (prefix, text) in enumerate(outline_items, start=1):
        set_explicit_numbered_text(find_paragraph(document, prefix), item_number, text)

    # 扩充附录公式速查表，所有新增公式继续使用可编辑的原生 Word 公式。
    formula_table = original_tables[6]
    add_formula_row(formula_table, "船长/水深弗劳德数", "Frₗ = U / √(gL),    Frₕ = U / √(gH)")
    add_formula_row(formula_table, "浅水判据", "H / L ≤ 0.30")
    add_formula_row(formula_table, "平滑过渡权重", "s = clamp[(Frₗ−0.08)/0.04,0,1],    w = s²(3−2s)")
    add_formula_row(formula_table, "潜艇兴波衰减", "Ew = exp(−gh/U²)")
    add_formula_row(formula_table, "浅水亚临界伸缩", "β = √(1−Frₕ²),    xβ = x/β")
    add_formula_row(formula_table, "浅水兴波压力", "pwave = [ρqU²/(Hβ)]·(xβ²−y²)/(rβ²)²")
    add_formula_row(formula_table, "过渡带动压合成", "pdyn = pbody + (1−w)plow,boundary + wpshallow,boundary")

    # 扩充术语速查表，新增概念均给出通俗解释和在代码中的作用。
    glossary_table = original_tables[7]
    add_glossary_row(glossary_table, "船长弗劳德数 FrL", "航速与以船长为尺度的重力波速度之比。", "判断低速、过渡和兴波分支")
    add_glossary_row(glossary_table, "水深弗劳德数 FrH", "航速与浅水重力波速度之比。", "判断亚临界、临界和超临界")
    add_glossary_row(glossary_table, "平滑过渡带", "在两个模型之间连续改变权重的航速区间。", "FrL=0.08～0.12，避免压力跳变")
    add_glossary_row(glossary_table, "浅水亚临界", "H/L≤0.30 且 FrH<0.95 的受限水深工况。", "启用线性深度平均兴波压力")
    add_glossary_row(glossary_table, "兴波衰减 Ew", "估算潜艇扰动传到自由表面后的剩余比例。", "Ew≤0.01 时可忽略水面兴波")

    set_paragraph_text(
        find_paragraph(document, "进一步学习可参考"),
        "进一步学习可参考 J. N. Newman《Marine Hydrodynamics》及经典势流、浅水长波和自由表面船舶水动力学章节。",
    )
    set_paragraph_text(
        find_paragraph(document, "若需要提高近场精度"),
        "若需要提高近场或自由表面精度，可进一步研究分布源/偶极子模型、线性自由表面面元法、临界与超临界浅水波理论、黏性自由液面 CFD 和模型试验。",
    )
    set_paragraph_text(
        find_paragraph(document, "最后记住"),
        "最后记住  新模型的关键不是“所有舰艇都加兴波”，而是先判工况：低速时用刚盖回转体，浅水亚临界时加入兴波，过渡带平滑混合；潜艇只有在深潜或低速时才能忽略水面兴波。",
    )

    # 同步核心属性，并使用临时文件替换，降低写入中断导致文档损坏的风险。
    document.core_properties.title = "水压场仿真原理与讲解指南"
    document.core_properties.subject = "舰艇与潜艇分工况水压场仿真原理"
    document.core_properties.comments = "2026-09-16：新增浅水判据、平滑过渡、浅水亚临界兴波与潜艇兴波衰减说明。"
    temporary_path = DOC_PATH.with_suffix(".updated.docx")
    document.save(temporary_path)
    temporary_path.replace(DOC_PATH)

    print(f"已更新：{DOC_PATH}")
    print(f"修改前备份：{BACKUP_PATH}")


if __name__ == "__main__":
    # 直接运行脚本即可重复生成与当前算法一致的原理文档。
    update_document()
