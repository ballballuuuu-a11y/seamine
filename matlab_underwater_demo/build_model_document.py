"""生成《水下光学成像模型原理与 MATLAB 演示》Word 文档。"""

from pathlib import Path
from datetime import date
from PIL import Image, ImageChops
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# 文档路径与配色均集中定义，便于后续统一维护。
ROOT = Path(__file__).resolve().parent
OUTPUT_DOCX = ROOT / "水下光学成像模型原理与MATLAB演示.docx"
FORMULA_DIR = ROOT / "doc_assets" / "formulas"
CROPPED_FORMULA_DIR = ROOT / "doc_assets" / "formulas_cropped"
DEMO_OUTPUT_DIR = ROOT / "output"

INK = RGBColor(31, 49, 66)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(102, 112, 123)
LIGHT_FILL = "E8EEF5"
PALE_FILL = "F4F6F9"
CODE_FILL = "F2F4F7"
WHITE = RGBColor(255, 255, 255)


# 这些字符串是文档中全部公式的 LaTeX 源码。
FORMULAS = {
    "f01_projection": r"s\tilde{\mathbf p}=\mathbf K[\,\mathbf R\mid\mathbf t\,]\tilde{\mathbf P}",
    "f02_intrinsic": r"\mathbf K=\begin{bmatrix}f_x&\gamma&c_x\\0&f_y&c_y\\0&0&1\end{bmatrix}",
    "f03_total": r"I_c(\mathbf x)=D_c(\mathbf x)+F_c(\mathbf x)+B_c(\mathbf x)+n_c(\mathbf x)",
    "f04_iop": r"c(\lambda)=a(\lambda)+b(\lambda)",
    "f05_transmission": r"T_\lambda(d)=\exp[-c(\lambda)d]",
    "f06_direct_general": r"D_c(\mathbf x)=\int S_c(\lambda)L_0(\mathbf x,\lambda)\exp[-c(\lambda)d(\mathbf x)]\,\mathrm d\lambda",
    "f07_direct_active": r"D_c(\mathbf x)=G_c(\mathbf x)\rho_c(\mathbf x)E_{0,c}\exp\{-c_c[d_s(\mathbf x)+d_c(\mathbf x)]\}",
    "f08_forward": r"F_c(\mathbf x)=\int_{\Omega}D_c(\boldsymbol\xi)h_{c,d(\mathbf x)}(\mathbf x-\boldsymbol\xi)\,\mathrm d\boldsymbol\xi",
    "f09_back_integral": r"B_c(\mathbf x)=\int_0^{d(\mathbf x)}q_c(s)\exp(-\beta_c^B s)\,\mathrm ds",
    "f10_back_simple": r"B_c(\mathbf x)=B_{\infty,c}\{1-\exp[-\beta_c^B d(\mathbf x)]\}",
    "f11_sensor": r"I_c=\mathcal Q\!\left\{g_c\!\left(\int S_c(\lambda)L(\lambda)\,\mathrm d\lambda\right)+\eta_c\right\}",
    "f12_demo": r"I_c=J_ct_c+\alpha_c(J_ct_c)*g_{\sigma(d)}+B_{\infty,c}(1-e^{-\beta_c^B d})+n_c",
    "f13_inverse": r"\widehat J_c=\mathrm{clip}\!\left(\frac{I_c-\widehat F_c-\widehat B_c}{\max(\widehat t_c,t_{\min})},0,1\right)",
    "f14_metrics": r"\mathrm{PSNR}=10\log_{10}\!\left(\frac{1}{\mathrm{MSE}}\right)",
}


def set_run_font(run, latin="Calibri", east_asia="宋体", size=11, color=INK, bold=None, italic=None):
    """统一设置中西文字体，避免不同 Word 渲染器替换字体。"""
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    """给表格单元格设置固定底色。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    """设置单元格内边距，单位为 DXA。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    """将首行标记为跨页重复表头。"""
    tr_pr = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    tr_pr.append(table_header)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    """设置表格固定宽度、缩进、网格列宽和每个单元格宽度。"""
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    total_width = sum(widths_dxa)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_width))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    """在页脚插入 Word 页码域。"""
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend((field_begin, instruction, field_end))
    set_run_font(run, size=9, color=MUTED)


def configure_document(doc):
    """应用 compact_reference_guide 预设及中文字体覆盖。"""
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Title", 28, DARK_BLUE, 0, 8),
        ("Subtitle", 14, MUTED, 0, 18),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # 采用真实列表样式，并显式设置缩进和段落节奏。
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header_p.add_run("水下光学成像模型 · 原理与 MATLAB 演示")
    set_run_font(header_run, east_asia="微软雅黑", size=9, color=MUTED, bold=True)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer_p.add_run("技术说明  |  ")
    set_run_font(footer_run, size=9, color=MUTED)
    add_page_number(footer_p)


def add_text(doc, text, bold_prefix=None, align=None, keep=False):
    """添加普通正文，可选择加粗开头标签。"""
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.keep_together = keep
    if bold_prefix and text.startswith(bold_prefix):
        label = paragraph.add_run(bold_prefix)
        set_run_font(label, bold=True)
        body = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(body)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(doc, text):
    """添加真实项目符号列表项。"""
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_number(doc, text):
    """添加真实编号列表项。"""
    paragraph = doc.add_paragraph(style="List Number")
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_callout(doc, label, body, fill=PALE_FILL):
    """使用一列表格创建简洁提示框。"""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    label_run = paragraph.add_run(f"{label}  ")
    set_run_font(label_run, east_asia="微软雅黑", color=DARK_BLUE, bold=True)
    body_run = paragraph.add_run(body)
    set_run_font(body_run)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths_dxa):
    """添加带固定几何尺寸的技术表格。"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_run_font(run, east_asia="微软雅黑", size=10, color=DARK_BLUE, bold=True)
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(str(value))
            set_run_font(run, size=9.5)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def crop_formula_images():
    """裁去公式图片的大面积白边，使 Word 排版更紧凑。"""
    CROPPED_FORMULA_DIR.mkdir(parents=True, exist_ok=True)
    for key in FORMULAS:
        source = FORMULA_DIR / f"{key}.png"
        target = CROPPED_FORMULA_DIR / f"{key}.png"
        with Image.open(source).convert("RGB") as image:
            background = Image.new("RGB", image.size, "white")
            difference = ImageChops.difference(image, background).convert("L")
            difference = difference.point(lambda value: 255 if value > 12 else 0)
            box = difference.getbbox()
            if box is None:
                raise RuntimeError(f"公式图片为空：{source}")
            left, top, right, bottom = box
            padding = 20
            box = (
                max(0, left - padding),
                max(0, top - padding),
                min(image.width, right + padding),
                min(image.height, bottom + padding),
            )
            image.crop(box).save(target)


def add_equation(doc, key, number, width=5.4):
    """插入由 LaTeX 渲染的公式图片和公式编号。"""
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(CROPPED_FORMULA_DIR / f"{key}.png"), width=Inches(width))
    inline_shape._inline.docPr.set("descr", f"公式（{number}）：{FORMULAS[key]}")
    number_paragraph = doc.add_paragraph()
    number_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    number_paragraph.paragraph_format.space_after = Pt(5)
    number_run = number_paragraph.add_run(f"（{number}）")
    set_run_font(number_run, size=9.5, color=MUTED)


def add_figure(doc, path, caption, width=6.3):
    """插入演示图，并保证题注与图片相邻。"""
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(path), width=Inches(width))
    inline_shape._inline.docPr.set("descr", caption)
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_before = Pt(2)
    caption_p.paragraph_format.space_after = Pt(8)
    caption_run = caption_p.add_run(caption)
    set_run_font(caption_run, size=9.5, color=MUTED, italic=True)


def add_code_block(doc, lines):
    """以固定宽度字体展示 MATLAB 运行命令。"""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CODE_FILL)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, latin="Consolas", east_asia="等线", size=9.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_reference(doc, number, citation, url=None):
    """添加编号参考文献和可复制的 DOI 或公开链接。"""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(f"[{number}] {citation}")
    set_run_font(run, size=9.5)
    if url:
        link_run = paragraph.add_run(f"  {url}")
        set_run_font(link_run, size=9, color=BLUE)


def add_cover(doc):
    """创建 editorial_cover 风格的首页。"""
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(84)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker_run = kicker.add_run("技术原理 · 模型审阅 · 可运行示例")
    set_run_font(kicker_run, east_asia="微软雅黑", size=11, color=BLUE, bold=True)
    kicker.paragraph_format.space_after = Pt(18)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("水下光学成像模型")
    set_run_font(title_run, east_asia="微软雅黑", size=29, color=DARK_BLUE, bold=True)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("从 Jaffe–McGlamery 三分量模型到 MATLAB 数值演示")
    set_run_font(subtitle_run, east_asia="微软雅黑", size=15, color=MUTED)

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line.paragraph_format.space_before = Pt(18)
    line.paragraph_format.space_after = Pt(68)
    line_run = line.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    set_run_font(line_run, size=10, color=BLUE)

    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.space_after = Pt(8)
    summary_run = summary.add_run("包括：完整成像链路、公式与假设、模型局限、反演方法、MATLAB 代码和实测演示图")
    set_run_font(summary_run, east_asia="微软雅黑", size=10.5, color=INK)

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_p.add_run(date.today().strftime("%Y 年 %m 月 %d 日"))
    set_run_font(date_run, size=10, color=MUTED)
    doc.add_page_break()


def build_document():
    """组织正文、公式、图片、附录和参考文献。"""
    crop_formula_images()
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    doc.add_heading("摘要与结论", level=1)
    add_callout(
        doc,
        "结论",
        "原图给出的“直接分量 + 前向散射 + 后向散射”框架在概念上正确，但不是完整的可计算物理模型。"
        "若用于算法说明，至少应补充光谱/颜色通道依赖、主动照明双程路径、空间变化散射核、后向散射路径积分、"
        "相机响应与噪声，并清楚标注均匀水体、线性 RGB、小视场等假设。",
    )
    add_text(
        doc,
        "本文在保留 Jaffe–McGlamery 三分量思想的基础上，将几何投影、水体固有光学量、直接透射、前向散射、"
        "后向散射和成像器响应串成一条完整链路。随后给出一个可直接运行的 MATLAB R2021b 示例，用合成清晰场景和"
        "逐像素深度图生成水下退化图像，并比较忽略与补偿前向散射时的恢复结果。",
    )
    add_text(doc, "文档边界：示例旨在讲清原理并验证分量关系，不是蒙特卡洛辐射传输求解器，也不用于反演真实水体的唯一参数。")

    doc.add_heading("1. 原模型完整性审阅", level=1)
    add_table(
        doc,
        ["环节", "原图已有内容", "需要补充或修正"],
        [
            ("几何成像", "针孔投影与外参 R、T", "内参应包含主点、斜切；需要说明畸变是否校正，尺度因子不能省略。"),
            ("直接分量", "Beer–Lambert 指数衰减", "衰减随波长变化；主动照明还包含光源到目标、目标到相机两段路径及几何照度。"),
            ("前向散射", "用高斯核卷积", "高斯核只是近似；真实 PSF 随距离、波长、水体和视场位置变化，严格写法是空间变积分。"),
            ("后向散射", "距离增长并趋于饱和", "应从视线方向的路径辐亮度积分出发；饱和式仅适用于近似均匀水体和简化照明。"),
            ("参数关系", "用一个消光系数描述衰减", "吸收 a、散射 b、束衰减 c=a+b 应区分；直接衰减与后向散射的宽带系数通常不相同。"),
            ("相机成像", "未显式给出", "应说明光谱响应、曝光、增益、白平衡、gamma/量化、暗电流和读出噪声。"),
        ],
        [1580, 2760, 5020],
    )
    add_text(
        doc,
        "因此，原描述适合作为“概念模型草图”，但若要支持仿真、参数估计或图像复原，需要采用下面的分层模型。",
    )

    doc.add_heading("2. 从三维目标到图像坐标", level=1)
    add_text(doc, "设空间点的齐次坐标为 P̃，像素齐次坐标为 p̃，针孔相机模型写为：")
    add_equation(doc, "f01_projection", 1, 5.0)
    add_text(doc, "其中 K 为相机内参矩阵，R、t 为世界坐标系到相机坐标系的旋转和平移，s 为齐次尺度：")
    add_equation(doc, "f02_intrinsic", 2, 3.7)
    add_bullet(doc, "f_x、f_y 是以像素为单位的等效焦距；c_x、c_y 是主点；γ 是像素轴斜切，一般可近似为 0。")
    add_bullet(doc, "若镜头畸变显著，应先应用径向/切向畸变模型，再进入水下光学退化模型。")
    add_bullet(doc, "平面玻璃舷窗或水下外壳会引入折射，精密测量时针孔模型应替换为折射相机模型。")

    doc.add_heading("3. 水下光传播与三分量成像", level=1)
    add_text(doc, "以下所有量应在相机线性响应域中定义，颜色通道 c∈{R,G,B}。总观测模型为：")
    add_equation(doc, "f03_total", 3, 5.9)
    add_text(doc, "D 为未发生散射而到达相机的直接分量，F 为目标光发生小角度前向散射形成的模糊分量，B 为水体粒子将环境光或人工光散射进视线的后向分量，n 为噪声。")

    doc.add_heading("3.1 水体固有光学量与透射率", level=2)
    add_text(doc, "吸收系数 a(λ) 和散射系数 b(λ) 的和构成束衰减系数 c(λ)：")
    add_equation(doc, "f04_iop", 4, 3.6)
    add_text(doc, "在均匀水体、单次路径且忽略多次散射的 Beer–Lambert 近似下：")
    add_equation(doc, "f05_transmission", 5, 3.6)
    add_text(doc, "红光在多数天然水体中衰减更快，因此远处目标通常偏青蓝。RGB 系数是光谱响应、照明光谱、目标反射率与水体特性的宽带合成量，不应简单当作纯水常数。")

    doc.add_heading("3.2 直接分量", level=2)
    add_text(doc, "对被动照明或已将入射照度吸收到 L₀ 的情况，颜色通道的直接信号可写成光谱积分：")
    add_equation(doc, "f06_direct_general", 6, 6.2)
    add_text(doc, "S_c(λ) 为相机第 c 通道的光谱敏感度，L₀ 是未经过目标到相机水程衰减时的目标辐亮度。对于灯光与相机分离的主动照明，更完整的近似为：")
    add_equation(doc, "f07_direct_active", 7, 6.2)
    add_text(doc, "d_s、d_c 分别为光源到目标和目标到相机的距离；G 包含余弦项、距离平方衰减、光束图案及镜头渐晕；ρ 为目标反射率。原图只写一个 d，会漏掉主动照明的入射水程。")

    doc.add_heading("3.3 前向散射分量", level=2)
    add_text(doc, "前向散射主要把邻域目标光重新分配到当前像素，表现为对比度下降、细节模糊和光晕。空间变化的形式可写为：")
    add_equation(doc, "f08_forward", 8, 6.0)
    add_text(doc, "h 是与颜色通道、距离、水体相函数和成像几何有关的点扩散函数（PSF）。只有在小视场、同一深度层且 PSF 近似平移不变时，才可简化为固定高斯卷积。MATLAB 示例采用“分层深度 + 不同 σ 的高斯核”逼近空间变化散射。")

    doc.add_heading("3.4 后向散射分量", level=2)
    add_text(doc, "后向散射来自相机与目标之间整条视线水体，而不是目标表面。一般形式是路径辐亮度积分：")
    add_equation(doc, "f09_back_integral", 9, 5.2)
    add_text(doc, "q_c(s) 汇集该位置的照明、体散射函数和几何关系。若水体均匀、照明场近似平稳，可得到常用饱和模型：")
    add_equation(doc, "f10_back_simple", 10, 5.4)
    add_text(doc, "B∞ 是无限远处的渐近水体颜色，βᴮ 是控制后向散射增长的宽带系数。Akkaynak 与 Treibitz 指出，βᴮ 不应强制等于直接信号的衰减系数 βᴰ。")

    doc.add_heading("3.5 相机响应与噪声", level=2)
    add_text(doc, "物理辐亮度经过传感器光谱响应、模拟增益、非线性映射和量化后才成为图像像素：")
    add_equation(doc, "f11_sensor", 11, 5.6)
    add_text(doc, "其中 g_c 表示曝光、增益、白平衡与相机响应曲线，Q 表示量化，η 包含光子散粒噪声、暗电流、读出噪声和压缩误差。进行物理反演前应尽可能使用 RAW 或线性化图像；直接在 sRGB/gamma 域相加减会破坏辐射度关系。")

    doc.add_heading("4. 可计算模型、参数与适用范围", level=1)
    add_table(
        doc,
        ["符号", "物理含义", "单位/范围", "获取方式"],
        [
            ("d(x)", "目标到相机距离", "m", "双目、结构光、声呐配准、SfM 或已知场景"),
            ("βᴰ_c", "直接信号宽带衰减系数", "m⁻¹，正数", "标定板多距离拟合或联合估计"),
            ("βᴮ_c", "后向散射增长系数", "m⁻¹，正数", "暗像素/空水区与距离联合拟合"),
            ("B∞,c", "渐近后向散射颜色", "线性强度 0–1", "远距离水体区域或参数拟合"),
            ("h/σ", "前向散射 PSF/尺度", "归一化核/像素", "水槽标定、点目标或经验模型"),
            ("S_c(λ)", "传感器通道光谱响应", "相对响应", "厂商数据或单色仪测量"),
        ],
        [1260, 3070, 1740, 3290],
    )
    add_text(doc, "该简化模型成立需要下列假设：")
    add_bullet(doc, "水体在局部空间内近似均匀，系数在单帧曝光期间不变。")
    add_bullet(doc, "以单次散射或弱多次散射近似为主，极浑浊水体的强多次散射不适用。")
    add_bullet(doc, "深度图与 RGB 图像已对齐，且相机图像已暗电流校正并线性化。")
    add_bullet(doc, "主动照明时光源位置、光束角和入射水程已知或被有效近似。")
    add_callout(doc, "可辨识性提醒", "单张未知深度的 RGB 图像通常不能唯一分离目标反射率、照明颜色、直接衰减与后向散射。需要深度、标定目标、多视角、偏振或可靠先验中的至少一种额外信息。")

    doc.add_heading("5. MATLAB 原理演示", level=1)
    add_text(doc, "示例以合成清晰场景 J、连续深度图 d 和每通道参数为输入，使用下式生成观测图像：")
    add_equation(doc, "f12_demo", 12, 6.2)
    add_text(doc, "其中 t_c=exp(-βᴰ_c d)，α_c 控制前向散射比例，gσ(d) 为随距离增宽的高斯核。该式便于直观展示三种分量，并不声称精确守恒每个光子的能量。")

    doc.add_heading("5.1 计算流程", level=2)
    add_number(doc, "构造包含色块、条纹和分辨率目标的清晰线性 RGB 场景，同时生成 2.4–9.8 m 的逐像素深度图。")
    add_number(doc, "按 R/G/B 三通道计算指数透射率和直接分量，红通道使用更大的衰减系数。")
    add_number(doc, "按深度层对直接分量进行高斯卷积，合成距离相关的前向散射模糊。")
    add_number(doc, "使用渐近水体颜色和独立 βᴮ 生成后向散射，再叠加小幅高斯噪声。")
    add_number(doc, "分别执行忽略前向散射的两分量恢复与使用已知前向分量的三分量恢复，计算 MAE、MSE 和 PSNR。")

    add_figure(doc, DEMO_OUTPUT_DIR / "underwater_model_overview.png", "图 1  MATLAB 实际运行的三分量成像与恢复总览", 6.25)
    add_figure(doc, DEMO_OUTPUT_DIR / "underwater_model_curves.png", "图 2  RGB 透射率、后向散射和各分量随距离变化的曲线", 6.25)

    doc.add_heading("5.2 结果解释", level=2)
    add_text(doc, "本次固定随机种子运行结果：两分量恢复 PSNR 为 32.70 dB、MAE 为 0.0195；三分量恢复 PSNR 为 39.42 dB、MAE 为 0.0058。")
    add_equation(doc, "f14_metrics", 13, 3.8)
    add_bullet(doc, "深度增大时红通道透射率最先下降，重现了典型的水下偏蓝绿色现象。")
    add_bullet(doc, "前向散射破坏高频细节；若反演中把它遗漏，恢复图仍有条纹模糊和局部光晕。")
    add_bullet(doc, "后向散射随距离增加并趋于 B∞，使远处区域对比度被水体颜色淹没。")
    add_bullet(doc, "三分量恢复采用了仿真中已知的 F、B 和 t，因此是“oracle 上界”，不代表真实单帧算法可以无误获得这些量。")

    doc.add_heading("5.3 运行方法与输出", level=2)
    add_text(doc, "在 MATLAB 命令窗口执行：")
    add_code_block(doc, [
        "cd('C:\\Users\\Administrator\\Desktop\\seamine\\matlab_underwater_demo');",
        "run_demo;",
    ])
    add_text(doc, "主程序不依赖 Image Processing Toolbox，已在 MATLAB R2021b 运行验证。输出目录包含清晰场景、深度图、水下观测图、两种恢复图、总览图、曲线图和 MAT 数据文件。")
    add_text(doc, "主要文件：underwater_optical_model_demo.m（模型与绘图）、run_demo.m（一键运行）、README.md（参数与输出说明）。")

    doc.add_heading("6. 从演示模型走向真实系统", level=1)
    add_text(doc, "若将模型用于真实相机或算法论文，建议按以下顺序扩展：")
    add_number(doc, "相机标定：完成内参、畸变、平面舷窗折射和 RAW 线性化标定。")
    add_number(doc, "几何输入：获取与 RGB 对齐的距离图，并建模光源到目标的入射路径。")
    add_number(doc, "光学参数：用多距离中性色卡或已知反射率目标分别拟合 βᴰ、βᴮ 和 B∞。")
    add_number(doc, "散射核：用点目标/细线目标在不同距离处标定空间变化 PSF，而不是固定高斯核。")
    add_number(doc, "不确定性：评估深度误差、参数耦合、噪声放大和透射率下限对恢复结果的影响。")
    add_number(doc, "验证：除 PSNR/SSIM 外，使用色差 ΔE、调制度传递函数、目标检测性能和真实水槽/海试数据。")
    add_callout(doc, "模型选择", "只做视觉增强时，可采用经验或学习方法；需要颜色定量、测量或跨水体泛化时，应优先采用光谱化、距离相关且区分 βᴰ 与 βᴮ 的物理模型。")

    doc.add_heading("附录 A：公式的 LaTeX 源码", level=1)
    add_text(doc, "正文全部公式均由下列 LaTeX 源码生成；复制到支持 LaTeX 的公式编辑器即可复用。")
    for index, source in enumerate(FORMULAS.values(), start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        label = paragraph.add_run(f"式源 {index:02d}  ")
        set_run_font(label, east_asia="微软雅黑", size=8.5, color=BLUE, bold=True)
        code = paragraph.add_run(source)
        set_run_font(code, latin="Consolas", east_asia="等线", size=8.2, color=INK)

    doc.add_heading("参考文献", level=1)
    add_reference(doc, 1, "McGlamery, B. L. A Computer Model for Underwater Camera Systems. Proceedings of SPIE, Vol. 208, 1980, pp. 221–231. DOI: 10.1117/12.958279.", "https://doi.org/10.1117/12.958279")
    add_reference(doc, 2, "Jaffe, J. S. Computer Modeling and the Design of Optimal Underwater Imaging Systems. IEEE Journal of Oceanic Engineering, 15(2), 1990, pp. 101–111. DOI: 10.1109/48.50695.", "https://doi.org/10.1109/48.50695")
    add_reference(doc, 3, "Mobley, C. D. Light and Water: Radiative Transfer in Natural Waters. Academic Press, San Diego, 1994. ISBN: 978-0-12-502750-2.", "https://openlibrary.org/books/OL1083771M/Light_and_water")
    add_reference(doc, 4, "Schechner, Y. Y.; Karpel, N. Recovery of Underwater Visibility and Structure by Polarization Analysis. IEEE Journal of Oceanic Engineering, 30(3), 2005, pp. 570–587. DOI: 10.1109/JOE.2005.850871.", "https://doi.org/10.1109/JOE.2005.850871")
    add_reference(doc, 5, "Akkaynak, D.; Treibitz, T. A Revised Underwater Image Formation Model. Proceedings of CVPR, 2018, pp. 6723–6732. DOI: 10.1109/CVPR.2018.00703.", "https://openaccess.thecvf.com/content_cvpr_2018/html/Akkaynak_A_Revised_Underwater_CVPR_2018_paper.html")
    add_reference(doc, 6, "Akkaynak, D.; Treibitz, T. Sea-Thru: A Method for Removing Water From Underwater Images. Proceedings of CVPR, 2019, pp. 1682–1691. DOI: 10.1109/CVPR.2019.00178.", "https://openaccess.thecvf.com/content_CVPR_2019/html/Akkaynak_Sea-Thru_A_Method_for_Removing_Water_From_Underwater_Images_CVPR_2019_paper.html")

    # 保存前设置核心属性，避免文档中出现临时脚本信息。
    doc.core_properties.title = "水下光学成像模型原理与 MATLAB 演示"
    doc.core_properties.subject = "Jaffe–McGlamery 三分量水下成像模型、完整性审阅与 MATLAB 示例"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "水下成像; Jaffe; McGlamery; 前向散射; 后向散射; MATLAB"
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build_document()
